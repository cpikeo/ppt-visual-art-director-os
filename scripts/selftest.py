#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""selftest.py · 核心测试（生产路径即测试路径）

四段：
  1) intelligence —— 设计判断：结论 / 权重 / 角色 / 焦点 / 构图（含否决理由）/ 媒体必要性 / 世界
  2) verify        —— 硬门：溢出、重叠、越界、空载荷、来源区、出处、口径
  3) production    —— 编译确定性、缓存复用、资产链、关键页取证
  4) cli + docs    —— 单入口 plan/check/qc、骨架携带判断、文档纪律
"""
from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent
ROOT = SCRIPTS.parent
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

PASS = 0
FAIL = 0


def check(name, ok, detail=""):
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"[PASS] {name}")
    else:
        FAIL += 1
        print(f"[FAIL] {name}" + (f" — {detail}" if detail else ""))


def run_vao(*args, cwd=None, env=None):
    # 显式 UTF-8：子进程输出恒为 UTF-8（vao.force_utf8_stdio），父进程必须按同一
    # 口径解码——用默认 text=True 时 Windows 会拿系统 ANSI 代码页解码管道，
    # 于是断言读到的是乱码或直接抛 UnicodeDecodeError。
    return subprocess.run([sys.executable, str(SCRIPTS / "vao.py"), *args],
                          capture_output=True, text=True, encoding="utf-8",
                          errors="replace", cwd=cwd or str(ROOT), env=env)


# ── fixtures ─────────────────────────────────────────────────────────
BRIEF = {
    "audience": "董事会与投资人",
    "decision": "批准 2026 自有内容预算",
    "tension": "投放成本上升，自有内容回报周期长",
    "quality_level": "advanced",
    "brand_colors": ["#1A3A5C", "#C8501E"],
    "slides": [
        {"id": "s01", "title": "自有内容占比过半", "content": "2026 内容战略提案"},
        {"id": "s02", "title": "投入自有内容是唯一防守",
         "content": "外部流量价格三年翻倍（来源：投放台账）"},
        {"id": "s03", "title": "自有内容占比 61%",
         "content": "61%，对比去年 38%（来源：内容资产台账）",
         "chart": {"type": "kpi", "value": 61}},
        {"id": "s04", "title": "18 个月落地节奏"},          # 证据页缺 content
        {"id": "s05", "title": "请批准 4,800 万专项",
         "content": "首期 1,200 万于 2027Q1 启动"},
    ],
}


def make_build(tmp: Path, *, chart: bool = True) -> Path:
    """生成一份合规 build.py（硬门全过的最小稿件）。"""
    colors = {"background": "#F2EDE4", "surface": "#F4EFE6", "primary": "#1A3A5C",
              "secondary": "#9A8C74", "accent": "#C8501E", "ink": "#191510",
              "muted": "#6E6A5E"}
    slides = []
    for i in range(5):
        els = []
        if i > 0:
            els.append({"id": "eyebrow", "type": "text", "x": 48, "y": 40, "width": 320,
                        "height": 24, "text": f"PAGE {i + 1}", "size": 12.5,
                        "color": "muted", "role": "eyebrow"})
            els.append({"id": "page_number", "type": "text", "x": 1208, "y": 40,
                        "width": 32, "height": 24, "text": str(i + 1), "size": 12.5,
                        "color": "muted", "role": "page_number", "align": "right"})
        els.append({"id": "title", "type": "text", "x": 48, "y": 96, "width": 1184,
                    "height": 56, "text": f"第{i + 1}页结论句", "size": 32,
                    "color": "ink", "role": "title", "bold": True})
        els.append({"id": "src", "type": "text", "x": 48, "y": 676, "width": 1100,
                    "height": 24, "text": "来源：测试数据 · 2025-12", "size": 12.5,
                    "color": "muted", "role": "source"})
        if chart and i == 2:
            els.append({"id": "chart_main", "type": "chart", "chart_kind": "line",
                        "x": 48, "y": 232, "width": 1184, "height": 400, "role": "focus",
                        "data": [{"label": "2023", "value": 38}, {"label": "2024", "value": 47},
                                 {"label": "2025", "value": 61}],
                        "highlight": 2, "show_values": True, "metric": "自有内容占比",
                        "source": "来源：内容资产台账", "unit": "%",
                        "period": "2023–2025", "basis": "全渠道加权"})
        slides.append({"id": f"s{i + 1:02d}",
                       "page_intent": {"insight": f"第{i + 1}页结论", "focus": "title"},
                       "source_zone": {"x": 48, "y": 672, "width": 1184, "height": 32},
                       "elements": els})
    build = tmp / "build.py"
    build.write_text(
        "# -*- coding: utf-8 -*-\nSPEC = {\n"
        f"  'canvas': {{'width': 1280, 'height': 720, 'grid_columns': 12, 'grid_unit': 8}},\n"
        f"  'theme': {{'colors': {colors!r}, 'fonts': {{'cn': 'Source Han Serif SC', 'latin': 'Georgia'}}}},\n"
        f"  'slides': {slides!r},\n}}\n", encoding="utf-8")
    return build


# ── 1 · intelligence（设计判断）───────────────────────────────────────
def test_intelligence(tmp: Path):
    import intelligence as intel

    # 固定关系正反样本就地保留，核心回归不依赖外部评测目录。
    relation_cases = [
        {"id": "previous-round-governance", "title": "四层汇报线明确决策归属", "content": "董事会决定战略；管理层负责预算；区域负责人协调资源；当地负责人执行。", "expect_relations": 3, "expect_role": "explain", "expect_composition": "linear_structure"},
        {"id": "responsibility", "title": "研发与交付的职责应分开", "content": "甲组负责研发；乙组负责交付。", "expect_relations": 2, "expect_role": "explain", "expect_composition": "linear_structure"},
        {"id": "reporting", "title": "审批归属必须清楚", "content": "区域负责人向总经理汇报；总经理向董事会汇报。", "expect_relations": 2, "expect_role": "explain", "expect_composition": "linear_structure"},
        {"id": "renamed-subjects", "title": "职责先于装饰", "content": "青岚组负责质量；澄川组负责交付。", "expect_relations": 2, "expect_role": "explain", "expect_composition": "linear_structure"},
        {"id": "same-subject", "title": "专人承担研发和交付", "content": "甲组负责研发；甲组负责交付。", "expect_relations": 0},
        {"id": "question", "title": "职责尚待确认", "content": "谁负责研发？谁负责交付？", "expect_relations": 0},
        {"id": "negation", "title": "当前职责仍未落实", "content": "甲组不负责研发；乙组未负责交付。", "expect_relations": 0},
        {"id": "proposal", "title": "分工还在讨论", "content": "建议甲组负责研发；乙组负责交付。", "expect_relations": 0},
        {"id": "unresolved-objects", "title": "分工还不完整", "content": "甲组负责；乙组执行。", "expect_relations": 0},
        {"id": "numeric-evidence", "title": "本期执行结果可核验", "content": "甲组执行 12 项；乙组执行 18 项。", "expect_relations": 0},
        {"id": "sequence", "title": "先研发再交付", "content": "先由甲组负责研发，然后由乙组负责交付。", "expect_relations": 0},
        {"id": "comparison", "title": "两组的交付率有差异", "content": "甲组负责研发；乙组负责交付。甲组交付率 90%，高于乙组 80%。", "expect_relations": 0, "expect_role": "compare"},
        {"id": "topic-only", "title": "四层汇报线与矩阵式分工", "content": "", "expect_relations": 0},
        {"id": "subset-control", "title": "多数试点选择续约", "content": "12 家试点门店中，9 家续约。", "expect_relations": 0, "expect_role": "prove"},
        {"id": "compound-control", "title": "结构性增长来自续约", "content": "我们必须保留这一增长来源。", "expect_relations": 0},
    ]
    for case in relation_cases:
        j = intel.page_judgment({**case, "asset": "none"}, index=2, total=7,
                                brief=BRIEF, world={}, warnings=[])
        actual = (len(j["understanding"].get("relations", [])), j["visual_role"]["role"],
                  j["composition"]["chosen"])
        check("relations: " + case["id"],
              actual == (case["expect_relations"], case.get("expect_role", actual[1]),
                         case.get("expect_composition", actual[2])), str(actual))
    relation_page = {**relation_cases[0], "asset": "none"}
    checked_relations = intel.think(dict(BRIEF, slides=[
        dict(relation_page, id=f"r{i}") for i in range(4)]))
    check("relations: 首末位置/跨页求变不能覆盖正文关系证据",
          all(p["judgment"]["visual_role"]["role"] == "explain"
              and p["judgment"]["composition"]["chosen"] == "linear_structure"
              for p in checked_relations["pages"]))
    rj = checked_relations["pages"][0]["judgment"]
    check("relations: 分工不伪装成汇报层级或时间轴",
          rj["focus"]["element_role"] == "structure_map"
          and "空白不代表时间" in rj["spatial"]["whitespace_duty"]
          and "推测汇报线" in rj["typography"]["discipline"])
    declared_relation = intel.think(dict(BRIEF, slides=[
        dict(relation_page, composition="editorial", focus="author_focus")]))["pages"][0]["judgment"]
    check("relations: 作者焦点与构图仍然优先",
          declared_relation["composition"]["chosen"] == "editorial"
          and declared_relation["focus"]["element_role"] == "author_focus")
    real_sequence = intel.page_judgment({"title":"按阶段落地", "content":"先研发，然后交付。", "asset":"none"},
                                       index=1,total=3,brief=BRIEF,world={},warnings=[])
    check("relations: 时间推进仍保留顺序空间职责",
          real_sequence["understanding"]["evidence"] == "sequence"
          and "空白是时间" in real_sequence["spatial"]["whitespace_duty"])

    bundle = intel.think(BRIEF)
    deck, pages = bundle["deck"], bundle["pages"]
    by_id = {p["id"]: p for p in pages}

    check("判断链：每页都有完整判断卡",
          all({"claim", "information_weight", "visual_role", "focus", "composition",
               "spatial", "media", "production"} <= set(p["judgment"]) for p in pages))

    # 结论抽取：declared 优先 / 标题优于署名行 / 缺内容 = absent
    s01 = by_id["s01"]["judgment"]
    check("claim: 封面取标题结论而不是署名行",
          s01["claim"]["text"] == "自有内容占比过半", str(s01["claim"]))
    check("type: 主语只取可排版的主张，不把权重判断理由贴到页面上",
          s01["typography"]["lead"] == s01["claim"]["text"]
          and "结论句" not in s01["typography"]["lead"]
          and "字距" in s01["typography"]["discipline"])
    s04 = by_id["s04"]["judgment"]
    check("claim: 证据页缺 content → absent，不脑补",
          s04["claim"]["source"] == "absent" and s04["claim"]["text"] is None)
    declared = intel.think(dict(BRIEF, slides=[{"id": "x", "title": "话题",
                                                "content": "数据",
                                                "insight": "作者写的结论"}]))
    check("claim: 作者声明的结论压过抽取",
          declared["pages"][0]["judgment"]["claim"]["source"] == "declared")

    # 信息权重：结论数字最大化、对照基数弱化、重复删除
    s03 = by_id["s03"]["judgment"]
    maxi = " ".join(s03["information_weight"]["maximize"])
    weak = " ".join(s03["information_weight"]["weaken"])
    check("weight: 结论数字最大化（61%）", "61%" in maxi)
    check("weight: 对照基数弱化（38% 小一档）", "38%" in weak)

    # 先出现的数字未必是结论：焦点、权重、排版必须共用同一结果值判定。
    rising = intel.think(dict(BRIEF, slides=[
        {"id": "rise", "title": "内容占比从 38% 升至 61%",
         "content": "2023 年 38%，2025 年 61%（来源：内容台账）"}]))["pages"][0]["judgment"]
    falling = intel.think(dict(BRIEF, slides=[
        {"id": "fall", "title": "单位成本从 2,800 元降至 1,800 元",
         "content": "同口径（来源：财务台账）"}]))["pages"][0]["judgment"]
    check("data: 升至/降至后的结果值为主，基期退后（不让首个数字误导读者）",
          rising["visual_role"]["role"] == "prove"
          and rising["understanding"]["evidence"] == "series"
          and "61%" in rising["focus"]["why"]
          and rising["typography"]["lead"] == "61%"
          and any("38%" in x for x in rising["information_weight"]["weaken"])
          and "1,800 元" in falling["focus"]["why"]
          and falling["typography"]["lead"] == "1,800 元",
          f"{rising['information_weight']} / {falling['information_weight']}")
    check("data: 量化箭头是变化，文字阶段箭头仍是序列；复合单位不被截断",
          intel.understand("38% → 61%")["evidence"] == "series"
          and intel.understand("试点 → 复制 → 全量")["evidence"] == "sequence"
          and intel.understand("从38%到61%")["evidence"] == "series"
          and intel.numerals("18个月 / 4,800万元 / 30分钟")[0]["unit"] == "个月"
          and intel.numerals("18个月 / 4,800万元 / 30分钟")[1]["unit"] == "万元")
    review = intel.think(dict(BRIEF, slides=[
        {"id": "compare", "title": "留存改善比曝光增长更稳健",
         "content": "内容回访率 71%，投放再触达率 42%（来源：回访台账）"},
        {"id": "subset", "title": "试点门店愿意继续共创",
         "content": "12 家试点门店中，9 家承诺提供记录（来源：回访台账）",
         "asset": "none"},
        {"id": "ask", "title": "请批准 4,800 万，启动 18 个月计划",
         "content": "请求批准 4,800 万，分季度验收（来源：预算草案）"}]))
    jc, jp, ja = [p["judgment"] for p in review["pages"]]
    check("data: 两个不同指标是比较，不伪装成同一指标的前后趋势/差值",
          jc["understanding"]["evidence"] == "comparison"
          and jc["visual_role"]["role"] == "compare"
          and "分母" in jc["focus"]["why"]
          and "基期" not in " ".join(jc["information_weight"]["weaken"])
          and jc["focus"]["type"] == "text"
          and any("分母" in q for q in jc["open_questions"]))
    chart_pair = intel.think(dict(BRIEF, slides=[
        {"id": "paired", "title": "两个样本的绝对量",
         "content": "甲组 30 人，乙组 18 人（来源：用户台账）",
         "chart": {"type": "bar", "values": [30, 18]}}]))["pages"][0]["judgment"]
    check("data: 作者给出的两个数图表是比较角色，不把第一数字误当趋势结果",
          chart_pair["visual_role"]["role"] == "compare"
          and chart_pair["focus"]["element_role"] == "comparison_field"
          and chart_pair["focus"]["type"] == "chart")
    mixed = intel.think(dict(BRIEF, slides=[
        {"id": "units", "title": "收入与人均成本对照",
         "content": "总收入 480 万元，人均成本 90 元（来源：财务表）",
         "chart": {"type": "bar", "values": [480, 90]}}]))["pages"][0]["judgment"]
    check("data: 单位不同不共用数量轴；曲线页编号不是第四个数据点",
          not mixed["understanding"]["same_unit"]
          and mixed["composition"]["chosen"] != "data_field"
          and any("单位不同" in q for q in mixed["open_questions"])
          and intel.understand("三年曲线1 96元→128元→205元")["distinct_numerals"] == 3)
    check("data: 9/12 子集突出 9、写明 12 分母，不误读为下降或摆并列卡",
          jp["understanding"]["evidence"] == "part_whole"
          and jp["focus"]["element_role"] == "ratio_main"
          and jp["typography"]["lead"] == "9 家"
          and jp["typography"]["support"] == "12 家"
          and jp["composition"]["chosen"] == "big_whitespace"
          and any(x["role"] == "claim_text" for x in jp["production"]["must_place"]))
    check("data: 预算和月份是两种量，18个月不是金额的历史基期",
          ja["typography"]["lead"] == "4,800 万"
          and ja["typography"]["support"] == "18 个月"
          and ja["focus"]["type"] == "text"
          and any("执行期限" in w for w in ja["information_weight"]["weaken"])
          and any(x["role"] == "claim_text" for x in ja["production"]["must_place"]))
    check("editorial: 不可比指标不共刻度，部分/总体不画趋势，请求金额/期限分层",
          "同口径" in jc["typography"]["discipline"]
          and "轴" in jp["typography"]["discipline"]
          and "期限" in ja["typography"]["discipline"])
    dup_text = "收入 12 亿，收入 12 亿（来源：财报）"
    dup = intel.information_weight(intel.understand(dup_text),
                                   {"text": "收入 12 亿", "source": "extracted"}, dup_text)
    check("weight: 同数字重复出现进删除清单",
          any("重复" in x for x in dup["delete"]), str(dup["delete"]))

    # 两组「各自从 A 到 B」的变化并排（证据表的一行）不再顶掉标题结论位：
    # 此前这句靠「变化词 + 数字」加分胜出，焦点与排版跟着落到基期数字上。
    two_moves = intel.think({
        "audience": "董事会", "decision": "批准投入",
        "slides": [{"id": "p1", "title": "我们的份额增速正在被本地玩家反超",
                    "content": "过去 12 个月本地头部玩家份额从 12% 升至 19%，我方从 9% 降至 7%；"
                               "获客成本我方 42 美元、对手 26 美元。（来源：第三方市场监测）"}]})
    two_moves_claim = two_moves["pages"][0]["judgment"]["claim"]
    check("claim: 两组变化并排是列举，作者写在标题位的结论仍然胜出",
          two_moves_claim["text"] == "我们的份额增速正在被本地玩家反超"
          and "标题" in two_moves_claim["why"], str(two_moves_claim))

    # 顺序形状由顺序词的数量决定：三段里程碑排期不应该被一个孤立的时长数字压成 number。
    plan_page = intel.think({
        "audience": "董事会", "decision": "批准排期",
        "slides": [{"id": "p1", "title": "必须在 2027 年一季度前完成三国落地",
                    "content": "先做基础设施与合规，然后复制内容与供应链，接着本地化运营团队，"
                               "最后做区域品牌投放。总计 18 个月，分三段里程碑。（来源：战略部排期）"}]})
    plan_judgment = plan_page["pages"][0]["judgment"]
    check("understand: 三段以上顺序词 → sequence（骨架走线性结构，而不是单数字大留白）",
          plan_judgment["understanding"]["evidence"] == "sequence"
          and plan_judgment["visual_role"]["role"] == "explain"
          and plan_judgment["composition"]["chosen"] == "linear_structure",
          str({"evidence": plan_judgment["understanding"]["evidence"],
               "role": plan_judgment["visual_role"]["role"],
               "composition": plan_judgment["composition"]["chosen"]}))

    # 视觉角色：开场 establish / 收尾 summarize / 对比 compare / 说服 persuade
    roles = {p["id"]: p["judgment"]["visual_role"]["role"] for p in pages}
    check("role: 开场 establish、收尾 summarize",
          roles["s01"] == "establish" and roles["s05"] == "summarize", str(roles))
    check("role: 对比内容 → compare，且给了替代角色",
          roles["s03"] == "compare" and by_id["s03"]["judgment"]["visual_role"]["alternatives"])

    # 焦点：焦点元素唯一 + 必须给出否决的替代项
    foc = s03["focus"]
    check("focus: 对比页焦点是两端对比场而不是卡片墙",
          foc["element_role"] == "comparison_field"
          and any("卡片" in r["why_not"] or "卡" in r["option"] for r in foc["rejected"]),
          str(foc["rejected"])[:120])
    check("focus: 数字结论页焦点是 kpi_main 并否决环形图",
          by_id["s05"]["judgment"]["focus"]["element_role"] in ("kpi_main", "hero_statement"))

    # 构图：必须回答「为什么不是别的形式」
    comp = by_id["s02"]["judgment"]["composition"]
    check("composition: 主张页 → big_whitespace 且否决项有效（只留真实竞争）",
          comp["chosen"] == "big_whitespace" and len(comp["rejected"]) <= 2
          and all(r["why_not"] and r["option"] != comp["chosen"]
                  for r in comp["rejected"]))
    check("composition: 数据证据页 → data_field",
          by_id["s03"]["judgment"]["composition"]["chosen"] in ("data_field",
                                                                "asymmetric_tension"))

    # 空间：路径/主区/留白职责非空
    sp = by_id["s02"]["judgment"]["spatial"]
    check("spatial: 阅读路径 + 留白职责齐备",
          bool(sp["reading_path"] and sp["whitespace_duty"] and sp["quiet_zone"]))

    # 媒体必要性：作者声明 > 图表抑制 > 必要性测试
    m_author = intel.media_necessity(intel.understand("产品发布会现场"),
                                     "establish", declared="required")
    check("media: 作者声明压过判断", m_author["decision"] == "required"
          and m_author["source"] == "author")
    m_chart = intel.media_necessity(intel.understand("客户现场 12 家"), "prove",
                                    has_chart=True)
    check("media: 页内已有图表 → 不出图（双焦点竞争）",
          m_chart["decision"] == "none")
    m_decor = intel.media_necessity(intel.understand("预算分配 60/25/15"), "summarize")
    check("media: 说不出「没有它会下降在哪」→ 不出图",
          m_decor["decision"] == "none" and "注意力" in m_decor["necessity"])
    m_subject = intel.media_necessity(intel.understand("产品发布：新一代设备与芯片"),
                                      "establish")
    # 词表只提示「有实体」，图像是否成立由作者决定：未声明 → hint（不出图、问作者）；声明 → 出图
    check("media: 开场有可指认实体，但未声明 → 词表提示，不自动出图",
          m_subject["decision"] == "none" and m_subject["source"] == "hint"
          and m_subject["suggested"] == "subject", str(m_subject))
    m_subject_decl = intel.media_necessity(intel.understand("产品发布：新一代设备与芯片"),
                                           "establish", declared="required")
    check("media: 作者声明 asset: required → 出图（来源 author）",
          m_subject_decl["decision"] == "required" and m_subject_decl["source"] == "author",
          str(m_subject_decl))

    # 判断质量：话题 ≠ 结论；结论句要被抽出来；图像只在「有现场」时成立
    topic = intel.extract_claim(
        {"title": "店长回访记录",
         "content": "12 家试点门店的店长在高峰时段如何使用自有内容（来源：回访记录）"})
    check("claim: 话题句（如何使用…）不当结论，宁可 absent",
          topic["source"] == "absent" and topic["text"] is None, str(topic))
    asserted = intel.extract_claim(
        {"title": "成本结构",
         "content": "单篇生产成本降至 1,800 元，低于去年同期（来源：财务台账）"})
    check("claim: 断言句（降至…）被真正提炼为结论",
          asserted["source"] == "extracted" and "1,800" in asserted["text"], str(asserted))
    info_only = intel.media_necessity(
        intel.understand("获客单价 96 元 → 128 元 → 205 元"), "prove")
    scene = intel.media_necessity(
        intel.understand("门店回访现场：店长在高峰时段的记录"), "prove")
    scene_decl = intel.media_necessity(
        intel.understand("门店回访现场：店长在高峰时段的记录"), "prove", declared="required")
    check("media: 纯数据页不出图；现场词只提示证词，作者声明后才出证词图",
          info_only["decision"] == "none" and scene["decision"] == "none"
          and scene["source"] == "hint" and scene["suggested"] == "witness"
          and scene_decl["decision"] == "required" and scene_decl["source"] == "author",
          f"{info_only} / {scene} / {scene_decl}")

    # 视觉世界：由内容推导，两个同类 deck 允许得到不同世界
    fin_human = dict(BRIEF, slides=[
        {"id": "s1", "title": "试点客户续约率 92%", "content": "12 家门店客户回访现场（来源：台账）"},
        {"id": "s2", "title": "门店真实反馈", "content": "客户的日常使用场景（来源：回访）"}])
    fin_data = dict(BRIEF, slides=[
        {"id": "s1", "title": "单位获客成本 128 元", "content": "口径：全渠道加权（来源：投放台账）"},
        {"id": "s2", "title": "三年成本曲线", "content": "2023–2025 逐年上升（来源：投放台账）"}])
    w_human = intel.think(fin_human)["deck"]["world"]
    w_data = intel.think(fin_data)["deck"]["world"]
    check("world: 世界由内容推导（不是风格预设查表）",
          w_human["source"].startswith(("subject", "neutral"))
          and w_human["visual_world"] and w_human["accent_why"])
    check("world: 同类内容允许产生不同世界（实体不同 → 材质/强调不同）",
          (w_human["name"], w_human["accent"]) != (w_data["name"], w_data["accent"])
          or w_human["chart_style"] != w_data["chart_style"], str(w_human["name"]))
    check("world: 每个判断都带理由", bool(w_human["reasons"].get("subject")))
    check("plan: 页级不重复存储（角色/证据只住判断卡，判断卡里不再回存原文与向量）",
          all("role" not in p and "evidence" not in p for p in pages)
          and all("declarations" not in p["judgment"] for p in pages)
          and all("text" not in p["judgment"]["understanding"]
                  and "modes" not in p["judgment"]["understanding"] for p in pages))
    # 资产：图像只由作者声明产生；词表命中只是 hint（不占额度、不出图）；声明的图全部进入清单
    decl = intel.think(dict(BRIEF, slides=[
        {"id": "s01", "title": "田间到仓的 48 小时", "content": "采收现场与冷链记录（来源：产地日报）", "asset": "required"},
        {"id": "s02", "title": "仓库现场的复盘会", "content": "分拣现场记录：9 个环节耗时（来源：现场复盘）", "asset": "required"},
        {"id": "s03", "title": "车间班组的交接记录", "content": "车间现场记录（来源：班组日志）", "asset": "required"},
        {"id": "s04", "title": "客户回访现场", "content": "12 家客户回访现场（来源：回访记录）", "asset": "reuse"},
        {"id": "s05", "title": "门店走访的真实反馈", "content": "华东 6 家门店走访记录（来源：走访纪要）"},
    ]))
    by_decl = {p["id"]: p["judgment"]["media"] for p in decl["pages"]}
    check("assets: 词表命中但未声明的页不进入图像清单（hint，问作者）",
          by_decl["s05"]["source"] == "hint" and "s05" not in decl["assets_hint"]["generate"],
          str(by_decl["s05"]))
    check("assets: 作者声明的出图全部进入清单，不被自动额度截断（4 张 > fast 额度 2）",
          decl["assets_hint"]["generate"] == ["s01", "s02", "s03", "s04"],
          str(decl["assets_hint"]))
    check("assets: 清单不再保留预算字段（deferred/cap 已随判断删除）",
          "deferred" not in decl["assets_hint"] and "cap" not in decl["assets_hint"])
    check("plan: 世界只留会被消费的键（无展示用常量/重复串）",
          not {"material", "motion", "chroma", "regime", "background",
               "rejected_worlds"} & set(w_human))

    # 构图：算子由内容形状决定（不是按页面类型查表），且与媒体/焦点自洽
    shaped = intel.think(dict(BRIEF, slides=[
        {"id": "s0", "title": "开场", "content": "内容战略（来源：档案）"},
        {"id": "a", "title": "把工艺从 7 步压到 3 步", "content": "采青 → 杀青 → 揉捻（来源：工艺卡）"},
        {"id": "b", "title": "复购 3.4 次", "content": "3.4 次/月（来源：门店台账）"},
        {"id": "c", "title": "激活率对比", "content": "44% → 71%（来源：试点报告）"}]))
    picked = {p["id"]: p["judgment"]["composition"]["chosen"] for p in shaped["pages"]}
    check("composition: 图像页构图即图像叙事（与焦点自洽）",
          all(p["judgment"]["composition"]["chosen"] == "image_narrative"
              for p in shaped["pages"]
              if p["judgment"]["media"]["decision"] == "required"
              and p["judgment"]["focus"]["type"] == "image"))
    check("composition: 内容形状不同 → 算子不同（序列/单数/对比）",
          picked["a"] == "linear_structure" and picked["b"] == "big_whitespace"
          and picked["c"] in ("asymmetric_tension", "data_field")
          and len(set(picked.values())) >= 3, str(picked))
    check("composition: 没有图像时不选图像叙事（判断不自相矛盾）",
          all(not (p["judgment"]["composition"]["chosen"] == "image_narrative"
                   and p["judgment"]["media"]["decision"] == "none")
              for p in shaped["pages"]))
    dup_num = intel.understand("复购 3.4 次 3.4 次/月（来源：台账）")
    check("understanding: 标题重复正文的数字只算一个证据",
          dup_num["evidence"] == "number" and dup_num["distinct_numerals"] == 1,
          str(dup_num["evidence"]))

    # 品牌色：一步覆盖世界种子
    seed = deck["theme"]
    check("theme: 品牌色落语义槽（primary→information, accent→accent）",
          seed["colors_seed"]["information"] == "#1A3A5C"
          and seed["colors_seed"]["accent"] == "#C8501E"
          and seed["seed_source"] == "brand_colors", str(seed))
    tokens = intel.theme_tokens({"foundation": "#FFFFFF", "information": "#FFFFFF",
                                 "supporting": "#888888", "accent": "#CC0000"})
    check("theme: 可读性兜底（墨色与纸面分得开）", tokens["ink"] in ("#141414", "#FFFFFF"))

    # 退役字段：写 design_direction 会被点名而不是静默套用
    w = intel.think(dict(BRIEF, design_direction="editorial"))["warnings"]
    check("plan: design_direction 退役并留痕",
          any(x.get("rule") == "direction_retired" for x in w), str(w)[:120])

    # 缺 deck 契约 / 证据页缺内容：两种警告都要开口
    rules = {x.get("rule") for x in bundle["warnings"]}
    check("plan: 缺 tension 与证据页缺 content 都有警告",
          "unresolved_content" in rules and "deck_contract" not in rules)

    # 骨架：判断连理由进注释（否决项不住这里，canonical owner 是 plan）；同 bundle 必得同文本
    sk1, sk2 = intel.build_skeleton(bundle), intel.build_skeleton(bundle)
    check("skeleton: 同 bundle 必得同文本", sk1 == sk2)
    check("skeleton: 注释携带结论/权重/角色/构图（不再复制否决项）",
          "结论[" in sk1 and "权重:" in sk1 and "角色:" in sk1
          and "构图:" in sk1 and "否决" not in sk1 and "媒体:" in sk1)
    check("skeleton: 不给坐标（几何留给作者）",
          "估算高度" not in sk1 and '"x":' not in sk1.split("slides")[0])
    sk_vars: dict = {}
    exec(compile(sk1, "judgment-skeleton.py", "exec"), sk_vars)
    check("skeleton: 统一字族由世界判断直接落到稿件，终稿默认全页取证",
          sk_vars["SPEC"]["theme"]["fonts"] == bundle["deck"]["world"]["fonts"]
          and "--speed strict" in sk1)
    long_source = "首句是有来源的事实。" + "信息不会为骨架长度而丢失。" * 25 + "\n末句：留给排版处理。"
    preserved = intel.think(dict(BRIEF, slides=[
        {"id": "raw", "title": "完整原文必须保留", "content": long_source}]))
    preserved_sk = intel.build_skeleton(preserved)
    check("skeleton: 多行/超长原文逐字留在可执行骨架，原文不是视觉主语",
          long_source in preserved["pages"][0]["ref"]
          and "首句是有来源的事实。" in preserved_sk
          and "末句：留给排版处理。" in preserved_sk
          and "…" not in preserved["pages"][0]["ref"]
          and compile(preserved_sk, "preserved.py", "exec") is not None)
    check("skeleton: 无图不添加空资产清单命令（含图命令交由 plan 决定）",
          "--assets-manifest" not in preserved_sk and "媒体:none" in preserved_sk)
    four_pages = intel.think(dict(BRIEF, slides=[
        {"id": f"p{i}", "title": f"第{i}页结论句", "content": f"证据 {i}（来源：台账）"}
        for i in range(1, 5)]))
    four_sk = intel.build_skeleton(four_pages)
    four_vars: dict = {}
    exec(compile(four_sk, "anchor-skeleton.py", "exec"), four_vars)
    anchors = [s.get("anchor") for s in four_vars["SPEC"]["slides"]]
    role_1 = four_pages["pages"][1]["judgment"]["visual_role"]["role"].upper()
    check("skeleton: 跨页锚以判断对象落进 spec（不是散文提醒），spec 不留第二个结论位",
          isinstance(anchors[1], dict) and anchors[1].get("page_number") == 2
          # 标题「第N页结论句」派不出主题：眉标省略，而不是输出角色词或英文枚举
          # （枚举 12 页里会有 7 页同为 PROVE；角色词全稿重复，读起来是模板）
          and "eyebrow" not in anchors[1] and anchors[1].get("eyebrow") != role_1
          and (anchors[0] is None or "page_number" not in anchors[0])
          and all("page_intent" not in s for s in four_vars["SPEC"]["slides"]),
          str(anchors))
    # 骨架提速：连续同动作收引用（信息不丢、可回溯）；落稿写法习语进头部
    repeat_deck = intel.think(dict(BRIEF, slides=[
        {"id": f"q{i}", "title": f"同类证据页{i}",
         "content": "履约成本 11.2 元，同比升 3%（来源：台账）"}
        for i in range(1, 7)]))
    rep_sk = intel.build_skeleton(repeat_deck)
    refs = re.findall(r"# …(.+?)同 (\S+)", rep_sk)
    ref_ids = re.findall(r"# ── (\S+) ·", rep_sk)
    check("skeleton: 连续同动作收引用（头页全行保留、引用可回溯、原文不丢）",
          len(refs) >= 3 and all(r in ref_ids for _, r in refs)
          and rep_sk.count("内容原文:") == 6
          and compile(rep_sk, "repeat-skeleton.py", "exec") is not None, str(refs)[:120])
    check("skeleton: 落稿写法（信封可压缩、判断逐页显式）写进骨架头部",
          "落稿写法" in rep_sk.split("SPEC")[0] and "逐页显式给值" in rep_sk)

    # 眉标：内容派生 + 作者声明压过推导 + 不输出半截词
    eyebrow_deck = intel.think(dict(BRIEF, slides=[
        {"id": f"e{i}", "title": f"东南亚市场{i}是我们唯一的结构性增长来源",
         "content": "履约成本 11.2 元（来源：自营财务口径）"} for i in range(1, 5)]))
    labels = [(p.get("anchor") or {}).get("eyebrow") for p in eyebrow_deck["pages"]]
    declared_eyebrow = intel.think(dict(BRIEF, slides=[
        {"id": "x1", "title": "标题", "content": "证据", "eyebrow": "单位经济"},
        {"id": "x2", "title": "标题", "content": "证据", "eyebrow": "清关与用工"},
        {"id": "x3", "title": "标题", "content": "证据", "eyebrow": "释放条件"},
        {"id": "x4", "title": "标题", "content": "证据", "eyebrow": "复盘节点"}]))
    check("眉标: 由本页结论派生（不是角色枚举），作者声明原样生效",
          all(lbl and not lbl.isupper() and len(lbl) <= 6 for lbl in labels)
          and len(set(labels)) == 1 and labels[0] == "东南亚市场"
          and [(p.get("anchor") or {}).get("eyebrow") for p in declared_eyebrow["pages"]]
          == ["单位经济", "清关与用工", "释放条件", "复盘节点"], str(labels))
    # 眉标不得**缝合数字两侧**：把数字删掉再断词，会把「中国 800 元以上」缝成
    # 「中国元以」——它看着像有信息，实际是噪声，比退回角色词更糟。
    # 边缘编号（东南亚市场1）可以去掉，夹在中间的数字一律不收。
    splice_bad = None
    for _t, _want in (("中国 800 元以上的葡萄酒，是唯一两位数增长的赛道", "葡萄酒"),
                      ("东南亚市场1是我们唯一的结构性增长来源", "东南亚市场"),
                      ("本轮 1.8 亿元里，67% 投向库存与藤园", "")):
        _got = intel.section_label({}, {"text": _t}, "prove")
        if _got != _want:
            splice_bad = (_t, _want, _got); break
    # 派不出主题就省略眉标：角色词在全稿重复，读起来是模板（见 section_label）
    check("眉标: 不缝合数字两侧的汉字；派不出主题时省略，不退回角色词",
          splice_bad is None, str(splice_bad))
    # 封面 / 收尾是**陈述位**：没有证据不等于没有主张。
    # 实测（CASE_006）：一句封面主张被判成「结构说明页」，焦点落到一张并不存在的
    # 层级图上，收尾页则被判 absent 并告警，逼作者给封面补数据。
    statement_deck = intel.think(dict(BRIEF, slides=[
        {"id": "cover", "title": "东南亚是我们未来五年唯一还能拿到结构性增长的战场"},
        {"id": "mid", "title": "复购率由 21.4% 升至 41.7%",
         "content": "同期补贴占收入比由 22.4% 降至 6.1%（来源：运营系统）"},
        {"id": "close", "title": "我们不是在赌一个市场，是在买一条可以复制到第三个国家的履约能力"}]))
    cov, clo = (statement_deck["pages"][0]["judgment"],
                statement_deck["pages"][-1]["judgment"])
    check("首末页: 陈述位按叙事位置取角色，不派发结构图、不误报告警",
          cov["visual_role"]["role"] == "establish"
          and clo["visual_role"]["role"] == "summarize"
          and cov["focus"]["element_role"] == "statement_text"
          and clo["focus"]["element_role"] != "structure_map"
          and cov["claim"]["text"].startswith("东南亚")
          and clo["claim"]["text"].startswith("我们不是在赌")
          and not [w for w in statement_deck["warnings"]
                   if w.get("rule") == "unresolved_content"],
          str([cov["visual_role"]["role"], clo["visual_role"]["role"],
               cov["focus"]["element_role"], clo["focus"]["element_role"]]))
    check("understanding: 「结构性增长」是复合名词，不产生结构/序列证据形态",
          intel.understand("东南亚是我们未来五年唯一还能拿到结构性增长的战场")["evidence"]
          in ("none", "proposition"))
    # 缺口的归因必须指向真实根因：内容缺失 vs 有内容但不成主张，两种修法不同
    #（实测：预算构成页标题是话题标签、正文完整，却被报「content 缺失」）。
    # 首末页是陈述位（没有 content 也按作者写下的话成立），所以缺口样本放在中段页。
    _filler = {"id": "z2", "title": "第二页结论已经写好", "content": "证据 5%"}
    def _unresolved_msg(slide):
        deck = intel.think({"slides": [{"id": "z1", "title": "开场结论句", "content": "证据 3%"},
                                       slide, _filler]})
        return next((w["msg"] for w in deck["warnings"]
                     if w.get("rule") == "unresolved_content"), "")
    _msg_topic = _unresolved_msg({"id": "t", "title": "本轮预算的用途与回收路径",
                                  "content": "新增预算 1.8 亿元：选址模型建设 3200 万元。"})
    _msg_empty = _unresolved_msg({"id": "e", "title": "预算用途", "content": ""})
    check("claim: absent 的告警归因分两种根因（内容缺失 / 有内容不成主张）",
          "没有可复述的结论" in _msg_topic and "content 缺失" not in _msg_topic
          and "content 缺失" in _msg_empty and "没有可复述的结论" not in _msg_empty,
          f"topic={_msg_topic[:40]} empty={_msg_empty[:40]}")

    # 焦点数字＝主张所指的那个量（不是第一个、也不是最大的那个）
    unit_econ = ("印尼雅加达自营仓：单均履约成本从 2024Q1 的 19.4 元降至 2025Q4 的 "
                 "11.2 元，降幅 42.3%；同期客单价由 68 元升至 83 元。泰国曼谷加盟："
                 "单均履约成本 16.8 元。")
    claim_ue = "三个市场里只有印尼跑通了单位经济，而且是靠自营履约跑通的"
    check("number_roles: 焦点取主张点名的那个指标（履约成本 11.2 元，不是客单价 83 元）",
          intel.number_roles(unit_econ)[1]["raw"] == "83 元"
          and intel.number_roles(unit_econ, claim_ue)[1]["raw"] == "11.2 元",
          str(intel.number_roles(unit_econ, claim_ue)[1]))
    risk = ("清关环节：口岸清关时效在 1.8 至 6.4 天之间波动，波动幅度 256%，"
            "库存周转天数在 21 至 47 天之间摆动。用工环节：一线分拣人力成本同比上升 18.6%。")
    check("number_roles: 区间下界不是可以立住的量（1.8 出局，取 6.4 天）",
          intel.number_roles(risk)[1]["raw"] == "1.8"
          and intel.number_roles(risk, "清关与用工是无法靠补贴解决的结构性约束")[1]["raw"]
          == "6.4 天")
    years = ("2026 年：收入 4.6 亿元，毛利率 -8.4%。2027 年：收入 12.8 亿元，"
             "毛利率 6.2%。2028 年：收入 27.4 亿元，毛利率 17.6%。")
    check("number_roles: 主张用序数点名某年时，焦点落在该年（第二年 → 2027 年的 12.8 亿元）",
          intel.number_roles(years)[1]["raw"] == "4.6 亿元"
          and intel.number_roles(years, "三年测算里最关键的是第二年")[1]["raw"] == "12.8 亿元")
    # 单位绑定：「个」是分类词，必须与计量名词同进同出（审计实测：主语量显示成「2.1 个」）
    check("numerals: 个百分点/个点/个季度 不被「个」截断（数字与单位不拆开）",
          [n["raw"] for n in intel.numerals("提升 2.1 个点、5 个百分点、缩短 2 个季度")]
          == ["2.1 个点", "5 个百分点", "2 个季度"]
          and intel.numerals("5 个百分点")[0]["unit"] == "个百分点"
          and intel.numerals("2.1 个点")[0]["unit"] == "个点"
          and intel.numerals("缩短 2 个季度")[0]["unit"] == "个季度"
          and intel.numerals("5 个方案")[0]["unit"] == "个",
          str(intel.numerals("提升 2.1 个点、5 个百分点、缩短 2 个季度")))
    # 目标/计划量是意图，不是「当前状态」（审计实测：下降序列的焦点被 2025 目标量顶掉）
    plan_text = ("2022 年 31%，2023 年 27%，2024 年 22%。"
                 "计划：2025 年试点季度回访，目标把复捐率拉回 28%。")
    check("number_roles: 目标/计划量不当选「序列当前值」（22% 是现状，28% 是目标）",
          intel.number_roles(plan_text, "捐赠人复捐率三年连续下降")[1]["raw"] == "22%",
          str(intel.number_roles(plan_text, "捐赠人复捐率三年连续下降")[1]))
    check("number_roles: 整段都是目标句时退回目标量；复合名词「目标客户」不算意图",
          intel.number_roles("目标把复捐率拉回 28%", "目标把复捐率拉回 28%")[1]["raw"] == "28%"
          and intel.number_roles("新增 30 万，目标客户 500 万。",
                                 "目标客户 500 万")[1]["raw"] == "500 万",
          str(intel.number_roles("新增 30 万，目标客户 500 万。", "目标客户 500 万")[1]))
    # 多量纲 + 主张无量 → 数字退为支撑，第一落点回到那句话
    mixed = intel.think(dict(BRIEF, slides=[
        {"id": "m1", "title": "真正的阻力不是竞争，是清关与用工这两件我们无法靠补贴解决的事",
         "content": risk}]))
    check("focus: 主张无量且证据横跨多量纲时，数字退为支撑（不拿波动下界当主语）",
          mixed["pages"][0]["judgment"]["focus"]["element_role"] == "statement_text")
    concentrated = intel.think(dict(BRIEF, slides=[
        {"id": "m2", "title": "三个市场里只有印尼跑通了单位经济",
         "content": unit_econ}]))
    check("focus: 页面由同一指标占住时数字仍是主语（集中度 ≥50% 不降级）",
          concentrated["pages"][0]["judgment"]["focus"]["element_role"] == "kpi_main")

    # 排版尺度：把纪律落成可核对的数字（行长/行距/内边距/断行）
    scale_page = intel.think(dict(BRIEF, slides=[
        {"id": "sc", "title": "真正的阻力不是竞争，是清关与用工这两件我们无法靠补贴解决的事",
         "content": risk}]))["pages"][0]["judgment"]["typography"]["scale"]
    check("typography: 保留内容相关尺度/断行，不注入固定字阶/行数/间距配额",
          scale_page["body_measure_cjk"] <= 45
          and not ({"hierarchy", "line_height", "padding", "lead_lines"} & set(scale_page))
          and scale_page["lead_break"]
          and "｜" in scale_page["lead_break"]
          and not scale_page["lead_break"].split("｜")[1].startswith(("，", "。"))
          and not scale_page["lead_break"].split("｜")[0][-1:].isdigit(),
          str(scale_page.get("lead_break")))

    raw_media = intel.think(dict(BRIEF, slides=[
        {"id": "photo", "title": "真实门店的回访记录",
         "content": "门店现场观察有来源（来源：实地记录）", "asset": "required"}]))
    check("skeleton: 直接调用 think 不谎称资产清单已经落盘",
          "尚未落盘" in intel.build_skeleton(raw_media))
    check("judgment: 受众/决定/阻力/排版/跨页节奏均由逐页判断串起",
          all(all(p["judgment"].get(key) for key in
                  ("audience", "decision", "tension", "typography", "rhythm"))
              for p in bundle["pages"])
          and all(p["judgment"]["typography"].get("lead")
                  and p["judgment"]["typography"].get("discipline")
                  and p["judgment"]["rhythm"].get("move") for p in bundle["pages"])
          and intel._FONTS["sans"]["cn"] == "Source Han Sans SC")

    # 反事实：去掉重复次数惩罚，不损失原本适合内容的构图。
    streak = intel.think(dict(BRIEF, slides=[
        {"id": f"d{i:02d}", "title": title, "content": content}
        for i, (title, content) in enumerate([
            ("请批准试点预算", "请于本周五前确认"),
            ("请批准复用预算", "请于本周五前确认"),
            ("请确认发布节奏", "请于本周五前确认"),
            ("请确认评审名单", "请于本周五前确认")], 1)]))
    streak_comps = [p["judgment"]["composition"]["chosen"] for p in streak["pages"]]
    streak_coh = streak["deck"]["coherence"]
    check("coherence: 整套回看存在（节奏序列与页数对齐，climax 有落点）",
          set(streak_coh) == {"compositions", "densities", "roles", "media",
                              "climax", "adjustments", "notes"}
          and len(streak_coh["compositions"]) == 4
          and len(streak_coh["roles"]) == 4 and len(streak_coh["media"]) == 4
          and streak_coh["climax"]["page"] == "d04")
    check("coherence: 连续同构图不再改选次优项",
          len(set(streak_comps)) == 1 and not streak_coh["adjustments"], str(streak_comps))
    check("coherence: 节奏仍记录停留，不捏造内容变化",
          all(p["judgment"]["rhythm"]["move"] == "hold" for p in streak["pages"][1:3]))
    check("coherence: 焦点连续同落点给出提醒（不硬改焦点）",
          bool(streak_coh["notes"]) and "headline" in streak_coh["notes"][0],
          str(streak_coh["notes"]))
    check("world: 单页单个词不决定整套世界（deck 级证据门槛）",
          streak["deck"]["world"]["source"] == "neutral_fallback",
          str(streak["deck"]["world"]["name"]))
    # 内容契合压过节奏：差距够大时不硬换
    stubborn = intel.think(dict(BRIEF, slides=[
        {"id": f"g{i:02d}", "title": f"三年曲线{i}",
         "content": "96 元 → 128 元 → 205 元（来源：投放台账）"} for i in range(1, 4)]))
    check("coherence: 内容强烈要求同一构图时不硬换（判断不是规则）",
          all(p["judgment"]["composition"]["chosen"] == "data_field"
              for p in stubborn["pages"])
          and not stubborn["deck"]["coherence"]["adjustments"])

    # 数字语义：日期/编号的碎片不是证据量（权重卡曾出现「次 -09」）
    dateless = intel.understand("18 个月分三段落地（来源：项目排期 2026-09）")
    check("understanding: 日期碎片（2026-09）不进证据量",
          [(n["raw"], n["unit"]) for n in dateless["numerals"]] == [("18 个月", "个月")]
          and dateless["distinct_numerals"] == 1, str(dateless["numerals"]))
    check("understanding: 多处日期形态（2026-09-30 / 2026/03 / 2026 年 9 月）都不算量",
          intel.evidence_numerals("2026-09-30 上线") == []
          and intel.evidence_numerals("回访 2026/03") == []
          and intel.evidence_numerals("2026 年 9 月上线") == []
          and [n["raw"] for n in intel.evidence_numerals("周期 18 个月")] == ["18 个月"])
    check("understanding: 出处标注里的数字不进证据（台账编号是档案坐标）",
          intel.understand("NPS 从 31 提升到 47（来源：客户回访 2026/03）")["distinct_numerals"] == 2)

    # 结论抽取：标题是结论位；数字列举不是结论；同分不由浮点噪声决定
    titled = intel.extract_claim({"title": "华东贡献 42% 增量，但复购率最低",
                                  "content": "华东 +42%、华南 +28%、华北 +15%（来源：BI 2026Q3）"})
    check("claim: 标题（作者的结论位）不被内容里的数字列举顶掉",
          titled["text"] == "华东贡献 42% 增量，但复购率最低", str(titled))
    listing_only = intel.extract_claim({"title": "华东区数据",
                                        "content": "华东 +42%、华南 +28%、华北 +15%"})
    check("claim: 只有数字列举的句子不当结论（没有主张就是 absent）",
          listing_only["text"] is None and listing_only["source"] == "absent",
          str(listing_only))
    topic_title = intel.extract_claim({"title": "成本结构",
                                       "content": "单篇生产成本降至 1,800 元，低于去年同期"})
    check("claim: 标题仍是话题标签时照旧从内容抽结论",
          topic_title["text"] == "单篇生产成本降至 1,800 元，低于去年同期", str(topic_title))

    # 运行数据最小化：判断卡只留被消费的键
    j0 = bundle["pages"][0]["judgment"]
    check("slim: 判断卡无解释性冗余字段",
          set(j0["media"]) == {"decision", "source", "function", "necessity"}
          and set(j0["claim"]) == {"text", "source", "why"}
          and set(j0["spatial"]) == {"reading_path", "primary_zone", "quiet_zone",
                                     "whitespace_duty"}
          and set(j0["production"]) == {"must_place", "must_not", "density_target"}
          and set(j0["composition"]) == {"chosen", "why", "rejected"})
    check("slim: 构图否决项至多 2 条（无竞争为空，只留真实竞争）",
          all(len(p["judgment"]["composition"]["rejected"]) <= 2
              and all(r["why_not"] and r["option"] != p["judgment"]["composition"]["chosen"]
                      for r in p["judgment"]["composition"]["rejected"])
              for p in pages))

    # 焦点：年份与编号不当第一落点
    year_focus = intel.focus_of(
        "summarize",
        {"evidence": "none", "numerals": [{"raw": "2026", "value": 2026.0,
                                           "unit": "", "year": True}],
         "distinct_numerals": 0, "comparison": False},
        {"text": "2026 年开局", "source": "extracted"}, {})
    check("focus: 年份不当第一落点（回到标题）",
          year_focus["element_role"] == "headline", str(year_focus["element_role"]))
    check("numerals: 字母紧贴的数字是编号不是证据（Q1 不计数）",
          intel.understand("2027Q1 启动")["distinct_numerals"] == 0
          and intel.understand("Q2 末交付 600 万")["distinct_numerals"] == 1)
    # 列表前缀「1）/2）/1. /(1)」是顺序枚举标记，不是可比较的数字——
    # 否则三大风险/三条规矩/四步法的第一行会错把「1」当成 KPI。
    check("numerals: 句首列表枚举序号（1）/1）/1. 不是证据量",
          # 1）2）3）被剥离 → 剩 60%、2家；③无后续数字。
          intel.understand("1）汇率波动覆盖60%；2）本地竞争并购2家；3）政策变化。")[
              "distinct_numerals"] == 2
          and intel.understand("(1) 试点 80% (2) 复制 100%")["distinct_numerals"] == 2
          and intel.understand("①A ②B ③C")["distinct_numerals"] == 0
          and intel.understand("1. 试点 2. 复制 3. 全量")["distinct_numerals"] == 0)
    check("claim: 列表项结论会剥掉序号前缀（不把「1）」当开头）",
          intel.extract_claim({"content": "1）汇率波动：对冲策略覆盖首年60%支出；2）…"})["text"]
          == "汇率波动：对冲策略覆盖首年60%支出")
    check("understanding: 箭头链是顺序的形状（盖过孤立数字）",
          intel.understand("18 个月：试点 → 复制 → 全量")["evidence"] == "sequence")

    # 结论：超长论述不当结论（可复述才算记住）
    long_claim = intel.extract_claim(
        {"title": "", "content": "结论是" + "甲" * 64 + "。结论是短句"})
    check("claim: 同等条件下短结论赢过超长论述（长度是记忆成本）",
          long_claim["text"] == "结论是短句", str(long_claim))

    # 信息权重真正抵达生产：删除清单变成明确禁令
    prod_w = intel.production_spec(
        {"element_role": "headline", "type": "text", "why": "w", "rejected": []},
        {"chosen": "big_whitespace", "rejected": []},
        {"decision": "none"}, {"numerals": []}, dup)
    check("weight: 删除清单进生产禁令（must_not 可执行）",
          any(str(m).startswith("删：") for m in prod_w["must_not"]))

    # 作者声明：写了画面主题就是要出图；未知构图不断规划
    subj = intel.think(dict(BRIEF, slides=[
        {"id": "x", "title": "安静的数字", "content": "口径稳定（来源：台账）",
         "asset_subject": "a calm archive room, empty left"}]))
    check("media: asset_subject 等价 required（声明不只留在回显里）",
          subj["pages"][0]["judgment"]["media"]["decision"] == "required"
          and subj["pages"][0]["judgment"]["media"]["source"] == "author"
          and "x" in subj["assets_hint"]["generate"])
    custom = intel.think(dict(BRIEF, slides=[
        {"id": "x", "title": "t", "content": "c", "composition": "custom_xyz"}]))
    check("plan: 作者未知构图不断规划（密度取中性）",
          custom["pages"][0]["judgment"]["composition"]["chosen"] == "custom_xyz"
          and custom["pages"][0]["judgment"]["production"]["density_target"] == "balanced")

    # 世界：整套证据性格取最常见的非空证据（「无证据」不盖过真证据）
    mixed = intel.think(dict(BRIEF, slides=[
        {"id": "c1", "title": "A", "content": "甲对比乙：61% 对 38%（来源：台账）"},
        {"id": "c2", "title": "B", "content": "丙对比丁：44% 对 29%（来源：台账）"},
        {"id": "n1", "title": "随笔", "content": "一些想法"}]))
    check("world: 图表性格跟随主导证据（对比）",
          mixed["deck"]["world"]["chart_style"]
          == "shared baseline, two-tone delta, direct labels",
          str(mixed["deck"]["world"]["chart_style"]))


# ── 2 · verify（硬门）────────────────────────────────────────────────
def test_verify(tmp: Path):
    import verify
    from verify import check_spec, normalize_spec
    from primitives import spec_fingerprint

    build = make_build(tmp)
    namespace: dict = {}
    exec(compile(build.read_text(encoding="utf-8"), str(build), "exec"), namespace)
    spec = namespace["SPEC"]

    ok = check_spec(spec)
    check("guard: 合规稿 0 error", ok["passed"], str(ok["checks"])[:200])
    check("guard: 合规稿零 warning（QA 不产生噪音）",
          not [c for c in ok["checks"] if c.get("level") == "warn"], str(ok["warnings"])[:160])

    def mutate(fn):
        import copy
        bad = copy.deepcopy(spec)
        fn(bad)
        return check_spec(bad)

    r = mutate(lambda s: s["slides"][0]["elements"][0].update({"height": 8}))
    check("guard: 文字溢出必拦（TEXT_OVERFLOW）",
          any(c["rule"] == "text_capacity" for c in r["checks"]))
    spaced_paras = mutate(lambda s: s["slides"][0]["elements"][0].update(
        {"text": "第一段正文\n第二段正文", "size": 20, "width": 400,
         "height": 59, "space_after": 22, "line_height": 1.35}))
    tracked = mutate(lambda s: s["slides"][0]["elements"][0].update(
        {"text": "MMMMMMMM", "size": 28, "width": 245, "height": 52,
         "max_lines": 1, "char_spacing": 16}))
    check("guard: 段距/字距真实占位，不能在编译前把会断行的标题错判 PASS",
          any(c["rule"] == "text_capacity" for c in spaced_paras["checks"])
          and any(c["rule"] == "text_capacity" for c in tracked["checks"]))
    invalid_margin = mutate(lambda s: s["slides"][0]["elements"][0].update(
        {"padding": -44}))
    invalid_leading = mutate(lambda s: s["slides"][0]["elements"][0].update(
        {"line_height": -1}))
    check("guard: 负边距/负行高在 guard 阶段 fail-closed，不靠编译异常止损",
          any(c["rule"] == "text_capacity" for c in invalid_margin["checks"])
          and any(c["rule"] == "text_capacity" for c in invalid_leading["checks"]))

    r = mutate(lambda s: (s["slides"][0]["elements"].append(
        {"id": "dup", "type": "text", "x": 48, "y": 96, "width": 400, "height": 56,
         "text": "重叠", "size": 32, "color": "ink"})))
    check("guard: 墨迹重叠必拦（OVERLAP）",
          any(c["rule"] == "overlap" for c in r["checks"]))

    r = mutate(lambda s: s["slides"][0]["elements"][0].update({"x": 1300}))
    check("guard: 越出画布必拦", any(c["rule"] == "safety" for c in r["checks"]))

    r = mutate(lambda s: s["slides"][0]["elements"][0].update({"width": 0}))
    check("guard: 退化几何必拦（元素不可见）",
          any(c["rule"] == "element_schema" for c in r["checks"]))

    r = mutate(lambda s: s["slides"][0]["elements"][0].update({"color": "secondry"}))
    check("guard: 未知色名必拦（静默回落不留痕）",
          any(c["rule"] == "color_token" for c in r["checks"]))

    r = mutate(lambda s: s["slides"][2]["elements"].append(
        {"id": "ghost_chart", "type": "chart", "chart_kind": "unknown_kind",
         "x": 48, "y": 300, "width": 400, "height": 200}))
    check("guard: 未知图表类型必拦",
          any(c["rule"] == "chart_type" for c in r["checks"]))

    r = mutate(lambda s: s["slides"][2]["elements"].append(
        {"id": "empty_chart", "type": "chart", "chart_kind": "waterfall",
         "x": 600, "y": 300, "width": 400, "height": 200, "data": []}))
    check("guard: 图表空载荷必拦", any(c["rule"] == "chart_payload" for c in r["checks"]))

    r = mutate(lambda s: s["slides"][0]["elements"].append(
        {"id": "intruder", "type": "text", "x": 60, "y": 676, "width": 300, "height": 24,
         "text": "侵入来源区", "size": 12.5, "color": "muted"}))
    check("guard: 主体侵入来源区必拦（SOURCE_COLLISION）",
          any(c["rule"] == "source_zone" for c in r["checks"]))

    r = mutate(lambda s: s["slides"].__setitem__(0, {"id": "s01", "elements": []}))
    check("guard: 空白页必拦", any(c["rule"] == "page_contract" for c in r["checks"]))
    r = mutate(lambda s: s.__setitem__("slides", []))
    check("guard: 空 deck 必拦", not r["passed"])

    # 数据完整性：label 缺失 / 非有限值 / 跨页单位打架
    def _chart(s):
        return next(e for e in s["slides"][2]["elements"] if e.get("type") == "chart")
    r = mutate(lambda s: _chart(s)["data"].append({"value": 9}))
    check("guard: 数据行缺 label 必拦",
          any(c["rule"] == "data_integrity" for c in r["checks"]))
    def _add_second_chart(s):
        import copy as _copy
        second = _copy.deepcopy(_chart(s))
        second.update({"id": "chart_page5", "x": 48, "y": 232, "unit": "万元"})
        s["slides"][4]["elements"] = [e for e in s["slides"][4]["elements"]
                                      if e.get("id") != "src"] + [second]
    r = mutate(_add_second_chart)
    check("guard: 同指标跨页单位不一致必拦",
          any(c["rule"] == "metric_consistency" for c in r["checks"]), str(r["checks"])[:200])
    r = check_spec(spec, rules={"require_provenance": True})
    check("guard: release 档出处齐全 → 不拦", r["passed"])
    missing_basis = json.loads(json.dumps(spec))
    next(e for e in missing_basis["slides"][2]["elements"]
         if e.get("type") == "chart").pop("basis")
    r = check_spec(missing_basis, rules={"require_provenance": True})
    check("guard: release 档缺比较口径必拦（DATA_INTEGRITY_FAIL）",
          any(c["rule"] == "data_provenance" for c in r["checks"]))

    # 归一化：8 单位吸附 + 色名规范化 + 幂等
    import copy
    messy = copy.deepcopy(spec)
    messy["slides"][0]["elements"][0]["x"] = 50
    messy["slides"][0]["elements"][0]["color"] = "#191510"
    normalized, report = normalize_spec(messy)
    check("normalize: 几何落网格、色值规范化",
          normalized["slides"][0]["elements"][0]["x"] % 4 == 0 and report["counts"])
    twice, _ = normalize_spec(normalized)
    check("normalize: 幂等（二次归一指纹不变）",
          spec_fingerprint(twice) == spec_fingerprint(normalized))

    # 回归（v6.5）：亚网格厚度不被抬高，同心关系不被吸附拆散。
    # 症状曾是——1px 发丝线被抬成 4px、骑线对象与线各挪各的，
    # 产物里对象永远悬在线的一侧，而硬门全绿（几何合法、不重叠、不溢出）。
    concentric = {"theme": {"colors": {}}, "slides": [{"id": "s", "elements": [
        {"id": "hairline", "type": "shape", "shape": "rect",
         "x": 96, "y": 280, "width": 400, "height": 1},
        {"id": "rider", "type": "shape", "shape": "rect",
         "x": 96, "y": 276, "width": 200, "height": 8},
    ]}]}
    norm_c, _ = normalize_spec(concentric)
    line_e, bar_e = norm_c["slides"][0]["elements"]
    check("normalize: 亚网格厚度原样保留（1px 发丝线不被抬成 4px）",
          line_e["height"] == 1, f'height={line_e["height"]}')
    check("normalize: 同心关系吸附后仍同心（线与骑线对象共中线）",
          abs((line_e["y"] + line_e["height"] / 2)
              - (bar_e["y"] + bar_e["height"] / 2)) < 1e-9,
          f'线中心={line_e["y"] + line_e["height"] / 2}，'
          f'对象中心={bar_e["y"] + bar_e["height"] / 2}')
    twice_c, _ = normalize_spec(norm_c)
    check("normalize: 同心吸附幂等（二次归一不再漂移）",
          spec_fingerprint(twice_c) == spec_fingerprint(norm_c))

    # 幂等对任意输入成立（不只是网格上的夹具）：位置不在网格、中心在网格的几何是
    # 上一趟「按中心吸附」的结果，再按边吸附会挪 ≤grid/2。随机 400 例实测旧写法
    # 307 例漂移——单次归一不动看不出来，但只要有第二个消费者重新归一（编译 / 预览 /
    # 复验共用同一份几何），对象就与 guard 认证的位置差 2px。
    import random as _rand_norm
    _rng_norm = _rand_norm.Random(20261020)
    _drift = None
    for _ in range(400):
        _e = {"type": "text", "id": "t", "x": _rng_norm.uniform(0, 600),
              "y": _rng_norm.uniform(0, 400), "width": _rng_norm.uniform(4, 600),
              "height": _rng_norm.uniform(4, 300), "size": 20, "text": "x", "max_lines": 2}
        _s = {"canvas": {"width": 1280, "height": 720}, "theme": {"colors": {}},
              "slides": [{"id": "s01", "elements": [_e]}]}
        _a, _ = normalize_spec(copy.deepcopy(_s))
        _b, _ = normalize_spec(copy.deepcopy(_a))
        _ea = _a["slides"][0]["elements"][0]
        _eb = _b["slides"][0]["elements"][0]
        if any(_ea[k] != _eb[k] for k in ("x", "y", "width", "height")):
            _drift = (_e, _ea, _eb)
            break
    check("normalize: 任意几何下幂等（二次归一不再位移）",
          _drift is None, str(_drift)[:180])

    # 未消费字段：作者写错的字段名被渲染器静默丢弃（`font_size` 写了但标题还是默认
    # 字号），而契约把这一类列为必拦。判据两档：近似已有字段 → error + 候选名；
    # 完全对不上的未知字段 → warn（可能是作者自己的注解，不假拦）。
    _u_base = {"canvas": {"width": 1280, "height": 720}, "theme": {"colors": {}}}
    _u_text = {"type": "text", "id": "t", "x": 100, "y": 100, "width": 600, "height": 120,
               "size": 32, "text": "结论一句话", "max_lines": 2, "line_height": 1.3}

    def _field_report(**extra):
        spec = {"canvas": {"width": 1280, "height": 720}, "theme": {"colors": {}},
                "slides": [{"id": "s01", "elements": [dict(_u_text, **extra)]}]}
        rep = check_spec(spec)
        return ([c for c in rep["checks"] if c["rule"] == "element_field_unknown"
                 and c["level"] == "error"],
                [c for c in rep["checks"] if c["rule"] == "element_field_unknown"
                 and c["level"] == "warn"])

    _typo_errs, _ = _field_report(font_size=40)
    check("guard: 疑似写错的字段名拦下并给候选（font_size → size）",
          bool(_typo_errs) and "size" in _typo_errs[0]["msg"], str(_typo_errs[:1])[:160])
    _typo_errs2, _ = _field_report(maxlines=3)
    check("guard: 下划线写错同样拦下（maxlines → max_lines）",
          bool(_typo_errs2) and "max_lines" in _typo_errs2[0]["msg"],
          str(_typo_errs2[:1])[:160])
    _far_errs, _far_warns = _field_report(note_for_me="作者自己的注解")
    check("guard: 对不上的未知字段只告警、不假拦（作者注解可留）",
          not _far_errs and bool(_far_warns), str(_far_warns[:1])[:160])
    _clean_errs, _clean_warns = _field_report()
    check("guard: 合法字段不产生未消费告警（零噪音）",
          not _clean_errs and not _clean_warns, str(_clean_warns[:1])[:160])

    # 文档↔代码：contract.md 里点名的字段，渲染器必须真的消费（否则作者照文档写却被拦，
    # 或者写了被静默丢弃）。这条钉把「规范说了、代码没做到」变成可执行的回归。
    import re as _re_doc
    _doc = (ROOT / "references" / "contract.md").read_text(encoding="utf-8")
    _known = verify.consumed_fields()
    _not_element = {"spec", "slides", "plan_path", "fix_plan", "draft", "release",
                    "matrix", "pie", "ranked_bar", "arrow", "image", "shape", "text",
                    "type", "id", "role", "name", "value", "display", "points", "layers",
                    "chart_field_unconsumed", "body", "latin", "family"}
    _doc_fields = set(_re_doc.findall(r"`([a-z_][a-z0-9_]*)`", _doc))
    _doc_missing = sorted(f for f in _doc_fields - _not_element if f not in _known)
    check("docs↔code: contract.md 点名的字段都在渲染器消费集里",
          not _doc_missing, str(_doc_missing))

    # 契约的填充四写法要求渐变 stops≥2。编译器会抛 COMPILE_FAIL，但 guard 应给出
    # 可修诊断（同一判定，更早、更省一次编译）。
    _grad = check_spec({"canvas": {"width": 1280, "height": 720}, "theme": {"colors": {}},
                        "slides": [{"id": "s", "elements": [
                            {"type": "shape", "id": "r", "shape": "rect", "x": 100, "y": 100,
                             "width": 400, "height": 200,
                             "fill": {"type": "gradient", "angle": 45,
                                      "stops": [{"color": "#112233"}]}}]}]})
    check("guard: 渐变 stops<2 在编译前拦下（契约四写法）",
          any(c["rule"] == "element_schema" and c["level"] == "error" for c in _grad["checks"]),
          str(_grad["checks"][:1])[:140])

    # 回归（v9.6.1）：LINE / ARROW 的零厚度轴必须保留为真正的横线/竖线。
    # 否则零宽竖线会被扩成网格宽度，PPT connector 从左上连到右下，出现轻微斜线。
    vertical = {"canvas": {"width": 1280, "height": 720},
                "theme": {"colors": {}}, "slides": [{"id": "axis", "elements": [
        {"id": "vertical_axis", "type": "shape", "shape": "line",
         "x": 93, "y": 278, "width": 0, "height": 239,
         "stroke": "#C8C2B5", "stroke_width": 1},
        {"id": "node", "type": "shape", "shape": "ellipse",
         "x": 88.5, "y": 273.5, "width": 9, "height": 9,
         "fill": "#713B43", "stroke": "none"},
    ]}]}
    norm_v, _ = normalize_spec(vertical)
    axis_e, node_e = norm_v["slides"][0]["elements"]
    check("normalize: 零宽竖线保留 width=0，长度按中心吸附",
          axis_e["width"] == 0 and axis_e["height"] == 240,
          f'width={axis_e["width"]}, height={axis_e["height"]}')
    check("normalize: 圆点与零宽竖线保持同一中线",
          abs(axis_e["x"] - (node_e["x"] + node_e["width"] / 2)) < 1e-9,
          f'线 x={axis_e["x"]}, 圆点中心={node_e["x"] + node_e["width"] / 2}')
    accepted_vertical = check_spec(norm_v)
    check("guard: 单轴为零的竖线合法",
          accepted_vertical["passed"])
    from compiler import compile_deck
    from pptx import Presentation
    vertical_pptx = tmp / "zero-width-vertical-line.pptx"
    compile_deck(norm_v, vertical_pptx)
    compiled_shapes = {sh.name: sh for sh in Presentation(str(vertical_pptx)).slides[0].shapes}
    compiled_axis, compiled_node = compiled_shapes["vertical_axis"], compiled_shapes["node"]
    check("compiler: 零宽竖线在 PPTX 中保持竖直且与圆点同心",
          compiled_axis.width == 0
          and compiled_axis.left == compiled_node.left + compiled_node.width // 2,
          f'line width={compiled_axis.width}, node center={compiled_node.left + compiled_node.width // 2}')
    zero_length = {"canvas": {"width": 1280, "height": 720},
                   "theme": {"colors": {}}, "slides": [{"id": "zero_line", "elements": [
        {"id": "zero_length_line", "type": "shape", "shape": "line",
         "x": 40, "y": 40, "width": 0, "height": 0,
         "stroke": "#C8C2B5", "stroke_width": 1},
    ]}]}
    norm_zero_line, _ = normalize_spec(zero_length)
    zero_line_report = check_spec(norm_zero_line)
    check("guard: 零长度线仍由几何门拦截",
          any(c["rule"] == "element_schema" and "线的 width 与 height 同时为 0" in c["msg"]
              for c in zero_line_report["checks"]))
    twice_v, _ = normalize_spec(norm_v)
    check("normalize: 零宽竖线归一化幂等",
          spec_fingerprint(twice_v) == spec_fingerprint(norm_v))

    horizontal = {"canvas": {"width": 1280, "height": 720},
                  "theme": {"colors": {}}, "slides": [{"id": "rule", "elements": [
        {"id": "horizontal_axis", "type": "shape", "shape": "line",
         "x": 93, "y": 278, "width": 239, "height": 0,
         "stroke": "#C8C2B5", "stroke_width": 1},
        {"id": "node", "type": "shape", "shape": "ellipse",
         "x": 208, "y": 273.5, "width": 9, "height": 9,
         "fill": "#713B43", "stroke": "none"},
    ]}]}
    norm_h, _ = normalize_spec(horizontal)
    axis_h, node_h = norm_h["slides"][0]["elements"]
    check("normalize: 零高横线保留 height=0，长度按中心吸附",
          axis_h["height"] == 0 and axis_h["width"] == 240,
          f'width={axis_h["width"]}, height={axis_h["height"]}')
    check("normalize: 圆点与零高横线保持同一中线",
          abs(axis_h["y"] - (node_h["y"] + node_h["height"] / 2)) < 1e-9,
          f'线 y={axis_h["y"]}, 圆点中心={node_h["y"] + node_h["height"] / 2}')
    check("guard: 单轴为零的横线合法",
          check_spec(norm_h)["passed"])
    twice_h, _ = normalize_spec(norm_h)
    check("normalize: 零高横线归一化幂等",
          spec_fingerprint(twice_h) == spec_fingerprint(norm_h))

    # 负尺寸/零面积不由归一化器偷偷修成可见对象，仍交给 guard 拦截。
    collapsed = {"canvas": {"width": 1280, "height": 720},
                 "theme": {"colors": {}}, "slides": [{"id": "degenerate", "elements": [
        {"id": "collapsed_rect", "type": "shape", "shape": "rect",
         "x": 20, "y": 20, "width": 0, "height": 20, "fill": "#713B43"},
    ]}]}
    norm_zero, _ = normalize_spec(collapsed)
    zero_report = check_spec(norm_zero)
    check("normalize: 零宽非线形保持退化并交由 guard 拦截",
          norm_zero["slides"][0]["elements"][0]["width"] == 0
          and any(c["rule"] == "element_schema" for c in zero_report["checks"]))

    # verdict：PASS / BLOCK 二态 + 分组修复包
    from verify import verdict
    good = verdict(spec, mode="release",
                   guard_report=ok, compile_report={"passed": True, "skipped": False,
                                                    "output_sha256": "a" * 64})
    check("verdict: 合规稿 PASS 且 release_eligible", good["passed"]
          and good["release_eligible"] and good["status"] == "PASS")
    bad_report = mutate(lambda s: s["slides"][0]["elements"][0].update({"height": 8}))
    bad = verdict(spec, mode="draft", guard_report=bad_report,
                  compile_report={"passed": False, "skipped": True})
    check("verdict: BLOCK 按根因分组（修复包自足）",
          not bad["passed"] and bad["failure_codes"] == ["TEXT_OVERFLOW"]
          and bad["fix_plan"]["groups"][0]["fix"])


# ── 3 · production（编译 / 缓存 / 资产 / 预览）────────────────────────
def test_production(tmp: Path):
    from compiler import compile_deck
    from primitives import ENGINE_SCOPES, file_digest, spec_view, compile_reuse, record_compile

    build = make_build(tmp)
    namespace: dict = {}
    exec(compile(build.read_text(encoding="utf-8"), str(build), "exec"), namespace)
    spec = namespace["SPEC"]

    out1, out2 = tmp / "a.pptx", tmp / "b.pptx"
    compile_deck(spec, out1)
    compile_deck(spec, out2)
    check("compiler: 同 spec 两次编译字节一致（确定性）",
          file_digest(out1) == file_digest(out2))
    # 系列色必须由引擎补齐，不能是 None：主题只声明 accent 时 `derive_tokens` 不派生
    # 色阶，此前 series_palette 直接吐 None → 编译器把 None 写进 line.color.rgb，
    # python-pptx 抛断言，整条序列的线宽/端点标记全落不下去（release 档 COMPILE_FAIL）。
    from primitives import RenderContext
    _minimal_theme = {"colors": {"background": "#FFFFFF", "ink": "#111111",
                                "muted": "#777777", "accent": "#AA0000"}}
    _pal = RenderContext(_minimal_theme, {"width": 1280, "height": 720}).series_palette(4, -1)
    check("primitives: 主题未派生色阶时系列色由引擎推导（非 None、互不相同）",
          all(c is not None for c in _pal) and len({str(c) for c in _pal}) == 4, str(_pal))
    minimal_deck = {"canvas": {"width": 1280, "height": 720}, "theme": _minimal_theme,
                    "slides": [{"id": "s01", "elements": [
                        {"type": "chart", "id": "c1", "chart_kind": "line", "x": 80, "y": 120,
                         "width": 640, "height": 320,
                         "categories": ["2023", "2024", "2025"],
                         "series": [{"name": "毛利率", "values": [58, 63, 67]}]}]}]}
    minimal_report = compile_deck(minimal_deck, tmp / "min_theme_chart.pptx", speed="fast")
    check("compiler: 该主题下多序列图表零告警编译（缺主色不等于页面编译失败）",
          minimal_report["passed"] and not minimal_report["warnings"],
          str(minimal_report["warnings"])[:200])

    from pptx import Presentation
    prs = Presentation(str(out1))
    texts = [sh.text_frame.text.replace("\u2009", "").replace("\u200a", "")
             for sl in prs.slides for sh in sl.shapes
             if sh.has_text_frame and sh.text_frame.text.strip()]
    check("compiler: 原生可编辑 PPTX（文本在对象模型里可改）",
          len(prs.slides) == 5 and any("第3页结论句" in t for t in texts))

    # 同一数值内容，两种载荷不能一边有原生标签、一边只有预览标签。
    from pptx.oxml.ns import qn
    base_chart = {"id":"labels", "type":"chart", "chart_kind":"column", "x":48,"y":220,
                  "width":1184,"height":360,"role":"focus", "show_values":True,
                  "label_size":24,"value_size":28,"source":"合成数据","unit":"项",
                  "period":"测试批次","basis":"分类计数","number_format":"0.0"}
    forms = [dict(base_chart, data=[{"label":"研发","value":12},{"label":"交付","value":8}]),
             dict(base_chart, categories=["研发","交付"],series=[{"name":"本期","values":[12,8]}]),
             dict(base_chart, categories=["研发","交付"],series=[{"name":"本期","values":[12,8]},
                                                               {"name":"对照","values":[10,6]}])]
    chart_spec = dict(spec, slides=[{"id":f"labels{i}","elements":[e]} for i,e in enumerate(forms)])
    labeled_out = tmp / "native-labels.pptx"
    compile_deck(chart_spec,labeled_out)
    labeled_prs = Presentation(labeled_out)
    native_charts = [sl.shapes[0].chart for sl in labeled_prs.slides]
    check("charts: data / 单 series / 多 series 均消费 show_values",
          all(c.plots[0].has_data_labels and c.plots[0].data_labels.show_value for c in native_charts))
    check("charts: 作者 px 字号进入原生标签，不再被 8/9/10pt 硬编码覆盖",
          all(c.category_axis.tick_labels.font.size.pt == 18
              and c.value_axis.tick_labels.font.size.pt == 18
              and c.plots[0].data_labels.font.size.pt == 21 for c in native_charts))
    check("charts: 中文 a:ea 与拉丁 a:latin 字体均有明确来源",
          all(c._chartSpace.xpath(".//a:ea")
              and {n.get("typeface") for n in c._chartSpace.xpath(".//a:ea")} == {"Source Han Serif SC"}
              and {n.get("typeface") for n in c._chartSpace.xpath(".//a:latin")} == {"Georgia"}
              for c in native_charts))
    check("charts: 两种载荷数值/格式不漂移且保持原生对象",
          list(native_charts[0].series[0].values) == list(native_charts[1].series[0].values) == [12,8]
          and all(c.plots[0].data_labels.number_format == "0.0" for c in native_charts))
    dense_chart = dict(base_chart, height=120, categories=list("甲乙丙丁戊己"),
                       series=[{"name":"本期","values":[1,2,3,4,5,6]}])
    from verify import check_spec
    dense_spec = dict(spec, slides=[{"id":"dense-labels","elements":[dense_chart]}])
    dense_guard = check_spec(dense_spec)
    check("charts: categories/series 不再绕过全页标签容量硬门",
          any(c.get("name") == "chart_label_collision" or c.get("rule") == "chart_label_collision"
              for c in dense_guard.get("checks",[]) if c.get("level") == "error"), str(dense_guard)[:300])
    hidden = dict(dense_chart,label_collision_policy="hide_redundant")
    disabled = dict(forms[1],show_values=False)
    hidden_spec = dict(spec,slides=[{"id":"hidden","elements":[hidden]},
                                  {"id":"disabled","elements":[disabled]}])
    hidden_out=tmp/"hidden-labels.pptx"; compile_deck(hidden_spec,hidden_out)
    check("charts: 作者禁用或声明隐藏重复值时不强加标签",
          all(not sl.shapes[0].chart.plots[0].has_data_labels for sl in Presentation(hidden_out).slides))
    import ghost as _chart_ghost
    from unittest.mock import patch
    with patch.object(_chart_ghost,"_font",wraps=_chart_ghost._font) as fonts:
        _chart_ghost.ghost_page(chart_spec["slides"][1],chart_spec,scale=1,supersample=1)
    check("ghost: 原生图表类目/数值遵循作者字号与字体角色",
          any(c.args[0] == 24 and c.kwargs.get("cjk") and c.kwargs.get("family") == "_serif"
              for c in fonts.call_args_list)
          and any(c.args[0] == 28 and c.kwargs.get("family") == "_serif" for c in fonts.call_args_list))
    check("ghost: 单 series 不把序列名缝进类别标签",
          [r["label"] for r in _chart_ghost._rows(forms[1])] == ["研发","交付"])

    # ── 9.31.1 登记缺口的回归：categories+series 路径上的 highlight 必须落 dPt ──
    # 缺口的形态是「写了强调、产物里什么都没有」（dPt=0 静默吞掉），而 ghost 预览
    # 渲染了它——预览/产物分叉是最坏的一种不一致。三种载荷形态一起锁：
    # data（既有路径不得回退）／单条 series（点级）／多序列（点级 + 序列级）。
    import zipfile as _zip
    def _dpt_count(pptx_path):
        with _zip.ZipFile(pptx_path) as z:
            charts = [n for n in z.namelist() if n.startswith("ppt/charts/chart")]
            xml = z.read(charts[0]).decode("utf-8")
        return xml.count("<c:dPt>"), xml

    hl_theme = {"colors": {"background": "#FFFFFF", "ink": "#111111", "muted": "#777777",
                           "primary": "#222222", "secondary": "#333333", "accent": "#AA0000"}}
    hl_base = {"id": "hl", "type": "chart", "x": 40, "y": 40, "width": 800, "height": 500,
               "chart_kind": "column", "categories": ["A", "B", "C", "D"], "role": "focus"}
    hl_cases = [
        ("data/点级（既有路径）", dict(hl_base, data=[{"label": c, "value": v} for c, v in
                                                  zip("ABCD", (10, 14, 18, 22))]),
         "C", 1),
        ("单条 series/类别名", dict(hl_base, series=[{"name": "本期", "values": [10, 14, 18, 22]}]),
         "C", 1),
        ("单条 series/整数", dict(hl_base, series=[{"name": "本期", "values": [10, 14, 18, 22]}]),
         2, 1),
        ("多序列/类别名", dict(hl_base, series=[{"name": "本期", "values": [10, 14, 18, 22]},
                                              {"name": "对照", "values": [5, 9, 13, 20]}]),
         "C", 2),
        ("多序列/序列名（不落 dPt，走序列级）",
         dict(hl_base, series=[{"name": "本期", "values": [10, 14, 18, 22]},
                               {"name": "对照", "values": [5, 9, 13, 20]}]), "对照", 0),
        ("写错名字：不猜强调对象",
         dict(hl_base, series=[{"name": "本期", "values": [10, 14, 18, 22]}]), "查无此类", 0),
    ]
    for tag, el, hl, expect in hl_cases:
        deck = {"canvas": {"width": 1280, "height": 720}, "theme": hl_theme,
                "slides": [{"id": "hl", "elements": [dict(el, highlight=hl)]}]}
        out = tmp / f"highlight-{abs(hash(tag))}.pptx"
        report = compile_deck(deck, out)
        count, xml = _dpt_count(out)
        check(f"highlight: {tag} → dPt={expect}",
              report["passed"] and count == expect
              and (expect == 0 or 'val="AA0000"' in xml),
              f"dPt={count} warnings={report['warnings'][:1]}")
    # accent 预留：点级强调时，任何序列基础色都不得占用 accent（强调只占一个位置）
    for tag, el, hl in (("单条", dict(hl_base, series=[{"name": "本期", "values": [10, 14, 18, 22]}]),
                         "C"),
                        ("多序列", dict(hl_base, series=[{"name": "本期", "values": [10, 14, 18, 22]},
                                                        {"name": "对照", "values": [5, 9, 13, 20]}]), "C")):
        deck = {"canvas": {"width": 1280, "height": 720}, "theme": hl_theme,
                "slides": [{"id": "hl", "elements": [dict(el, highlight=hl)]}]}
        out = tmp / f"highlight-accent-{abs(hash(tag))}.pptx"
        compile_deck(deck, out)
        _, xml = _dpt_count(out)
        series_colors = set(re.findall(r'<c:ser>.*?<a:srgbClr val="([0-9A-F]{6})"', xml, re.S))
        check(f"highlight: {tag} 序列基础色不占用 accent（强调只占一个位置）",
              series_colors and "AA0000" not in series_colors, str(series_colors))
    # 预览/产物一致：同一次解析决定两边强调；改强调目标只改预览的落色位置，
    # 写错名字时两边都不产生强调（不替作者猜）。
    def _hl_image(el):
        deck = {"canvas": {"width": 1280, "height": 720}, "theme": hl_theme,
                "slides": [{"id": "hl", "elements": [el]}]}
        return _chart_ghost.ghost_page(deck["slides"][0], deck, scale=1, supersample=1).convert("RGB")
    _probe = dict(hl_base, series=[{"name": "本期", "values": [10, 14, 18, 22]}])
    img_c = _hl_image(dict(_probe, highlight="C"))
    img_b = _hl_image(dict(_probe, highlight="B"))
    img_none = _hl_image(dict(_probe))
    img_bogus = _hl_image(dict(_probe, highlight="查无此类"))
    check("highlight: 预览强调位置随解析结果移动（不是画一个凑数的点）",
          img_c.tobytes() != img_b.tobytes() and img_c.tobytes() != img_none.tobytes())
    check("highlight: 写错名字时预览与产物都不产生强调",
          img_bogus.tobytes() == img_none.tobytes())
    # 展开顺序属于结构：产物是「类别优先」分组柱，预览必须同序（旧实现按序列优先，
    # 把同一类别的两条并排画成相隔很远的两组，作者会读到不存在的分组关系）。
    multi_rows = _chart_ghost._rows(dict(hl_base, series=[
        {"name": "甲", "values": [10, 14, 18, 22]},
        {"name": "乙", "values": [5, 9, 13, 20]}]))
    check("ghost: 多序列展开按类别优先（与产物分组柱同序）",
          [r["label"] for r in multi_rows][:4] == ["A · 甲", "A · 乙", "B · 甲", "B · 乙"]
          and [r["value"] for r in multi_rows][:4] == [10, 5, 14, 9],
          str([r["label"] for r in multi_rows]))
    # 环心 KPI：产物里是两个可编辑文本框，预览此前把孔留空——作者看不到自己写的
    # 读数是否存在、是否挤到孔外。像素级断言（孔心区域必须有墨迹）。
    donut = {"id": "d", "type": "chart", "x": 40, "y": 40, "width": 600, "height": 480,
             "chart_kind": "donut", "center_value": "1.8 亿元", "center_label": "总预算",
             "data": [{"label": "门店投入", "value": 14000}, {"label": "选址模型", "value": 3200},
                      {"label": "运营储备", "value": 800}]}
    donut_deck = {"canvas": {"width": 1280, "height": 720}, "theme": hl_theme,
                  "slides": [{"id": "d", "elements": [donut]}]}
    donut_img = _chart_ghost.ghost_page(donut_deck["slides"][0], donut_deck,
                                        scale=1, supersample=1).convert("L")
    import numpy as _np
    hole_patch = _np.asarray(donut_img)[255:305, 250:430]
    check("ghost: 环心 KPI 与产物同形（孔内有读数，不再留空）", int(hole_patch.min()) < 120,
          f"孔心最暗像素={int(hole_patch.min())}")
    from primitives import resolve_chart_highlight
    check("highlight: 消歧单一真源（命名先于数字；整数按序列数/类别数分派）",
          resolve_chart_highlight({"highlight": "C"}, ["A", "B", "C"], ["本期"]) == (-1, 2)
          and resolve_chart_highlight({"highlight": "本期"}, ["A"], ["本期", "对照"]) == (0, -1)
          and resolve_chart_highlight({"highlight": 2}, ["A", "B", "C"], ["本期"]) == (-1, 2)
          and resolve_chart_highlight({"highlight": 1}, ["A", "B"], ["甲", "乙"]) == (1, -1)
          and resolve_chart_highlight({"highlight": 9}, ["A"], ["甲"]) == (-1, -1)
          and resolve_chart_highlight({}, ["A"], None) == (-1, -1))

    from ghost import ghost_page
    native = {'canvas': {'width': 500, 'height': 220},
              'theme': {'colors': {'background': '#FFFFFF', 'ink': '#000000',
                                   'primary': '#000000', 'secondary': '#555555'}},
              'slides': [{'id': 'editorial', 'background': '#FFFFFF', 'elements': [
                  {'id': 'label', 'type': 'shape', 'shape': 'rect', 'x': 40, 'y': 40,
                   'width': 260, 'height': 100, 'fill': {'type': 'none'},
                   'stroke': {'type': 'none'}, 'text': 'ALIGN',
                   'text_size': 30, 'text_color': '#000000', 'padding': 40,
                   'align': 'left', 'text_anchor': 'top'},
                  {'id': 'tracking', 'type': 'text', 'text': '编辑排版',
                   'x': 350, 'y': 80, 'width': 140, 'height': 76, 'size': 24,
                   'color': '#000000', 'char_spacing': 6}]}]}
    out_editorial = tmp / 'editorial.pptx'
    compile_deck(native, out_editorial, speed='fast')
    prs_editorial = Presentation(str(out_editorial))
    sh1, sh2 = list(prs_editorial.slides[0].shapes)[:2]
    proof = ghost_page(native['slides'][0], native, scale=1, supersample=1)
    ink_x = [x for y in range(40, 140) for x in range(0, 310)
             if all(c < 120 for c in proof.getpixel((x, y)))]
    check("editorial: 空心形状不变卡片，内边距/文字主轴与原生 PPTX 对齐",
          round(sh1.text_frame.margin_left / 9525) == 40
          and min(ink_x) >= 78 and proof.getpixel((200, 100)) == (255, 255, 255))
    spc = [run._r.get_or_add_rPr().get('spc')
           for p in sh2.text_frame.paragraphs for run in p.runs]
    check("editorial: 作者显式声明的中文字距进入可编辑 OOXML，而非只对拉丁文生效",
          spc == ['600'], str(spc))

    styled_shape = {'id': 'shape_label', 'type': 'shape', 'shape': 'rect',
                    'x': 35, 'y': 20, 'width': 430, 'height': 174,
                    'fill': {'type': 'none'}, 'stroke': {'type': 'none'},
                    'text': 'AXIS\nBASELINE', 'padding': 12, 'text_size': 28,
                    'text_color': '#000000', 'text_opacity': .5, 'text_bold': True,
                    'text_wrap': False, 'text_line_height': 1.85,
                    'align': 'left', 'text_anchor': 'middle'}
    styled_text = {'id': 'same_text', 'type': 'text', 'text': styled_shape['text'],
                   'x': 35, 'y': 20, 'width': 430, 'height': 174, 'padding': 12,
                   'size': 28, 'color': '#000000', 'opacity': .5, 'bold': True,
                   'wrap': False, 'line_height': 1.85,
                   'align': 'left', 'anchor': 'middle'}
    def styled_spec(e):
        return dict(native, slides=[{'id': 'styled', 'background': '#FFFFFF',
                                     'elements': [e]}])
    shape_deck, text_deck = styled_spec(styled_shape), styled_spec(styled_text)
    label_image = ghost_page(shape_deck['slides'][0], shape_deck, scale=1, supersample=1)
    equivalent_image = ghost_page(text_deck['slides'][0], text_deck, scale=1, supersample=1)
    label_pptx = tmp / 'shape-text.pptx'
    compile_deck(shape_deck, label_pptx, speed='fast')
    label_native = Presentation(str(label_pptx)).slides[0].shapes[0]
    native_alpha = label_native.text_frame.paragraphs[0].runs[0]._r.find('.//' + qn('a:alpha'))
    check("editorial: 形状文字的粗细/行距/不换行/透明度与原生文本语义相同",
          label_image.tobytes() == equivalent_image.tobytes()
          and label_native.text_frame.paragraphs[0].runs[0].font.bold is True
          and native_alpha is not None and native_alpha.get('val') == '50000'
          and round(label_native.text_frame.paragraphs[0].line_spacing / 9525, 1) == 51.8)

    # 预览诚实度：声明「不渲染」的名单之外，每一个图表类型都必须真的画出来。
    # 断言方式是在占位函数上装探针：走占位分支的类型 == 声明表里的类型，
    # 一个不多一个不少。漏声明 → 预览画的是不存在的图形（作者按假图改构图）；
    # 声明了却仍渲染 → 声明表在骗人。
    import ghost as _ghost
    from primitives import CHART_KINDS
    base_spec = {'canvas': {'width': 480, 'height': 260},
                 'theme': {'colors': {'background': '#FFFFFF', 'ink': '#000000',
                                      'primary': '#111111', 'secondary': '#555555',
                                      'accent': '#9C3B26'}},
                 'fonts': {'cn': 'Source Han Sans SC', 'latin': 'Arial'}}
    placeholder_calls: list = []
    _real_unrendered = _ghost._draw_unrendered

    def _probe_unrendered(d, box, kind, scale):
        placeholder_calls.append(str(kind))
        return _real_unrendered(d, box, kind, scale)

    _ghost._draw_unrendered = _probe_unrendered
    try:
        for kind in sorted(CHART_KINDS):
            element = {'id': 'c', 'type': 'chart', 'chart_kind': kind, 'x': 20, 'y': 20,
                       'width': 440, 'height': 220, 'value': 7, 'label': 'A',
                       'data': [{'label': 'A', 'value': 3}, {'label': 'B', 'value': 7},
                                {'label': 'C', 'value': 5}]}
            slide = {'id': 'k', 'background': '#FFFFFF', 'elements': [element]}
            _ghost.ghost_page(slide, dict(base_spec, slides=[slide]), scale=1, supersample=1)
    finally:
        _ghost._draw_unrendered = _real_unrendered
    abstracted = set(placeholder_calls)
    check("ghost: 「不渲染」声明表与实际渲染一一对应（预览不替假图背书）",
          abstracted == set(_ghost.PREVIEW_ABSTRACT) <= set(CHART_KINDS),
          f"声明={sorted(_ghost.PREVIEW_ABSTRACT)} 实际占位={sorted(abstracted)}")

    # 图表单一规范：五维事实只在 primitives.CHART_SPEC 定义一次，三层集合全部派生。
    # 任一层再各写一份（guard 的 limits、compiler 的路由集、ghost 的占位表），
    # 就会重现「放行但截尾 / 预览与产物不一致」；这里逐维对拍派生关系。
    from primitives import (CHART_LIMITS as _LIMITS, CHART_SPEC as _SPEC,
                            NUMERIC_CHART_KINDS as _NUM, ROWS_CHART_KINDS as _ROWSK,
                            chart_rows as _chart_rows)
    import verify as _verify_mod, compiler as _compiler_mod
    check("chart_spec: 图表分类单一来源（guard/compiler/ghost 全部派生自 CHART_SPEC）",
          set(_compiler_mod.NATIVE_CHART_TYPES) == {k for k, v in _SPEC.items() if v["renderer"] == "native"}
          and _compiler_mod.SHAPE_CHARTS == {k for k, v in _SPEC.items() if v["renderer"] == "shape"}
          and _compiler_mod.ELEMENT_PAYLOAD_CHARTS == {k for k, v in _SPEC.items() if v["payload"] in ("points", "layers")}
          and _compiler_mod.NUMERIC_UNIT_KINDS == _NUM == _verify_mod.NUMERIC_CHART_KINDS
          and set(_verify_mod.ROWS_CHART_KINDS) == set(_ROWSK)
          and _verify_mod.CHART_LIMITS == _LIMITS == {k: v["limit"] for k, v in _SPEC.items() if v["limit"] is not None}
          and set(_ghost.PREVIEW_ABSTRACT) == {k for k, v in _SPEC.items() if v["preview"] == "abstract"},
          f"spec={len(_SPEC)} kinds")
    check("chart_rows: 编译与预览共用唯一行解析（desc/无效值/display 同口径）",
          _compiler_mod._rows is _chart_rows
          and _chart_rows({"data": [{"label": "A", "value": "3", "desc": "注"},
                                    {"label": "B", "value": None}]})[0]["desc"] == "注"
          and _ghost._rows({"data": [{"label": "A", "value": 3}]}) == _chart_rows({"data": [{"label": "A", "value": 3}]}))

    # 风险色（negative）由主题推导：含负值的图表不再因「主题未声明风险色」被硬门挡住，
    # 作者显式声明的风险色仍然压过推导。
    neg_spec = {'canvas': {'width': 480, 'height': 260},
                'theme': {'colors': dict(base_spec['theme']['colors'])},
                'slides': [{'id': 'neg', 'background': '#FFFFFF', 'elements': [
                    {'id': 'loss', 'type': 'chart', 'chart_kind': 'bar', 'x': 20, 'y': 20,
                     'width': 440, 'height': 220, 'show_values': True, 'metric': '净收益',
                     'unit': '万元', 'source': '来源：财务测算', 'period': '2027–2029',
                     'basis': '含折现',
                     'data': [{'label': 'M6', 'value': -820}, {'label': 'M36', 'value': 1480}]}]}]}
    neg_out = tmp / 'negative.pptx'
    neg_report = compile_deck(neg_spec, neg_out, speed='fast')
    check("compiler: 六 token 主题也能画负值图（风险色由主题派生，不逼作者补一轮）",
          neg_report["passed"] and not neg_report["warnings"], str(neg_report["warnings"]))
    from primitives import RenderContext, contrast, _rgb_to_oklab, _tuple
    ctx_neg = RenderContext(neg_spec["theme"], neg_spec["canvas"])
    derived = ctx_neg.chart_role("negative")
    check("compiler: 派生风险色与纸面的对比度足够（不是看不见的一档）",
          derived is not None
          and contrast("#" + str(derived), neg_spec["theme"]["colors"]["background"]) >= 3.0)
    # 强调色本身就是红时，风险色必须继续加深到与它分得开：否则「亏损」与
    # 「被强调的答案」在读数上是同一个色，颜色不再区分任何东西。
    red_accent = json.loads(json.dumps(neg_spec))
    red_accent["theme"]["colors"]["accent"] = "#9C3B26"
    red_ctx = RenderContext(red_accent["theme"], red_accent["canvas"])
    gap = sum((i - j) ** 2 for i, j in zip(
        _rgb_to_oklab(_tuple(str(red_ctx.chart_role("negative")))),
        _rgb_to_oklab(_tuple(red_ctx.colors["accent"])))) ** 0.5
    check("compiler: 红强调色主题下风险色与强调色 OKLab 可分（风险 ≠ 答案）",
          gap >= 0.12, f"OKLab 距离={gap:.3f} ≠ {red_ctx.chart_role('negative')} / "
                       f"{red_ctx.colors['accent']}")
    declared_spec = json.loads(json.dumps(neg_spec))
    declared_spec["theme"]["chart_palette"] = {"negative": "#6A1B1B"}
    declared = RenderContext(declared_spec["theme"], declared_spec["canvas"]).chart_role("negative")
    check("compiler: 作者声明的风险色压过派生",
          str(declared).upper().endswith("6A1B1B"), str(declared))

    # 预览与产物必须说同一件事：负值条既要有长度（零基），也要是风险色。
    neg_image = ghost_page(neg_spec["slides"][0], neg_spec, scale=1, supersample=1).convert("RGB")
    role = RenderContext(neg_spec["theme"], neg_spec["canvas"]).chart_role("negative")
    want = tuple(int(v, 16) for v in (str(role).zfill(6)[i:i + 2] for i in (0, 2, 4)))
    want = tuple(round(c * 0.95 + 255 * 0.05) for c in want)
    band = [neg_image.getpixel((x_col, 60)) for x_col in range(150, 240)]
    hit = next((c for c in band if sum(abs(a - b) for a, b in zip(c, want)) <= 6), None)
    check("ghost: 负值条在预览里有长度且用与产物相同的风险色（预览不替亏损改口径）",
          hit is not None, f"want={want} band={band[:4]}")

    missing = [name for files in ENGINE_SCOPES.values() for name in files
               if not (Path(__file__).resolve().parent / name).is_file()]
    check("engine: 指纹 scope 引用的文件都存在（删改模块必须同步）", not missing, str(missing))

    work = tmp / "cache"
    work.mkdir(exist_ok=True)
    report = compile_deck(spec, out1, speed="fast")
    view = spec_view(spec, base_path=tmp)
    record_compile(work, out1, view, report)
    reused = compile_reuse(work, out1, view, fast_probe=True, speed="fast")
    check("cache: 同投影 + 同产物字节 ⇒ 复用编译（不重编）",
          bool(reused and reused.get("reused")))
    changed = json.loads(json.dumps(spec))
    changed["slides"][0]["elements"][0]["text"] = "改过的标题"
    check("cache: spec 变了不命中",
          compile_reuse(work, out1, spec_view(changed, base_path=tmp),
                        fast_probe=True, speed="fast") is None)
    # 编译产物与 speed 无关（字节一致已由 media 组锁定）：缓存不再按档分区，
    # strict 终稿后切 fast 迭代（或反向）直接复用，不整稿重编。
    check("cache: 编译缓存跨 speed 复用（产物字节与档无关，不重复编译）",
          compile_reuse(work, out1, view, fast_probe=False, speed="strict") is not None
          and (compile_reuse(work, out1, view, fast_probe=False, speed="strict") or {}).get("performance", {}).get("speed") == "strict")

    # strict 探测已算出当前产物 SHA；attestation 复用，不重复触发摘要/状态探测。
    from vao import _compile_step
    import primitives
    cached_out = tmp / "strict-cache.pptx"
    first_report, _ = _compile_step(spec, cached_out, speed="strict", spec_path=str(build))
    from unittest.mock import patch
    with patch.object(primitives, "file_digest", wraps=primitives.file_digest) as digest_spy:
        cached_report, cache_timing = _compile_step(
            spec, cached_out, speed="strict", spec_path=str(build))
    hashed_output = [call for call in digest_spy.call_args_list
                     if Path(call.args[0]).expanduser().resolve() == cached_out.resolve()]
    check("cache: strict 命中对 PPTX 只做一次内容哈希",
          first_report.get("output_sha256") == cached_report.get("output_sha256")
          and cached_report.get("reused") and len(hashed_output) == 1
          and cache_timing.get("compile_ms") == 0)

    # 资产链：媒体必要性 → 清单；坏图 release 必 BLOCK
    from intelligence import think
    from assets import build_manifest, image_qc, qc_retry_decision
    bundle = think(BRIEF)
    manifest = build_manifest(BRIEF, bundle)
    generated = {sid for a in manifest["assets"] if a.get("decision") == "generate"
                 for sid in a.get("slide_ids") or []}
    check("assets: 只有通过必要性测试的页进清单",
          generated <= set(bundle["assets_hint"]["generate"])
          and all("necessity" in s or s.get("reason") for s in manifest["skipped_pages"]),
          str(manifest["assets"])[:160])
    from assets import asset_fingerprint, asset_card
    page = bundle["pages"][0]
    card, contract = asset_card(page, BRIEF, bundle["deck"], "asset-x")
    check("assets: 提示词确定性 + 带负向清单",
          asset_fingerprint(card, contract) == asset_fingerprint(card, contract)
          and "no " in (card.get("negative") or [""])[0] if card.get("negative") else True)
    from PIL import Image, ImageDraw
    img = tmp / "img.png"
    canvas = Image.new("RGB", (1280, 720), (240, 236, 228))
    ImageDraw.Draw(canvas).ellipse((700, 200, 1000, 500), fill=(120, 90, 60))
    canvas.save(img)
    from unittest.mock import patch
    image_spec = {"canvas": {"width": 1280, "height": 720}, "theme": {},
                  "slides": [{"id": "one_image", "elements": [
                      {"id": "photo", "type": "image", "src": str(img),
                       "x": 80, "y": 120, "width": 520, "height": 320,
                       "fit": "cover"}]}]}
    image_snapshot: dict = {}
    image_digest: dict = {}
    path_read_count = {str(img.resolve()): 0}
    original_read_bytes = Path.read_bytes
    def count_image_read(path, *args, **kwargs):
        key = str(Path(path).expanduser().resolve())
        if key in path_read_count:
            path_read_count[key] += 1
        return original_read_bytes(path, *args, **kwargs)
    with patch.object(Path, "read_bytes", count_image_read):
        spec_view(image_spec, base_path=tmp, image_bytes=image_snapshot, digests=image_digest)
        compile_deck(image_spec, tmp / "one_image.pptx", speed="fast",
                     spec_path=str(build), image_bytes=image_snapshot)
    from primitives import digest_bytes
    image_key = str(img.resolve())
    diag1_read = path_read_count[image_key]
    diag1_sha = image_digest.get(image_key) == digest_bytes(image_snapshot.get(image_key, b""))
    check("media: 指纹预读的源图字节由编译复用（同轮仅一次文件读取）",
          diag1_read == 1 and diag1_sha,
          f"read_count={diag1_read} (expect 1), sha_match={diag1_sha}, "
          f"snapshot_keys={sorted(image_snapshot.keys())}, digest_keys={sorted(image_digest.keys())}")
    # 模拟 QC 已有 SHA、但因 QC cache hit 没把图像快照装入内存：compile 读一次后
    # 把源 bytes 放回共享快照，page-key 不再为 preview 读取第二次。
    image_snapshot_cached: dict = {}
    image_digest_cached = {image_key: image_digest[image_key]}
    path_read_count[image_key] = 0
    media_spec_for_preview = {**image_spec, "_image_bytes": image_snapshot_cached,
                              "_base_path": str(tmp)}
    media_slide = image_spec["slides"][0]
    with patch.object(Path, "read_bytes", count_image_read):
        spec_view(image_spec, base_path=tmp, image_bytes=image_snapshot_cached,
                  digests=image_digest_cached)
        compile_deck(image_spec, tmp / "one_image_cached_qc.pptx", speed="fast",
                     spec_path=str(build), image_bytes=image_snapshot_cached)
        from ghost import _page_key
        _page_key(media_slide, media_spec_for_preview, 0.5, 1, 1)
    # 拆分组合断言，Windows CI 失败时能直接看到哪个环节断链：
    #   1) 读盘次数（1 = compiler 首次读，后续 preview 走快照）
    #   2) compiler 读后 canonical key 是否回灌到共享快照
    #   3) 共享快照里的字节 SHA 是否与预缓存的 digest 一致
    diag_read = path_read_count[image_key]
    diag_snap = bool(image_snapshot_cached.get(image_key))
    diag_sha = image_digest_cached.get(image_key) == digest_bytes(image_snapshot_cached.get(image_key, b""))
    extra_keys = sorted(set(image_snapshot_cached.keys()) - {image_key})
    check("media: QC 缓存 SHA → 编译读一次 → preview 复用同一图像字节",
          diag_read == 1 and diag_snap and diag_sha,
          f"read_count={diag_read} (expect 1), snapshot_hit={diag_snap}, sha_match={diag_sha}, "
          f"canonical_key={image_key!r}, extra_keys_in_snapshot={extra_keys}")

    # 图片变换编码口径唯一（v9.12）：PNG 无损，zlib 级别不改变任何解码像素；
    # fast/strict 两档的产物字节必须完全一致（v5.9 的「差异只在耗时与字节」
    # 双口径在此收口）。用需要 cover 变换的图（比例不符）逼出重编码路径。
    _ratio_src = tmp / "ratio_src.png"
    Image.new("RGB", (100, 60), (200, 40, 40)).save(_ratio_src)
    _ratio_spec = {"canvas": {"width": 1280, "height": 720}, "theme": {},
                   "slides": [{"id": "r1", "elements": [
                       {"id": "pic", "type": "image", "src": str(_ratio_src),
                        "x": 40, "y": 40, "width": 80, "height": 80, "fit": "cover"}]}]}
    from compiler import compile_deck as _compile_deck
    _deck_fast, _deck_strict = tmp / "img_fast.pptx", tmp / "img_strict.pptx"
    _compile_deck(_ratio_spec, _deck_fast, speed="fast", spec_path=str(build))
    _compile_deck(_ratio_spec, _deck_strict, speed="strict", spec_path=str(build))
    check("media: 图片变换后 fast 与 strict 产物字节一致（编码口径唯一）",
          file_digest(_deck_fast) == file_digest(_deck_strict))

    safe = {"x": 0.06, "y": 0.08, "width": 0.34, "height": 0.78}
    qc = image_qc(img, safe_rect=safe, text_is_dark=True)
    decision_ok = qc_retry_decision(qc, attempt=0, phase="release")
    check("assets: 合规画心 QC 通过（无阻断项，动作 accept）",
          qc.get("status") == "ok" and decision_ok.get("action") == "accept",
          str(qc.get("checks"))[:160])
    fast_qc = image_qc(img, safe_rect=safe, text_is_dark=True, max_side=960)
    check("assets: fast/strict 的硬缝判定一致（不因降采样制造重试）",
          fast_qc.get("status") == qc.get("status")
          and next(c for c in fast_qc["checks"] if c["check"] == "hard_seam")["status"]
          == next(c for c in qc["checks"] if c["check"] == "hard_seam")["status"])
    bad_img = tmp / "bad.png"          # 文字安全区被高对比纹理压住 = 不可用
    canvas2 = Image.new("RGB", (1280, 720), (245, 243, 240))
    paint = ImageDraw.Draw(canvas2)
    for i in range(0, 460, 12):
        paint.rectangle([i, 0, i + 6, 720], fill=(30, 30, 30))
    canvas2.save(bad_img)
    bad = image_qc(bad_img, safe_rect=safe, text_is_dark=True)
    decision = qc_retry_decision(bad, attempt=1, phase="release")
    check("assets: 坏图 release 档不自动重出（判负交 QA 收口）",
          bad.get("status") != "ok" and decision.get("action") in ("block", "fail",
                                                                   "hold", "none"),
          f"{bad.get('status')} {decision}")

    # 离线 qc 环与 check 同判据（CASE_004 实测教训）：清单声明 16:9 的资产，
    # 4:3 成图必须在 qc 环就被拦——否则作者拿到「qc 全绿 → check BLOCK」的
    # 无效往返，最短反馈环变成最长的弯路。
    import vao as _vao
    _ratio_img = tmp / "qc-loop-ratio.png"
    Image.new("RGB", (1024, 768), (214, 206, 192)).save(_ratio_img)
    _m = {"schema": "vao-assets-v2", "assets_dir": str(tmp),
          "assets": [{"asset_id": "qc-loop-ratio", "decision": "generate",
                      "prompt": "test", "negative": ["test"], "slide_ids": ["s1"],
                      "ratio": "16:9", "asset_role": "background",
                      "safe_area": {"x": 0.06, "y": 0.08, "width": 0.34, "height": 0.78},
                      "meta": {"text_color": "dark"}}]}
    _mpath = tmp / "qc_loop_manifest.json"
    _mpath.write_text(json.dumps(_m, ensure_ascii=False), encoding="utf-8")
    _loop = _vao.run_qc(str(_mpath), str(tmp), speed="strict")
    _loop_r = _loop["results"][0]
    check("qc: 离线环同判比例（清单 16:9 对 4:3 成图必须 retry）",
          "aspect_ratio" in _loop_r.get("blocking", []),
          json.dumps(_loop_r, ensure_ascii=False)[:160])

    # prepare 落盘：plan/brief 一致性走文件凭证，不再有 dict 级二次比对
    from assets import prepare_manifest
    import copy as _copy
    prepped = prepare_manifest(_copy.deepcopy(manifest), {"different": True}, bundle,
                               str(tmp / "brief.yml"), None, str(tmp / "m.json"))
    check("assets: prepare 不因 brief dict 不一致抛错（文件凭证是唯一口径）",
          prepped.get("schema") == "vao-assets-v2")

    # 解码缓存：只缓存未裁切底图（同源不同裁切各算各的，不互相污染）
    from compiler import _fit_image_bytes
    import io as _io
    half = Image.new("RGB", (100, 50), (255, 0, 0))
    half.paste(Image.new("RGB", (50, 50), (0, 0, 255)), (50, 0))
    buf = _io.BytesIO()
    half.save(buf, "PNG")
    blob = buf.getvalue()
    dcache: dict = {}
    left_png = _fit_image_bytes(blob, 50, 50, "cover", crop=(0, 0, 0.5, 0),
                                decode_cache=dcache, cache_key="k")
    right_png = _fit_image_bytes(blob, 50, 50, "cover", crop=(0.5, 0, 0, 0),
                                 decode_cache=dcache, cache_key="k")
    left_px = Image.open(_io.BytesIO(left_png)).convert("RGB").getpixel((25, 25))
    right_px = Image.open(_io.BytesIO(right_png)).convert("RGB").getpixel((25, 25))
    check("compiler: 同源不同裁切像素各自正确（缓存不污染）",
          left_px[0] > 200 and right_px[2] > 200, f"{left_px} vs {right_px}")

    # 关键页取证：封面 / 密数据 / 收尾
    from ghost import (PAGE_CACHE_DIR, _page_key, key_pages, ghost_deck,
                       key_page_roles, key_selection, make_contact_sheet)
    pages = key_pages(spec["slides"], limit=3)
    labels = key_page_roles(spec["slides"], pages)
    check("ghost: 关键页取证含封面与收尾",
          1 in pages and len(spec["slides"]) in pages and
          set(labels.values()) <= {"cover", "closing", "hero_image", "dense_data",
                                   "section", "content"}, str(labels))
    rendered_images: list = []
    page_stats: dict = {}
    rendered = ghost_deck(spec, tmp / "preview", pages=pages, scale=0.4, supersample=1,
                          images_out=rendered_images, stats=page_stats)
    check("ghost: 预览只画关键页（页级缓存内建）",
          len(rendered) == len(pages) and all(Path(p).is_file() for p in rendered))
    page_cache_files = list((tmp / "preview" / PAGE_CACHE_DIR).glob("*.png"))
    check("ghost: 页缓存与交付预览复用同一份 PNG 编码",
          bool(rendered) and any(Path(rendered[0]).read_bytes() == p.read_bytes()
                                 for p in page_cache_files))
    sheet_stats: dict = {}
    contact = make_contact_sheet(rendered, tmp / "preview" / "contact.png",
                                 images=rendered_images,
                                 cache_dir=tmp / "preview" / PAGE_CACHE_DIR,
                                 stats=sheet_stats)
    sheet_cache = list((tmp / "preview" / PAGE_CACHE_DIR).glob("sheet-*.png"))
    check("ghost: 联络表只编码一次并以字节副本缓存",
          bool(contact and Path(contact).is_file() and sheet_cache
               and Path(contact).read_bytes() == sheet_cache[0].read_bytes()))
    sel_pages, sel_roles = key_selection(spec["slides"], limit=3)
    check("ghost: 单次取证与两次调用结论一致（结构只扫一遍）",
          sel_pages == pages and sel_roles == labels)
    # 字面回落必须可查：本机缺 CJK 字面时页图仍「成功」产出，但中文字形是缺的。
    # 预览记录不带这一条，读图的人就分不清「排版如此」与「本机没字体」。
    import ghost as ghost_mod
    import primitives as font_facts
    saved_candidates = dict(ghost_mod._FONT_CANDIDATES)
    saved_bold_candidates = dict(ghost_mod._FONT_BOLD_CANDIDATES)
    saved_stamp, saved_fallbacks = font_facts._PREVIEW_FONT_STAMP, dict(ghost_mod._FONT_FALLBACKS)
    saved_font_cache = dict(ghost_mod._FONT_CACHE)
    # 候选身份变更是测试输入，不依赖宿主是否安装任何字库。
    probe_font = tmp / "identity-probe.ttf"
    probe_font.write_bytes(b"identity fixture, not a usable font")
    ghost_mod._FONT_CANDIDATES["cjk"] += (str(probe_font),)
    ghost_mod._font_stamp(refresh=True)
    key_with_fonts = _page_key(spec["slides"][0], spec, 0.4, 1, 1)
    try:
        ghost_mod._FONT_CANDIDATES.update({role: () for role in saved_candidates})
        ghost_mod._FONT_BOLD_CANDIDATES.update({role: () for role in saved_bold_candidates})
        font_facts._PREVIEW_FONT_STAMP = None
        ghost_mod._FONT_CACHE.clear()          # 字面已缓存时不会走回落分支
        ghost_mod._FONT_FALLBACKS.clear()
        ghost_mod.ghost_page(spec["slides"][0], spec, scale=0.4, supersample=1)
        fallbacks = ghost_mod.font_fallbacks()
        check("ghost: 无字面可用时记账（预览不得静默丢掉中文字形）",
              sum(fallbacks.get(r, 0) for r in ("cjk", "cjk_serif")) >= 1, str(fallbacks))
        check("ghost: 字面身份进页缓存键（换字体必须重画）",
              _page_key(spec["slides"][0], spec, 0.4, 1, 1) != key_with_fonts)
    finally:
        ghost_mod._FONT_CANDIDATES.update(saved_candidates)
        ghost_mod._FONT_BOLD_CANDIDATES.update(saved_bold_candidates)
        font_facts._PREVIEW_FONT_STAMP = saved_stamp
        ghost_mod._FONT_CACHE.clear()
        ghost_mod._FONT_CACHE.update(saved_font_cache)
        ghost_mod._FONT_FALLBACKS.clear()
        ghost_mod._FONT_FALLBACKS.update(saved_fallbacks)

    media_src = tmp / "preview-source.png"
    Image.new("RGB", (32, 32), (180, 40, 40)).save(media_src)
    blob_a = media_src.read_bytes()
    media_slide = {"id": "media", "elements": [
        {"id": "photo", "type": "image", "src": str(media_src)}]}
    media_spec_a = {"canvas": {"width": 1280, "height": 720}, "theme": {},
                    "_image_bytes": {str(media_src): blob_a}}
    key_a = _page_key(media_slide, media_spec_a, 0.5, 1, 1)
    Image.new("RGB", (32, 32), (40, 40, 180)).save(media_src)
    blob_b = media_src.read_bytes()
    media_spec_b = {"canvas": {"width": 1280, "height": 720}, "theme": {},
                    "_image_bytes": {str(media_src): blob_b}}
    key_b = _page_key(media_slide, media_spec_b, 0.5, 1, 1)
    check("ghost: 同一路径的图片字节变化会使页缓存失效", key_a != key_b)
    # 与先画透明整幅图再合成的旧路径逐像素比对：省掉背景合成不能改变任何墨迹。
    from ghost import _draw_text, _font, _rgba, _font_family
    from primitives import RenderContext
    ctx = RenderContext(spec.get("theme"), spec.get("canvas"))
    text_elem = {"id": "probe", "type": "text", "x": 12, "y": 24,
                 "width": 160, "height": 80, "text": "61%", "size": 32, "color": "ink"}
    same = True
    for opacity in (None, 0.52):
        bg_color = (244, 239, 230, 255)
        expected = Image.new("RGBA", (240, 130), bg_color)
        old_overlay = Image.new("RGBA", expected.size, (0, 0, 0, 0))
        ImageDraw.Draw(old_overlay, "RGBA").text(
            (12, 24), "61%", font=_font(32, family=_font_family(text_elem, ctx)),
            fill=_rgba(ctx, "ink", opacity))
        expected.alpha_composite(old_overlay)
        actual = Image.new("RGBA", expected.size, bg_color)
        _draw_text(actual, {**text_elem, "opacity": opacity}, ctx, 1)
        same = same and actual.tobytes() == expected.tobytes()
    check("ghost: 不透明直画与半透明旧合成路径像素一致（速度不能改变设计）", same)

    # 字体角色与 compiler 共用 RenderContext.families；没有 font 字段也要读主题。
    import ghost as _g
    serif_ctx = RenderContext({"fonts": {"cn": "Noto Serif CJK SC", "latin": "Arial"}}, {})
    sans_ctx = RenderContext({"fonts": {"cn": "Noto Sans CJK SC", "latin": "Georgia"}}, {})
    check("ghost: 中文主题默认宋体不会静默画成黑体",
          _g._font_family({"text": "中文"}, serif_ctx) == "_serif")
    check("ghost: Latin font/family 覆盖不改变中文主题字面",
          all(_g._font_family(e, serif_ctx) == "_serif" for e in
              ({"text": "中文 AI 2026", "font": "latin"},
               {"text": "中文", "family": "Arial"}))
          and _g._font_family({"text": "中文", "font": "latin"}, sans_ctx) == "")
    check("ghost: 纯拉丁遵循主题与显式覆盖（sans-serif 不是 serif）",
          _g._font_family({"text": "Evidence"}, sans_ctx) == "_serif"
          and _g._font_family({"text": "Evidence", "font": "latin"}, serif_ctx) == ""
          and _g._font_family({"text": "Evidence", "family": "sans-serif"}, sans_ctx) == ""
          and _g._font_family({"text": "Evidence", "family": "Times New Roman"}, serif_ctx) == "_serif")
    alias_ctx = RenderContext({"fonts": {"body": "宋体", "display": "Georgia"}}, {})
    check("ghost: body/display 别名与编译器一致",
          all(_g._font_family({"text": t}, alias_ctx) == "_serif" for t in ("中文", "English")))
    # 跨平台锁住 TTC index 参数，不靠 CI 恰好安装某字体；缺字库仍走原回退。
    from unittest.mock import patch
    saved_fonts = dict(_g._FONT_CACHE)
    try:
        with patch.object(Path, "exists", return_value=True), \
                patch.dict(_g._FONT_CANDIDATES, {
                    "cjk": ("/fonts/NotoSansCJK-Regular.ttc",),
                    "cjk_serif": ("/fonts/NotoSerifCJK-Regular.ttc",)}), \
                patch.dict(_g._FONT_BOLD_CANDIDATES, {
                    "cjk": ("/fonts/NotoSansCJK-Bold.ttc",),
                    "cjk_serif": ("/fonts/NotoSerifCJK-Bold.ttc",)}), \
                patch.object(_g.ImageFont, "truetype") as load_font:
            _g._FONT_CACHE.clear()
            for family in ("", "_serif"):
                for bold in (False, True):
                    _g._font(20, cjk=True, family=family, bold=bold)
            check("ghost: Noto CJK 常规/粗体/衬线四路径均取 SC index 2",
                  len(load_font.call_args_list) == 4
                  and all(c.kwargs.get("index") == 2 for c in load_font.call_args_list))
        with patch.object(Path, "exists", return_value=True), \
                patch.dict(_g._FONT_CANDIDATES, {"cjk": ("C:/Windows/Fonts/msyh.ttc",)}), \
                patch.object(_g.ImageFont, "truetype") as load_font:
            _g._FONT_CACHE.clear()
            _g._font(20, cjk=True)
            check("ghost: 其他 TTC 不误套 Noto 索引", load_font.call_args.kwargs.get("index") == 0)
        with patch.object(Path, "exists", return_value=False), \
                patch.object(_g.ImageFont, "load_default") as fallback:
            _g._FONT_CACHE.clear()
            _g._font(20, cjk=True)
            _g._font(20, cjk=True)
            check("ghost: 缺字体安全回退且进程内复用", fallback.call_count == 1)
    finally:
        _g._FONT_CACHE.clear()
        _g._FONT_CACHE.update(saved_fonts)

    # CJK 折行：二分拟合与逐字贪心同断点（禁则修正之后仍同断点）。
    # visible_width 随前缀长度单调，二分与线性扫描找到同一个「最长可容纳前缀」；
    # 禁则只在其上做局部回退。随机混合文本 × 字号 × 宽度 × tracking 对拍内联的
    # 旧线性参考实现，输出必须逐字符一致——这条不变式一破，预览换行位置就变，
    # 像素证据不再可信。
    from ghost import (_wrap_lines as _wrap, _has_cjk as _has_cjk2,
                       _cjk_break_fix as _fix, _NO_LINE_START as _nls,
                       _UNIT_HEAD as _unit_head, _UNIT_INNER as _unit_inner)
    # 禁则单一来源：预览折行与判断层断行建议共用 primitives 的同一份字符串，
    # 任一层再抄一份就会出现「建议不断、预览断了」的分叉。
    from primitives import NO_LINE_END_CJK as _PLE, NO_LINE_START_CJK as _PLS
    import intelligence as _intel_mod
    check("typography: CJK 禁则单一来源（预览与判断层共用同一份）",
          _nls == _PLS and _intel_mod._NO_LINE_START_CJK == _PLS
          and _intel_mod._NO_LINE_END_CJK == _PLE,
          "ghost/intelligence 必须取自 primitives.NO_LINE_*_CJK")

    def _linear_cjk(text, font, width, tracking=0.0):
        limit = max(1, width)
        def vw(value):
            return font.getlength(value) + max(0, len(value) - 1) * tracking
        out = []
        for raw in str(text).split("\n"):
            if not raw:
                out.append("")
                continue
            pos, n = 0, len(raw)
            while pos < n:
                seg = raw[pos:]
                if vw(seg) <= limit:
                    out.append(seg)            # 末行整段可容纳
                    break
                if len(seg) == 1:
                    out.append(seg)
                    break
                lo = 1                          # 逐字扫描：找到最长可容纳前缀
                while lo < len(seg) and vw(seg[:lo + 1]) <= limit:
                    lo += 1
                lo = _fix(seg, lo)              # 与二分路径同一套禁则回退
                out.append(seg[:lo].rstrip())
                pos += lo
        return out

    # 禁则本身：行首不得出现禁则字符；数字与单位不得被拆到两行
    _rule_font = _font(18, cjk=True)
    _rule_text = ("口岸清关时效在 1.8 至 6.4 天之间波动，波动幅度 256%，"
                  "库存周转天数在 21 至 47 天之间摆动，「次日达」无法售卖。")
    _rule_bad = None
    for _w in (220, 260, 300, 340, 420, 520, 640, 800):
        _lines = _wrap(_rule_text, _rule_font, _w)
        for _ln in _lines[1:]:
            if _ln[:1] and _ln[0] in _nls:
                _rule_bad = ("行首标点", _w, _ln); break
            _prev = _lines[_lines.index(_ln) - 1]
            if _prev[-1:].isdigit() and _unit_head.match(_ln.strip()):
                _rule_bad = ("数字单位拆行", _w, _prev, _ln); break
        if _rule_bad:
            break
    check("ghost: 中文禁则折行（行首无禁则标点、数字与单位同行）",
          _rule_bad is None, str(_rule_bad)[:160])

    # 断口常常落在**空格**上（中英混排的「9 个」「11.2 元」）：此时真正的邻居是
    # 空格两边的实字，不是断口那一个字符。旧实现取 seg[lo] / seg[lo-1]，量到的
    # 是两个空格，禁则静默失效——预览看着合规，产物里读出来是「9」「个」两个信息。
    # 实测（CASE_006 第 4 页）：旧预览断成「…可控节点 9 | 个，复制周期…」。
    _sp_font = _font(15, cjk=True)
    _sp_text = ("印尼雅加达自营：首城盈利 14 个月，履约成本 11.2 元，可控节点 9 个，"
                "复制周期 5 个月。")
    _sp_bad = None
    for _w in range(200, 470, 3):
        _ls = _wrap(_sp_text, _sp_font, _w)
        for _a, _b in zip(_ls, _ls[1:]):
            _l, _r = _a.rstrip(), _b.lstrip()
            if _l[-1:].isdigit() and _unit_head.match(_r):
                _sp_bad = ("数字与量词被空格隔开后拆行", _w, _a, _b); break
            if _r[:1] and _r[0] in _nls:
                _sp_bad = ("行首禁则标点", _w, _b); break
        if _sp_bad:
            break
    check("ghost: 禁则看断口两侧的实字（空格不是邻居，数字与量词不隔空拆行）",
          _sp_bad is None, str(_sp_bad)[:160])

    # 双字单位不腰斩：断在「个｜月」「万｜元」「小｜时」中间，单位被读成两个词，
    # 「14 个月」会排成「14 个 / 月」——比数字与单位分家更常见，因为断口落在
    # 单位的两个汉字之间，任何只看断口两侧是不是「数字 + 单位」的规则都抓不到。
    _du_font = _font(15, cjk=True)
    _du_text = "首城盈利周期 14 个月，投入 3,800 万美元，分拣时长 45 分钟。"
    _du_bad = None
    for _w in range(180, 520, 2):
        _ls = _wrap(_du_text, _du_font, _w)
        for _a, _b in zip(_ls, _ls[1:]):
            _l, _r = _a.rstrip(), _b.lstrip()
            if _l and _r and (_l[-1] + _r[0]) in _unit_inner:
                _du_bad = ("双字单位被腰斩", _w, _a, _b); break
        if _du_bad:
            break
    check("ghost: 双字单位内部不断行（个月/万元/分钟/小时不读成两个词）",
          _du_bad is None, str(_du_bad)[:160])

    # 数字内部不是断点：「回收期从 2.6 年拉到 3.9 年」曾被断成「2.」/「6 年」，
    # 「压到 18%」曾被断成「1」/「8%」。行首行尾都像正常文本，靠读行发现不了——
    # 而 PowerPoint 不在数字内部断行，预览断了就在替真实排版说谎（实测 E2E 发现）。
    _num_font = _font(18, cjk=True)
    _num_text = "订单波动超过 ±20%，回收期从 2.6 年拉到 3.9 年，投入 4,800 万元。"
    _num_bad = None
    for _w in range(120, 620, 2):
        _ls = _wrap(_num_text, _num_font, _w)
        for _a, _b in zip(_ls, _ls[1:]):
            _l, _r = _a.rstrip(), _b.lstrip()
            if _l[-1:] in (".", ",") and _r[:1].isdigit():
                _num_bad = ("小数被断成两行", _w, _a, _b); break
            if _l[-1:].isdigit() and _r[:1] in ("%", "％"):
                _num_bad = ("百分号与数字拆行", _w, _a, _b); break
        if _num_bad:
            break
    check("ghost: 数字内部不断行（2.6 / 18% / 4,800 不被拆到两行）",
          _num_bad is None, str(_num_bad)[:160])

    import random as _random
    _rng = _random.Random(20261013)
    _pool = "本议案请求批准3800万美元首期预算，覆盖新加坡、雅加达；份额从7%到19%（括号）「引号」混合latin words在内。"
    _fonts = [(_font(s, cjk=True, bold=b), s) for s in (11, 17, 34, 84) for b in (False, True)]
    _wrap_bad = None
    for _ in range(600):
        _text = "".join(_rng.choice(_pool) for _ in range(_rng.randint(1, 50)))
        if not _has_cjk2(_text):
            continue
        _f, _s = _rng.choice(_fonts)
        _w = _rng.choice([1, 7, 37, 120, 496, 900])
        _t = _rng.choice([0, 0, 0.5, 1.5, 3.0])
        if _wrap(_text, _f, _w, _t) != _linear_cjk(_text, _f, _w, _t):
            _wrap_bad = (_text, _s, _w, _t)
            break
    check("ghost: CJK 折行二分拟合与逐字贪心同断点（600 组随机对拍）",
          _wrap_bad is None, str(_wrap_bad)[:160])

    # 导入契约（primitives 顶部注释承诺过的自检）：纯预览进程不得拖入
    # pptx / lxml / compiler。parse_token 曾经由 RGBColor 把整条 pptx 导入链
    # 拉进 PIL 预览路径（实测 ~150ms/轮），这里锁死：谁改回去，立刻失败。
    noleak_spec = tmp / "noleak_spec.json"
    noleak_spec.write_text(json.dumps({
        "canvas": {"width": 1280, "height": 720},
        "theme": {"colors": {"background": "#F1EFEA", "ink": "#12120F",
                             "muted": "#8A8984", "accent": "#6E6A63"}},
        "slides": [{"id": "s1", "background": "#F1EFEA", "elements": [
            {"id": "t", "type": "text", "x": 48, "y": 40, "width": 400, "height": 40,
             "text": "纯预览", "size": 20, "color": "ink"}]}]}, ensure_ascii=False),
        encoding="utf-8")
    noleak_out = tmp / "noleak_preview"
    probe = (
        "import sys, json;"
        f"sys.path.insert(0, {str(SCRIPTS)!r});"
        "from primitives import parse_token;"
        "c, _ = parse_token('#F1EFEA');"
        "assert str(c) == 'F1EFEA' and tuple(c) == (241, 239, 234), c;"
        "from ghost import ghost_deck;"
        f"ghost_deck(json.load(open({str(noleak_spec)!r}, encoding='utf-8')), "
        f"{str(noleak_out)!r}, pages=[1], scale=0.5);"
        "bad = sorted(m for m in ('pptx', 'pptx.dml.color', 'lxml', 'compiler') "
        "if m in sys.modules);"
        "print('LEAK:' + ','.join(bad) if bad else 'CLEAN')"
    )
    r_probe = subprocess.run([sys.executable, "-c", probe],
                             capture_output=True, text=True, encoding="utf-8")
    check("primitives: 纯预览进程零 pptx / lxml / compiler（导入契约成立）",
          r_probe.returncode == 0 and r_probe.stdout.strip() == "CLEAN",
          (r_probe.stdout + r_probe.stderr).strip()[:200])
    # 编译边界不变式：pptx 已加载的进程里 parse_token 必须产真 RGBColor
    # （`*.rgb =` 的 isinstance 校验只收真对象；等价元组只存在于零 pptx 进程）。
    from pptx.dml.color import RGBColor
    from primitives import parse_token as _parse_token
    c_loaded, _ = _parse_token("#123456")
    check("primitives: pptx 已加载时 parse_token 产真 RGBColor（编译边界类型不变）",
          isinstance(c_loaded, RGBColor), type(c_loaded).__name__)


# ── 4 · cli + docs ───────────────────────────────────────────────────
def test_cli(tmp: Path):
    brief = tmp / "brief.yml"
    import yaml
    brief.write_text(yaml.safe_dump(BRIEF, allow_unicode=True), encoding="utf-8")

    # Windows 默认 ANSI 代码页（管道重定向时 Python 用它写 stdout）不得让命令崩掉：
    # cp1252 下 `print("plan: 5 页…")` 曾经直接 UnicodeEncodeError、退出码 2，
    # 产物还没落盘就结束——这正是 CI Windows 作业全红的根因。
    import os as _os
    ascii_pipe = dict(_os.environ, PYTHONIOENCODING="cp1252")
    wide = tmp / "ascii-plan.json"
    r_ascii = run_vao("plan", str(brief), "--out", str(wide), env=ascii_pipe)
    check("cli: 非 UTF-8 控制台/管道下 plan 不崩且产物落盘（Windows 默认代码页）",
          r_ascii.returncode == 0 and wide.is_file() and "plan:" in r_ascii.stdout,
          (r_ascii.stderr or r_ascii.stdout)[-200:])
    r_strict = run_vao("plan", str(brief), "--out", str(tmp / "ascii-plan2.json"),
                       env=dict(_os.environ, PYTHONIOENCODING="ascii"))
    check("cli: 连 ASCII 管道也不崩（输出恒为 UTF-8）",
          r_strict.returncode == 0 and "视觉世界" in r_strict.stdout,
          (r_strict.stderr or "")[-200:])
    plan = tmp / "plan.json"
    skeleton = tmp / "build.py"
    assets_out = tmp / "assets.json"
    r = run_vao("plan", str(brief), "--out", str(plan), "--skeleton", str(skeleton),
                "--assets-out", str(assets_out), "--assets-dir", str(tmp / "gen"), "--json")
    check("cli: plan 退出码 0 且产物齐", r.returncode == 0 and plan.is_file()
          and skeleton.is_file() and assets_out.is_file(), r.stderr[-200:])
    plan_stdout = json.loads(r.stdout)
    timing = plan_stdout.get("performance") or {}
    import hashlib
    check("cli: plan 阶段计时可分解且 brief 读入/哈希凭证一致",
          timing.get("total_ms", 0) >= timing.get("intelligence_ms", 0)
          and "brief_load_ms" in timing and timing.get("reference_files_read") == 0
          and plan_stdout.get("workflow", {}).get("brief_file_sha256")
          == hashlib.sha256(brief.read_bytes()).hexdigest(),
          str(timing))
    payload = json.loads(plan.read_text(encoding="utf-8"))
    check("cli: plan.json 携带逐页判断卡（否决项有效、无竞争可空）",
          all(set(p["judgment"]["composition"]) == {"chosen", "why", "rejected"}
              and all(r["why_not"] for r in p["judgment"]["composition"]["rejected"])
              for p in payload["pages"]))
    check("cli: plan.json 无旧字段残留（direction/direction_execution/media）",
          "direction" not in payload["deck"] and "direction_execution" not in payload["deck"]
          and all("media" not in p for p in payload["pages"]))
    sk = skeleton.read_text(encoding="utf-8")
    check("cli: 骨架写「为什么」但不再复制否决（canonical owner 收口）",
          "为什么" in sk and "结论[" in sk and "否决" not in sk)

    # 一次 CLI plan：图是判断所得才准备资产；没有图不产生空清单/额外命令。
    media_brief = tmp / "media-brief.yml"
    media_brief.write_text(yaml.safe_dump(dict(BRIEF, slides=[
        {"id": "witness", "title": "门店回访揭示真实的使用障碍",
         "content": "门店现场与店长在高峰时段的记录（来源：走访纪要）",
         "asset": "required", "asset_subject": "店长在门店高峰时段真实记录"}]),
        allow_unicode=True), encoding="utf-8")
    media_dir = tmp / "media"
    media_plan = media_dir / "plan.json"
    media_sk = media_dir / "build.py"
    media_run = run_vao("plan", str(media_brief), "--out", str(media_plan),
                        "--skeleton", str(media_sk), "--json")
    image_json = json.loads(media_run.stdout) if media_run.returncode == 0 else {}
    auto_manifest = media_dir / "asset_manifest.json"
    noimg_plan = tmp / "noimg-plan.json"
    noimg_sk = tmp / "noimg-build.py"
    noimg_run = run_vao("plan", str(brief), "--out", str(noimg_plan),
                         "--skeleton", str(noimg_sk), "--json")
    noimg_json = json.loads(noimg_run.stdout) if noimg_run.returncode == 0 else {}
    check("cli: 图片被判断为必要时，同一次 plan 交付可定位的资产清单+构图骨架",
          media_run.returncode == 0 and auto_manifest.is_file()
          and image_json.get("workflow", {}).get("assets_manifest_path")
          == str(auto_manifest.resolve())
          and "--assets-manifest" in media_sk.read_text(encoding="utf-8")
          and len(json.loads(auto_manifest.read_text(encoding="utf-8")).get("assets", [])) == 1,
          str(image_json.get("workflow", {})))
    # --asset-cache（提示词缓存）此前在读取路径上引用了未导入的 json 模块，
    # 属于「只在冷启动加缓存时才炸」的真实崩溃点：这里把读写两个方向都钉住。
    cache_file = media_dir / "prompt-cache.json"
    cache_args = ("--assets-out", str(auto_manifest), "--asset-cache", str(cache_file))
    cache_first = run_vao("plan", str(media_brief), "--out", str(media_plan),
                          "--skeleton", str(media_sk), *cache_args, "--json")
    cache_second = run_vao("plan", str(media_brief), "--out", str(media_plan),
                           "--skeleton", str(media_sk), *cache_args, "--json")
    cache_meta = json.loads(auto_manifest.read_text(encoding="utf-8")).get("prompt_cache", {}) \
        if cache_first.returncode == 0 else {}
    check("cli: --asset-cache 读写闭环（第二次 plan 命中提示词缓存，不再 NameError）",
          cache_first.returncode == 0 and cache_second.returncode == 0
          and cache_file.is_file()
          and json.loads(cache_file.read_text(encoding="utf-8")).get("schema")
          == "vao-asset-prompt-cache-v3" and cache_meta.get("hits") == 1,
          f"rc={cache_second.returncode} hits={cache_meta.get('hits')} "
          f"{cache_second.stderr[-160:]}")

    check("cli: 无图一次 plan 没有资产清单和清单验证参数",
          noimg_run.returncode == 0 and not (tmp / "asset_manifest.json").exists()
          and not noimg_json.get("workflow", {}).get("assets_manifest_path")
          and "--assets-manifest" not in noimg_sk.read_text(encoding="utf-8"))

    invalid_brief = tmp / "invalid-media.yml"
    invalid = dict(BRIEF, slides=[dict(BRIEF["slides"][0],
                                   asset="required", asset_role="cosmic_unrecognized")])
    invalid_brief.write_text(yaml.safe_dump(invalid, allow_unicode=True), encoding="utf-8")
    invalid_dir = tmp / "invalid-media"
    invalid_run = run_vao("plan", str(invalid_brief), "--out", str(invalid_dir / "plan.json"),
                          "--skeleton", str(invalid_dir / "build.py"), "--json")
    check("cli: 不合法图片角色 fail-closed，不能留下声称已有清单的半成品 plan/骨架",
          invalid_run.returncode == 2 and "asset_role" in invalid_run.stderr
          and not (invalid_dir / "plan.json").exists()
          and not (invalid_dir / "build.py").exists()
          and not (invalid_dir / "asset_manifest.json").exists(), invalid_run.stderr[-200:])

    filled = make_build(tmp)
    out = tmp / "deck.pptx"
    stale_packet = out.with_suffix(".repair.json")
    stale_packet.write_text('{"status":"BLOCKED"}', encoding="utf-8")
    r = run_vao("check", str(filled), str(out), "--mode", "release", "--speed", "fast",
                "--json")
    check("cli: release PASS、产出预览/清单且不留多余/过期 repair 包",
          r.returncode == 0 and out.is_file()
          and out.with_suffix(".manifest.json").is_file() and not stale_packet.exists(),
          r.stdout[-200:])
    release_result = json.loads(r.stdout)
    release_timing = release_result.get("timing") or {}
    check("cli: 端到端计时包含预览且区分 compile/cache/QA",
          release_result.get("status") == "PASS"
          and release_timing.get("total_ms", 0) >= release_timing.get("ghost_ms", 0)
          and {"compile_ms", "cache_probe_ms", "qa_ms", "render_pages_drawn",
               "image_source_snapshots", "image_snapshot_bytes", "asset_qc_reused"}
          <= set(release_timing), str(release_timing))
    first = out.stat().st_mtime_ns

    r2 = run_vao("check", str(filled), str(out), "--mode", "release", "--speed", "fast")
    check("cli: 二次 release 缓存复用（未重编）",
          r2.returncode == 0 and "缓存复用" in r2.stdout and out.stat().st_mtime_ns == first,
          r2.stdout[-200:])
    unused = run_vao("check", str(filled), str(out), "--mode", "release",
                     "--assets-manifest", str(tmp / "nonexistent-asset.json"), "--json")
    unused_timing = (json.loads(unused.stdout).get("timing") or {}) if unused.returncode == 0 else {}
    check("cli: 无图 release 不读取无关清单，也不触发额外像素验证",
          unused.returncode == 0 and unused_timing.get("asset_bind_ms") == 0
          and unused_timing.get("qc_ms") == 0 and out.stat().st_mtime_ns == first,
          str(unused_timing))

    # 空分支不能测出伪绑定成本；模拟调度/时钟开销，避免依赖一次恰好舍入为 0。
    import contextlib
    import io
    import itertools
    import time
    import vao
    from unittest.mock import patch
    real_clock, ticks = time.perf_counter, itertools.count()
    with patch.object(vao.time, "perf_counter", side_effect=lambda: real_clock() + next(ticks) * .001), \
            contextlib.redirect_stdout(io.StringIO()):
        no_bind, no_bind_code = vao.run_check(str(filled), str(out), mode="release",
                                             assets_manifest=str(tmp / "nonexistent-asset.json"))
    check("cli: 调度开销不使未执行的资产绑定计时变成非零",
          no_bind_code == 0 and no_bind["performance"]["asset_bind_ms"] == 0
          and no_bind["performance"]["qc_ms"] == 0)

    preview_dir = tmp / "deck_preview"
    contact = preview_dir / "ghost-contact-sheet.png"
    contact.write_bytes(b"corrupt preview, not a PNG")
    repaired = run_vao("check", str(filled), str(out), "--mode", "release", "--speed", "fast",
                       "--json")
    repair_timing = (json.loads(repaired.stdout).get("timing") or {}) if repaired.returncode == 0 else {}
    from PIL import Image
    png_valid = False
    try:
        with Image.open(contact) as sample:
            sample.verify()
        png_valid = True
    except (OSError, ValueError):
        pass
    # 证据被破坏必须重贴，但页面没变就不该重画：逐页缓存命中 → 联络表重建 + 零重绘。
    # （此前断言还要求整份 bundle 不复用——那会把「什么都没变时复用页面」也判成失败。）
    check("cli: 交付 PNG 被破坏不能再凭 marker 假 PASS；复用未变页图修复联络表",
          repaired.returncode == 0 and png_valid
          and repair_timing.get("render_pages_drawn") == 0
          and repair_timing.get("compile_ms") == 0
          and repair_timing.get("ghost_ms", 0) > 0)
    small = run_vao("check", str(filled), str(out), "--mode", "release", "--speed", "fast",
                    "--ghost-pages", "2", "--json")
    small_timing = (json.loads(small.stdout).get("timing") or {}) if small.returncode == 0 else {}
    small_marker = json.loads((preview_dir / "ghost.meta.json").read_text(encoding="utf-8")) \
        if (preview_dir / "ghost.meta.json").is_file() else {}
    check("cli: 5→2 页取证不复用旧的 5 页凭证，不留过期预览图片",
          small.returncode == 0 and small_timing.get("render_pages") == 2
          and not small_timing.get("render_bundle_reused")
          and small_marker.get("request", {}).get("key_page_limit") == 2
          and len(list(preview_dir.glob("ghost-[0-9]*.png"))) == 2,
          str(small_timing))
    restored = run_vao("check", str(filled), str(out), "--mode", "release", "--speed", "fast",
                       "--ghost-pages", "5", "--json")
    restore_timing = (json.loads(restored.stdout).get("timing") or {}) if restored.returncode == 0 else {}
    check("cli: 2→5 页重新扩大证据范围，逐页缓存复用而非重画/重编",
          restored.returncode == 0 and restore_timing.get("render_pages") == 5
          and restore_timing.get("render_pages_drawn") == 0
          and restore_timing.get("compile_ms") == 0
          and len(list(preview_dir.glob("ghost-[0-9]*.png"))) == 5)

    broken = tmp / "broken.py"
    namespace: dict = {}
    exec(compile(filled.read_text(encoding="utf-8"), str(filled), "exec"), namespace)
    spec = json.loads(json.dumps(namespace["SPEC"], default=str))
    spec["slides"][0]["elements"][0]["height"] = 8
    broken.write_text(f"SPEC = {spec!r}\n", encoding="utf-8")
    r = run_vao("check", str(broken), str(tmp / "bad.pptx"), "--mode", "draft")
    packet = json.loads((tmp / "bad.repair.json").read_text(encoding="utf-8"))
    check("cli: BLOCK 退出码 2 + 分组修复包",
          r.returncode == 2 and packet["failure_codes"] == ["TEXT_OVERFLOW"]
          and packet["fix_plan"]["groups"], r.stdout[-200:])
    no_manifest = tmp / "image-build.py"
    with_pic = json.loads(json.dumps(spec))
    with_pic["slides"][0]["elements"].append(
        {"id": "missing_picture", "type": "image", "src": "asset:missing",
         "x": 580, "y": 240, "width": 540, "height": 350, "role": "witness"})
    no_manifest.write_text(f"SPEC = {with_pic!r}\n", encoding="utf-8")
    image_out = tmp / "missing-image.pptx"
    without = run_vao("check", str(no_manifest), str(image_out), "--mode", "release", "--json")
    blocked = json.loads(without.stdout) if without.returncode == 2 else {}
    check("cli: 含图但缺清单阻断于编译之前，不输出假的合格 PPTX",
          without.returncode == 2 and not image_out.exists()
          and (blocked.get("timing") or {}).get("compile_ms") == 0
          and "ASSET" in " ".join(blocked.get("failure_codes", [])),
          str(blocked.get("failure_codes", [])))

    empty = tmp / "empty.yml"
    empty.write_text("audience: 董事会\ndecision: 批准预算\nslides: []\n", encoding="utf-8")
    r = run_vao("plan", str(empty), "--out", str(tmp / "p.json"))
    check("cli: 空 slides 早失败并给修法", r.returncode != 0
          and ("slides" in r.stderr or "slides" in r.stdout), r.stderr[-200:])

    # CLI 面就是这套体系对外暴露的全部动作。多一个子命令 = 多一条 Agent 会调用、
    # 却没有文档与消费者支持的路径（`preview` 曾是这种残骸：SKILL 明说别跑它）。
    # 断言在这里锁住，新增入口必须是有意决定，而不是顺手。
    r_help = run_vao("--help")
    import re as _re
    _m = _re.search(r"\{([\w,]+)\}", r_help.stdout)
    listed = set(_m.group(1).split(",")) if _m else set()
    check("cli: 入口面恰好是 plan/check/qc（无第二执行路径）",
          listed == {"plan", "check", "qc"}, str(sorted(listed)))
    check("cli: 已删除的 preview 子命令确实不存在",
          run_vao("preview", "x.py").returncode != 0)


def test_delivery(tmp):
    from compiler import compile_deck
    from verify import check_spec
    from pptx import Presentation
    from pptx.enum.chart import XL_LEGEND_POSITION
    import intelligence as intel
    base = {"id": "chart", "type": "chart", "x": 64, "y": 100,
            "width": 1100, "height": 480, "label_size": 24,
            "categories": ["低负载", "高负载"],
            "series": [{"name": "路径甲", "values": [12, 30]},
                       {"name": "路径乙", "values": [16, 38]}]}
    deck = {"canvas": {"width": 1280, "height": 720},
            "theme": {"fonts": {"cn": "Noto Serif CJK SC", "latin": "Arial"},
                      "colors": {"background": "#FFFFFF", "ink": "#182B3A", "muted": "#536574",
                                 "primary": "#233649", "secondary": "#7B919E", "accent": "#BB5D26"}}, "slides": []}
    for kind in ("line", "column", "area", "bar"):
        for show in (False, True):
            deck["slides"].append({"id": f"{kind}-{show}", "elements": [dict(base, chart_kind=kind, show_values=show)]})
    path = tmp / "identity.pptx"
    report = compile_deck(deck, path)
    check("identity: 原生多序列编译无警告", report["passed"] and not report.get("warnings"))
    for slide, native in zip(deck["slides"], Presentation(path).slides):
        chart = native.shapes[0].chart
        check("identity: " + slide["id"], chart.has_legend
              and chart.legend.position == XL_LEGEND_POSITION.BOTTOM
              and chart.legend.include_in_layout is False
              and chart.legend.font.size.pt == 18
              and chart.legend.font._element.xpath("a:ea/@typeface") == ["Noto Serif CJK SC"]
              and [s.name for s in chart.series] == ["路径甲", "路径乙"]
              and [list(s.values) for s in chart.series] == [[12,30], [16,38]]
              and chart.plots[0].has_data_labels == slide["elements"][0]["show_values"])
    for kind in ("pie", "donut", "donut_composition"):
        for show in (False, True):
            e = {k:v for k,v in base.items() if k not in ("series", "categories")}
            e.update(chart_kind=kind, show_values=show, data=[{"label": f"类别{i}", "value": i+1} for i in range(7)])
            pie = dict(deck, slides=[{"id": "composition", "elements": [e]}])
            report = compile_deck(pie, path)
            c = Presentation(path).slides[0].shapes[0].chart
            labels = c.plots[0].data_labels if c.plots[0].has_data_labels else None
            check(f"identity: {kind}/{show} 隐藏值仍保留七个类别身份",
                  report["passed"] and not report.get("warnings") and labels is not None and labels.show_category_name
                  and labels.show_value == show and not c.has_legend
                  and len(c.series[0].values) == 7)
    def errors(element):
        report = check_spec(dict(deck, slides=[{"id": "probe", "elements": [element]}]))
        return [c for c in report["checks"] if c.get("level") == "error"]
    for names in (("", "乙"), ("甲", "甲")):
        e = dict(base, chart_kind="line", series=[{"name": n, "values": [1,2]} for n in names])
        check("identity: 拒绝缺失/歧义序列名 " + str(names), bool(errors(e)))
    check("identity: 单序列不强求冗余名称", not errors(dict(base, chart_kind="line", series=[{"values": [1,2]}])))
    check("identity: 类别名不可为空", bool(errors(dict(base, chart_kind="line", categories=["", "乙"]))))
    for kind, field, limit in (("process_flow","data",7), ("timeline","data",7), ("steps","data",6),
            ("matrix","points",12), ("architecture","layers",3), ("ranked_bar","data",8),
            ("progress_bar","data",6), ("stacked_bar","data",8), ("big_number_row","data",5), ("bubble","data",12)):
        e = {k:v for k,v in base.items() if k not in ("series", "categories")}
        e["chart_kind"] = kind
        # 数据行只带渲染器消费的键；x/y 只属于 matrix.points（未消费字段门会拦截数据行上的 x/y）。
        # matrix 的点必须是 points[{x,y,label}]：缺 x/y 时编译器只会抛裸 KeyError，
        # 所以 guard 的「非有限值」硬门也管它（下方另有专门负例）。
        items = ([{"label": f"节点{i}", "x": (i % 5) / 4, "y": (i % 7) / 6} for i in range(limit+1)]
                 if field == "points" else
                 [f"层{i}" if field == "layers" else {"label": f"节点{i}", "value": i+1} for i in range(limit+1)])
        e[field] = items[:limit]
        fits = not errors(e)
        e[field] = items
        check("payload: " + kind + " 边界完整/超量阻断", fits and bool(errors(e)))
    matrix_bad = {k: v for k, v in base.items() if k not in ("series", "categories")}
    matrix_bad.update(chart_kind="matrix", points=[{"label": "缺纵轴", "x": 0.2}])
    check("payload: matrix 缺 x/y 硬门阻断（裸 KeyError 不是可修的诊断）", bool(errors(matrix_bad)))
    from verify import release_manifest
    verification = release_manifest(deck, {})["verification"]
    check("evidence: 未打开/编辑/独立审阅不冒充验收",
          verification.get("target_powerpoint") == verification.get("native_editing") == "not_verified"
          and verification.get("independent_visual_review") == "not_performed")
    for content in ("甲组负责研发；乙组负责交付。", "必须先核验证据再发布。", "效率从 12% 提升到 18%。"):
        brief = {"slides": [{"id": str(i), "title": "核验依据", "content": content, "asset": "none"} for i in range(6)]}
        bundle = intel.think(brief)
        check("rhythm: 重复不能覆盖关系/观点/数据判断 " + content,
              all(p["judgment"]["composition"] == intel.page_judgment(brief["slides"][i], index=i, total=6,
                  brief=brief, world=bundle["deck"]["world"], warnings=[])["composition"] for i,p in enumerate(bundle["pages"])))


def test_plan_boundaries(tmp):
    import intelligence as intel
    import vao, assets
    from unittest.mock import patch
    from PIL import Image
    import contextlib, io
    base = {"audience": "评审", "decision": "发布", "tension": "核验证据",
            "slides": [{"id": "page", "title": "保留原文", "content": "信息不可损失", "asset": "none"}]}
    for field in ("title", "claim", "id", "audience", "tension", "eyebrow"):
        b = json.loads(json.dumps(base))
        value = '先验证\n再交付 """ # C:\\notes\r限定条件'
        b["slides"][0][field] = value
        sk = intel.build_skeleton(intel.think(b)); scope = {}
        exec(compile(sk, field, "exec"), scope)
        check("plan: 多行/引号安全生成 " + field, len(scope["SPEC"]["slides"]) == 1)
    b = dict(base, audience='讨论 """ 与路径 C:\\new')
    check("plan: deck 说明不能终止 Python 字符串", compile(intel.build_skeleton(intel.think(b)), "deck", "exec") is not None)
    brief = tmp / "brief.json"; brief.write_text(json.dumps(base), encoding="utf-8")
    plan = tmp / "plan.json"; work = tmp / "author.py"
    old = "SPEC = {'slides': [{'elements': [{'type': 'text', 'text': '作者已完成'}]}]}\n"
    work.write_text(old, encoding="utf-8")
    with contextlib.redirect_stdout(io.StringIO()): result = vao.run_plan(str(brief), str(plan), str(work))
    side = Path(result["workflow"]["skeleton_path"])
    check("plan: 不执行/不覆盖任何现有作者文件", work.read_text(encoding="utf-8") == old and side != work and side.exists())
    stamp = side.stat().st_mtime_ns
    with contextlib.redirect_stdout(io.StringIO()): again = vao.run_plan(str(brief), str(plan), str(work))
    check("plan: 同输入副本不重写、不产生更多文件", again["workflow"]["skeleton_path"] == str(side.resolve()) and side.stat().st_mtime_ns == stamp)
    side.write_text("# 第二份作者工作", encoding="utf-8")
    blocked = False
    try: vao.run_plan(str(brief), str(plan), str(work))
    except FileExistsError: blocked = True
    check("plan: 副本也被保护，不按代码格式猜是否已填稿", blocked and side.read_text(encoding="utf-8") == "# 第二份作者工作")
    cli = run_vao("plan", str(brief), "--out", str(plan), "--json")
    compact = json.loads(cli.stdout)
    check("plan: JSON 仅回摘要，完整判断仍落盘", cli.returncode == 0 and "pages" not in compact
          and compact["page_count"] == 1 and "pages" in json.loads(plan.read_text(encoding="utf-8"))
          and len(cli.stdout.encode()) < 4096)

    image = tmp / "original.png"; Image.new("RGB", (1024,768), (214,206,192)).save(image)
    entry = {"asset_id": "named", "decision": "generate", "expected_filename": image.name,
             "prompt": "test", "negative": ["test"], "slide_ids": ["s1"], "ratio": "16:9",
             "allow_crop": True, "meta": {"text_color": "#111111"},
             "asset_role": "background", "asset_role_source": "author"}
    manifest = tmp / "assets.json"
    def save(entry):
        manifest.write_text(json.dumps({"assets": [entry]}), encoding="utf-8")
    save(entry); original = manifest.read_bytes()
    with patch.object(assets, "image_qc", wraps=assets.image_qc) as measure:
        offline = vao.run_qc(str(manifest), speed="strict")
    kw = measure.call_args.kwargs
    check("qc: 同一解析器消费文件名/HEX 字色/裁切声明", Path(measure.call_args.args[0]) == image.resolve()
          and kw["text_is_dark"] is True and kw["allow_crop"] is True)
    report, _ = assets.qc_report(manifest, phase="draft", speed="strict", persist=False)
    check("qc: 离线与生产同判且不写报告/重试状态", offline["results"][0]["action"] == report["results"][0]["policy"]["action"]
          and manifest.read_bytes() == original and not manifest.with_suffix(".qc.json").exists())
    save(dict(entry, decision="existing", origin={"kind":"original","path":str(image),"source":"test fixture"}))
    check("qc: 既有素材按 origin.path 查找，不误报缺失", vao.run_qc(str(manifest), speed="strict")["results"][0]["status"] != "missing")
    assets.qc_report(manifest, phase="draft", speed="strict")
    qc_file = manifest.with_suffix(".qc.json"); prior = qc_file.read_bytes()
    with patch.object(assets, "image_qc", wraps=assets.image_qc) as measure:
        vao.run_qc(str(manifest), speed="strict")
    check("qc: 有效报告复用测量且不改写凭证", not measure.called and qc_file.read_bytes() == prior)
    import os
    stat = image.stat(); blob = bytearray(image.read_bytes()); blob[-1] ^= 1; image.write_bytes(blob)
    os.utime(image, ns=(stat.st_atime_ns, stat.st_mtime_ns))
    with patch.object(assets, "image_qc", wraps=assets.image_qc) as measure:
        vao.run_qc(str(manifest), speed="strict")
    check("qc: strict 同 size/mtime 换字节不复用旧测量", measure.called and qc_file.read_bytes() == prior)
    save(dict(entry, expected_filename="missing.png"))
    missing = run_vao("qc", "--assets-manifest", str(manifest), "--speed", "strict")
    check("qc: 缺素材退出码非零，不能假成功", missing.returncode == 2 and json.loads(missing.stdout)["status"] == "BLOCKED")
    save(dict(entry, expected_filename="ambiguous.webp"))
    image.with_name("ambiguous.png").write_bytes(image.read_bytes())
    image.with_name("ambiguous.jpg").write_bytes(image.read_bytes())
    ambiguous = run_vao("qc", "--assets-manifest", str(manifest))
    check("qc: 多候选不能任取第一个", ambiguous.returncode == 2 and "多个候选" in ambiguous.stderr)


def test_release_evidence(tmp: Path):
    """Evidence is ordered, complete and fail-closed; arrows retain direction."""
    import copy
    import contextlib
    import io
    import vao
    from verify import check_spec, release_manifest, preview_issues
    from compiler import compile_deck
    from pptx import Presentation
    from ghost import ghost_page
    from PIL import ImageChops

    spec = vao.load_spec(make_build(tmp, chart=False))[0]
    spec['slides'][0]['id'] = 'z-open'
    spec['slides'][1]['id'] = 'a-evidence'
    build = tmp / 'ordered.json'
    build.write_text(json.dumps(spec), encoding='utf-8')
    output = tmp / 'ordered.pptx'
    with contextlib.redirect_stdout(io.StringIO()):
        result, code = vao.run_check(str(build), str(output), mode='release', speed='strict')
    check('release: 非字典序叙事 ID 严档一次收口', code == 0 and result['release_eligible'])
    renamed = copy.deepcopy(spec)
    for s in renamed['slides']:
        s['id'] = 'renamed-' + s['id']
    build.write_text(json.dumps(renamed), encoding='utf-8')
    with contextlib.redirect_stdout(io.StringIO()):
        renamed_result, renamed_code = vao.run_check(str(build), str(output), mode='release', speed='strict')
    check('release: 相同 PPTX 字节但 ID 改名不复用旧页绑定', renamed_code == 0
          and renamed_result['compile']['output_sha256'] == result['compile']['output_sha256'])
    build.write_text(json.dumps(spec), encoding='utf-8')
    with contextlib.redirect_stdout(io.StringIO()):
        result, code = vao.run_check(str(build), str(output), mode='release', speed='strict')
    mf = json.loads(output.with_suffix('.manifest.json').read_text(encoding='utf-8'))
    ghost = mf['ghost_preview']
    check('release: 持久清单保留逐页证据而非只留目录', ghost['slide_ids'] == [s['id'] for s in spec['slides']] and len(ghost['pages']) == 5)
    for name, mutate in (
        ('缺 spec 戳', lambda q: q.pop('source_spec_hash')),
        ('缺 PPTX 双凭证', lambda q: q.update(compile={'passed': True})),
        ('缺 PPTX 摘要', lambda q: q['compile'].pop('output_sha256')),
    ):
        q = copy.deepcopy(result); mutate(q)
        m = release_manifest(spec, q, ghost_preview=ghost)
        check('release: ' + name + ' 必须 BLOCK', m['status'] == 'BLOCKED' and not m['release_eligible'])
    check('release: 同页 ID 的旧预览不能冒充当前产物', release_manifest(
        spec, result, ghost_preview={**ghost, 'output_sha256':'0'*64})['status'] == 'BLOCKED')
    m = release_manifest(spec, result)
    check('release: 缺预览必须 BLOCK', m['status'] == 'BLOCKED')
    ids = ghost['slide_ids']
    for name, mutation in (
        ('重复页面', {'slide_ids': [ids[0]]*5}),
        ('重复文件', {'pages': [ghost['pages'][0]]*5}),
        ('错误顺序', {'slide_ids': ids[::-1]}),
        ('未知范围', {'scope': 'invented'}),
    ):
        check('release: ' + name + ' 凭证拒收', bool(preview_issues({**ghost, **mutation}, ids)))
    subset = dict(ghost, scope='key_pages', count=1, slide_ids=[ids[0]],
                  pages=ghost['pages'][:1], key_pages={ids[0]:'cover'})
    check('release: fast 合法抽样保留接口', not preview_issues(subset, ids))
    check('release: strict 不接受抽样', release_manifest(spec, result, ghost_preview=subset)['status'] == 'BLOCKED')
    with contextlib.redirect_stdout(io.StringIO()):
        expired, code = vao.run_check(str(build), str(output), mode='release', speed='strict', deadline=0.001)
    check('release: deadline 不升级为可发布', code == 2 and not expired['release_eligible'])
    # Restore a valid manifest then fail BEFORE parsing; stale PASS must not survive.
    with contextlib.redirect_stdout(io.StringIO()):
        vao.run_check(str(build), str(output), mode='release', speed='strict')
        build.write_text('{broken', encoding='utf-8')
        failed, code = vao.run_check(str(build), str(output), mode='release', speed='strict')
    stale = json.loads(output.with_suffix('.manifest.json').read_text(encoding='utf-8'))
    check('release: 输入异常覆盖旧 PASS 清单', code == 2 and stale['status'] == 'BLOCKED' and not stale['release_eligible'])
    check('release: 目标环境状态不被自动升级', mf['verification']['target_powerpoint'] == 'not_verified' and mf['verification']['native_editing'] == 'not_verified')
    dup = copy.deepcopy(spec); dup['slides'][1]['id'] = dup['slides'][0]['id']
    check('identity: 重复页 ID 在编译前 BLOCK', not check_spec(dup)['passed'])
    dup = copy.deepcopy(spec); dup['slides'][0]['elements'][1]['id'] = dup['slides'][0]['elements'][0]['id']
    check('identity: 重复元素 ID 在编译前 BLOCK', not check_spec(dup)['passed'])

    elements = [dict(id=name, type='shape', shape='arrow', x=100, y=100,
                     width=w, height=h, stroke='#A34225', stroke_width=4)
                for name,w,h in [('horizontal',160,0),('vertical',0,160),('diagonal',160,160)]]
    elements += [dict(elements[0],id='ordinary',shape='line',y=320),
                 dict(elements[0],id='invisible',y=360,stroke={'type':'none'})]
    arrows = dict(spec, slides=[dict(id='arrows',elements=elements)])
    path = tmp/'arrows.pptx'; report=compile_deck(arrows,path)
    shapes=Presentation(path).slides[0].shapes
    check('arrow: 三方向箭头是原生可编辑连接线', report['passed'] and all(len(shapes[i]._element.xpath('.//a:tailEnd[@type="triangle"]')) == 1 for i in range(3)))
    check('arrow: 普通线不误加箭头', not shapes[3]._element.xpath('.//a:tailEnd'))
    check('arrow: 显式无描边保持不可见', bool(shapes[4]._element.xpath('.//a:ln/a:noFill')))
    for element in elements[:3]:
        slide=dict(id='arrow-preview',elements=[element])
        actual=ghost_page(slide,arrows)
        plain=ghost_page(dict(slide,elements=[dict(element,shape='line')]),arrows)
        check('arrow: ghost 方向端点 '+element['id'], ImageChops.difference(actual.convert('RGB'),plain.convert('RGB')).getbbox() is not None)


def test_boundary_typography(tmp: Path):
    """中文混排的行尾纪律：数字与单位不许被断到两行。

    判据来自既往一次真实版面观察：「平均配送半径是 186」留在行尾、「公里」推到下一行——
    数字与单位拆开，正是 SKILL.md §3 禁止的那件事。
    混排间隙（U+2009 细空格）本身是合法断行机会，所以数字与单位之间必须用
    不断行空格；这条纪律由 primitives 的单位词表**一处**决定，编译器与幽灵预览共用。
    """
    import primitives as prim

    gapped = prim.insert_script_gaps("配送半径是186公里")
    check("混排: 数字与单位之间用不断行空格（细空格是断行机会）",
          prim.NO_BREAK_SPACE in gapped and "186" + prim.NO_BREAK_SPACE + "公里" in gapped,
          repr(gapped))
    check("混排: 作者自己写的空格也被绑定（「186 公里」同样不许断）",
          prim.insert_script_gaps("平均半径是 186 公里")
          == "平均半径是 186" + prim.NO_BREAK_SPACE + "公里",
          repr(prim.insert_script_gaps("平均半径是 186 公里")))
    check("混排: 非单位的普通空格不受影响（不制造粘在一起的假词）",
          prim.insert_script_gaps("覆盖 华东 与 华南") == "覆盖 华东 与 华南"
          and prim.insert_script_gaps("我们 做了 三件事") == "我们 做了 三件事")
    check("混排: 间隙插入幂等（二次处理不再漂移）",
          prim.insert_script_gaps(gapped) == gapped)
    # 容量估算必须与产物实际写的字符一致，否则 Guard 会算错行数。
    check("混排: 容量估算认识不断行空格（不被当成一个汉字）",
          0.5 < prim.text_units(prim.NO_BREAK_SPACE) < 0.6
          and prim.text_units(prim.NO_BREAK_SPACE) == prim.text_units(" "),
          str(prim.text_units(prim.NO_BREAK_SPACE)))
    # 单位词表只有一处：预览的禁则判据必须与编译器写出的间隙同源。
    import ghost
    for probe in ("公里", "个月", "万元", "%", "家"):
        check(f"混排: 预览禁则与编译共用同一份单位表（{probe}）",
              bool(ghost._UNIT_HEAD.match(probe)) and probe in prim.UNIT_WORDS)
    # 真实产物里也必须是不换行：编译后从 PPTX 里读回文本核对。
    from compiler import compile_deck
    import vao
    spec, _ = vao.load_spec(make_build(tmp, chart=False))
    spec["slides"][0]["elements"].append(
        {"id": "nb", "type": "text", "x": 40, "y": 600, "width": 800, "height": 40,
         "role": "body", "text": "平均配送半径是 186 公里", "size": 16})
    out = tmp / "nb.pptx"
    compile_deck(spec, str(out))
    from pptx import Presentation
    text = "\n".join(shape.text_frame.text for shape in
                     Presentation(str(out)).slides[0].shapes if shape.has_text_frame)
    check("混排: 交付的 PPTX 里数字与单位之间是不换行空格",
          "186" + prim.NO_BREAK_SPACE + "公里" in text, repr(text[-80:]))


def test_design_judgment(tmp: Path):
    """设计判断回归：每条用例对应一次真实成品/渲染审计暴露出来的错判。

    这些不是「更多规则」，而是把已经修对的判断钉住——每个断言的期望值都来自
    一次可复现的误判现场（接受基准 brief_v1 的 s01/s02/s04/s05/s08 与条件从句页）。
    """
    import intelligence as intel

    # 1 · 时间坐标不进证据量：封面写着「投委会 2026.10」，四位年 + 点分月日是日期。
    #     旧实现把它读成 2026.1，封面于是「有一个量要立住」，焦点/权重/构图/字阶
    #     一起落到日期上（实测：封面最大的元素成了 2026.10）。
    check("numerals: 点分日期（2026.10 / 2026.10.14）不是证据量",
          intel.evidence_numerals("华东区 2024—2026 经营复盘 · 投委会 2026.10") == []
          and intel.evidence_numerals("数据截至 2026.10.14") == []
          and [n["raw"] for n in intel.evidence_numerals("单价 12.5 元")] == ["12.5 元"],
          str(intel.evidence_numerals("华东区 2024—2026 经营复盘 · 投委会 2026.10")))
    cover = intel.page_judgment(
        {"id": "c1", "title": "扩张的答案不在开店速度，在单店质量",
         "content": "华东区 2024—2026 经营复盘 · 投委会 2026.10",
         "insight": "扩张的答案不在开店速度，在单店质量", "asset": "none"},
        index=0, total=8, brief={"audience": "投委会", "decision": "批准扩张预算"},
        world=intel.derive_world({"audience": "投委会"}, [intel.understand("x")]),
        warnings=[])
    check("判断: 封面是陈述位，日期不当焦点（焦点回到主张本人）",
          cover["visual_role"]["role"] == "establish"
          and cover["focus"]["element_role"] == "statement_text"
          and cover["understanding"]["evidence"] == "none",
          str({k: cover[k] for k in ("visual_role", "focus")}))

    # 2 · 主张说的是哪个指标，就锚哪个数字：标题讲增速连续收窄，正文末句讲的是
    #     单店月均营收——旧实现让正文的「数字＋变化」把标题结论顶掉，主语变成
    #     另一个指标的金额（实测：增速页的头号数字成了 37.5 万元）。
    mixed = intel.extract_claim(
        {"title": "同店销售额仍在增长，但增速已经连续三个季度收窄",
         "content": "同店销售额同比增速：2025Q4 为 8.6%，2026Q1 为 6.1%，2026Q2 为 3.4%，"
                    "2026Q3 为 2.2%。单店月均营收从 41.2 万元降至 37.5 万元。"})
    check("claim: 正文换了指标时不顶掉标题结论（同一度量才谈得上更强）",
          mixed["text"] == "同店销售额仍在增长，但增速已经连续三个季度收窄", str(mixed))
    stronger = intel.extract_claim(
        {"title": "华东区数据",
         "content": "结论：增速收窄的根因是店长成熟度。单店月均营收降至 37.5 万元。"})
    check("claim: 带结论措辞的正文句仍然胜出（不被指标一致性挡住）",
          "根因" in (stronger["text"] or ""), str(stronger))

    # 3 · 序列位置坐标不是第二个量：「第 3 个月 × 月均营收（万元）」是坐标轴，
    #     不是两种单位并置。旧实现判成单位不同 → 丢掉 data_field，还多报一条
    #     并不存在的口径告警（实测：同一张图的同口径数据被判不可比）。
    curve = intel.understand(
        "成熟店与新店的营收差距在开业第 9 个月才收敛 "
        "开业后各月单店月均营收（万元）：第 3 个月，成熟店 41.2，新店 27.2；"
        "第 6 个月，成熟店 41.8，新店 33.6。")
    check("data: 序列位置坐标（第 N 个月）不把同口径序列判成单位不同",
          curve["same_unit"] is True and curve["evidence"] == "comparison", str(curve))
    check("data: 真跨量纲的页面仍然判不可比（尺子没被放松）",
          intel.understand("总收入 480 万元，人均成本 90 元")["same_unit"] is False)

    # 4 · 请求的动作自己带的量赢过后面的终值：「批准本轮 30 家，同步把培训基地
    #     从 2 个扩到 3 个」的主语是要批的 30 家，不是「扩到 3 个」。
    ask = "批准本轮 30 家，同步把培训基地从 2 个扩到 3 个"
    _, ask_primary = intel.number_roles(ask, ask)
    check("focus: 要决定的那个量赢过后面出现的终值（30 家，而不是 3 个）",
          ask_primary and ask_primary["raw"] == "30 家", str(ask_primary))

    # 5 · 序列没有点名取值时，同分同单位取**末值**：讲「增速连续三个季度收窄」，
    #     主语是当前的 2.2%，不是三年前的 8.6%（旧实现取同分里的第一个）。
    _, latest = intel.number_roles(
        "同店销售额同比增速：2025Q4 为 8.6%，2026Q1 为 6.1%，2026Q2 为 3.4%，2026Q3 为 2.2%。",
        "增速已经连续三个季度收窄")
    check("focus: 序列同分同单位时取末值（当前状态，不是基期）",
          latest and latest["raw"] == "2.2%", str(latest))

    # 6 · 主张自己的量不算「重复」：结论点明落点、证据里必然再出现一次。
    #     旧实现把它列进删除清单，等于叫作者删掉本页的主张（实测：收敛页的
    #     删除清单里出现「9 个月 的重复出现」）。
    gap_text = ("开业后各月：第 3 个月，成熟店 41.2，新店 27.2；"
                "第 9 个月，成熟店 42.1，新店 38.9。")
    gap_page = intel.information_weight(
        intel.understand(gap_text),
        {"text": "成熟店与新店的营收差距在开业第 9 个月才收敛", "source": "extracted"},
        gap_text)
    check("weight: 结论点名的坐标不进删除清单（主张与证据的必要呼应）",
          not any("重复" in x for x in gap_page["delete"]), str(gap_page["delete"]))
    dup_text = "收入 12 亿，收入 12 亿（来源：财报）"
    dup_again = intel.information_weight(intel.understand(dup_text),
                                        {"text": "收入 12 亿", "source": "extracted"}, dup_text)
    check("weight: 同一句话里的数字重复照旧进删除清单",
          any("重复" in x for x in dup_again["delete"]), str(dup_again["delete"]))

    # 7 · 「单位不同」只在**主张**层面成立：页面里同时出现家数、月数与百分比，
    #     不等于这一页要把它们画到同一条数量轴上（实测：人效对比页的未决项来自
    #     正文的门店家数，主张其实只有 34% 一个量）。
    chart_brief = {"audience": "投委会", "decision": "批准扩张"}
    chart_page = intel.think({**chart_brief, "slides": [{
        "id": "u1", "title": "开店速度已经超过组织承接能力",
        "content": "2026 年新开 62 家门店，同比 38 家增长 63%；同期店长培养周期从 5.2 个月"
                   "拉长到 8.7 个月。新店首月人效比成熟店低 34%。"}]})["pages"][0]["judgment"]
    check("judgment: 主张单一量时不误报口径未决（正文混量纲不算）",
          not any("单位不同" in q for q in (chart_page["open_questions"] or [])),
          str(chart_page["open_questions"]))
    mixed_claim = intel.think({**chart_brief, "slides": [{
        "id": "u2", "title": "总收入与人均成本对比：480 万元 vs 90 元",
        "content": "总收入 480 万元，人均成本 90 元（来源：财务表）"}]})["pages"][0]["judgment"]
    check("judgment: 主张自己跨量纲时照旧报口径未决（尺子没被放松）",
          any("单位不同" in q for q in (mixed_claim["open_questions"] or [])),
          str(mixed_claim["open_questions"]))

    # 8 · 单个「先」是优先级不是流程：「先把周期压回去，再谈扩张」是决策页，
    #     不是一个时间轴（旧实现据一个虚词判出 sequence + linear_structure）。
    decision = intel.think({**chart_brief, "slides": [{
        "id": "d1", "title": "先把店长培养周期压回 6 个月，再谈第三轮扩张",
        "content": "批准本轮 30 家，同步把培训基地从 2 个扩到 3 个；"
                   "若 2027Q1 店长培养周期仍高于 7 个月，暂停后续批次。"}]})["pages"][0]["judgment"]
    check("judgment: 单个「先/再」不构成流程（决策页回到收束位）",
          decision["visual_role"]["role"] == "summarize"
          and decision["composition"]["chosen"] != "linear_structure",
          str({"role": decision["visual_role"]["role"],
               "comp": decision["composition"]["chosen"]}))
    check("judgment: 条件从句里的比较词不把收尾页判成两端对比",
          decision["understanding"]["evidence"] != "comparison", str(decision["understanding"]))

    # 9 · 眉标是被截断的从句时不输出：「答案不在开店速度」曾是眉标「答案不」。
    for text, forbid in (("扩张的答案不在开店速度，在单店质量", "答案不"),
                         ("同店销售额仍在增长，但增速已经连续三个季度收窄", "同店销售额仍")):
        got = intel.section_label({}, {"text": text}, "prove")
        check(f"眉标: 不以副词收尾的截断从句（{forbid}）", got != forbid and len(got) <= 6, got)
    check("眉标: 干净名词短语照旧派得出来（收尾字表不被过度收紧）",
          intel.section_label({}, {"text": "供应链的响应半径决定了开业节奏的上限"}, "prove")
          == "供应链", intel.section_label({}, {"text": "x"}, "prove"))

    # 10 · 单位表覆盖实务量词：配送半径 186 公里不能读成无单位的 186。
    check("numerals: 公里/吨/平方米等单位进入读数（主语不丢量纲）",
          [(n["raw"], n["unit"]) for n in intel.evidence_numerals("配送半径 186 公里")]
          == [("186 公里", "公里")]
          and intel.evidence_numerals("产量 12 万吨")[0]["unit"] == "万吨")

    # 11 · 主张讲「差」时，排版纪律点名差本身：两条近重合的线各带一套数值标签，
    #      会在同一个 x 上挤到一起且认不出归属（真实渲染实测）。
    gap_judgment = intel.page_judgment(
        {"id": "g1", "title": "成熟店与新店的营收差距在开业第 9 个月才收敛",
         "content": "开业后各月单店月均营收（万元）：第 3 个月，成熟店 41.2，新店 27.2；"
                    "第 6 个月，成熟店 41.8，新店 33.6；第 9 个月，成熟店 42.1，新店 38.9。",
         "asset": "none"},
        index=4, total=8, brief=chart_brief,
        world=intel.derive_world(chart_brief, [intel.understand("x")]), warnings=[])
    check("typography: 讲差的一页要求画差值或只标端点（不给两条线各配标签）",
          "差值" in gap_judgment["typography"]["discipline"],
          gap_judgment["typography"]["discipline"])

    # 12 · 趋势词不是序列：序列至少两个数据点。封面「接管增长」与致谢页「内容增长负责人」
    #      曾因一个措辞被判成数据序列，连带封面变成证明页并派发「坐标轴」指令。
    check("series: 无数字的趋势措辞不构成序列",
          intel.understand("自有内容正在接管增长，预算应该跟上")["evidence"] != "series")
    check("series: 两个以上数据点仍是序列",
          intel.understand("2023 年 38%，2024 年 47%，2025 年 61%")["evidence"] == "series")
    deck_plan = intel.think({
        "audience": "董事会", "decision": "批准预算", "tension": "成本持续上升",
        "slides": [
            {"id": "s01", "title": "自有内容正在接管增长，预算应该跟上", "content": "2027 内容战略提案"},
            {"id": "s02", "title": "外部流量价格三年翻倍", "content": "获客成本从 186 元涨到 412 元（来源：投放台账）"},
            {"id": "s03", "title": "谢谢，欢迎质询", "content": "联系人：内容增长负责人"}]})
    jd = {p["id"]: p["judgment"] for p in deck_plan["pages"]}
    check("judgment: 封面建立主张，不被判成证明页",
          jd["s01"]["visual_role"]["role"] == "establish", jd["s01"]["visual_role"]["role"])
    check("judgment: 无图表的封面不派发「坐标轴与刻度」",
          not any("坐标轴" in w for w in jd["s01"]["information_weight"]["weaken"]))
    check("judgment: 致谢/问答页没有结论时不追问补句（也不编造结论）",
          jd["s03"]["claim"]["source"] == "frame" and jd["s03"]["claim"]["text"] is None
          and not any("先写出一句" in q for q in (jd["s03"].get("open_questions") or []))
          and "先补一句" not in str(jd["s03"]["claim"].get("why")))
    check("judgment: 一句话收束的致谢页让留白承担分量",
          jd["s03"]["composition"]["chosen"] == "big_whitespace", jd["s03"]["composition"]["chosen"])
    # 13 · 时间序列不是比较：按年份排列的同一指标（有图表）仍是 series；分期不是两个并置选项
    year_series = intel.understand("2023 年 38%，2024 年 47%，2025 年 61%", has_chart=True)
    check("series: 按年份排列的图表数据是时间序列，不被当成比较",
          year_series["evidence"] == "series" and not year_series["comparison"],
          year_series["evidence"])
    phased = intel.understand("一期 1,200 万（2027Q1）；二期 1,800 万（2027Q3）；三期 1,800 万（2028Q1）。"
                              "预计 2028Q2 累计自有流量收入超过投入")
    check("series: 分期投入是先后承诺（sequence），不是「超过投入」式比较",
          phased["evidence"] == "sequence" and not phased["comparison"], phased["evidence"])
    skeleton_text = intel.build_skeleton(deck_plan)
    check("skeleton: 通用排版纪律只在头部声明一次，逐页不重复",
          skeleton_text.count(intel.BASE_DISCIPLINE) == 1, str(skeleton_text.count(intel.BASE_DISCIPLINE)))
    check("skeleton: 跨页锚的规则只在头部说明一次，逐页不重复",
          skeleton_text.count("跨页锚（anchor") == 1,
          str(skeleton_text.count("跨页锚（anchor")))
    # 14 · 断言谓语结尾的标题是结论，不是话题标签（探针：「差距明显」被判 absent，并要求补句）
    check("claim: 以断言谓语结尾的标题（差距明显）是结论，不追问补句",
          intel._title_assertive("三个市场的进入成本与回本周期差距明显")
          and not intel._title_assertive("华东区数据"),
          str(intel._title_assertive("三个市场的进入成本与回本周期差距明显")))
    # 15 · 「对比」是页面主题标签（封面「对比评估」），不把封面判成比较页；「N 倍」才是比较
    check("compare: 「对比评估」作为主题标签不触发比较",
          not intel.understand("对比评估")["comparison"])
    check("compare: 「N 倍」是比较断言",
          intel.understand("高端线售价是泰国的 1.6 倍")["comparison"])
    # 16 · 规格列举不是序列；同单位也需要明确时间/变化轴
    spec = intel.understand("规格：3 次烧成，1280℃，0.6–0.8 毫米")
    check("series: 不同单位的规格列举不被判为序列",
          spec["evidence"] != "series", spec["evidence"])
    # 17 · 群体/角色名词不是可拍的证词主体（「伙伴」「团队」「客户」不触发配图）
    partner = intel.understand("选择越南，需要先补齐两个渠道伙伴")
    check("media: 「伙伴」这类群体名词不构成证词，不强制配图",
          intel.media_necessity(partner, "persuade")["decision"] != "required",
          str(intel.media_necessity(partner, "persuade")))
    # 18 · 首页元信息行是标签，不参与角色判定（探针：封面「流程评审」被判成流程页）
    meta_cover = intel.page_judgment(
        {"id": "m1", "title": "发布周期能从 6 周压到 4 周，前提是测试前移", "content": "2026 Q4 发布流程评审"},
        index=0, total=6, brief={"audience": "评审会", "decision": "同意"},
        world=intel.derive_world({"audience": "评审会"}, [intel.understand("x")]), warnings=[])
    # 本断言的职责是「元信息行不参与形状判断」：它不得让封面变成流程页。
    # 角色本身由内容先决（「从 6 周压到 4 周」是量化主张 → 证明页），与上一条数据页
    # 规则同源；此前把断言绑在 role == establish 上，是因为端点变化识别不全、
    # 该页证据被判成 none（10 月 14 日修正）。判据改为形状：证据不是 sequence、
    # 构图不是线性结构、焦点不是流程/结构位、结论仍是作者写下的那句主张。
    check("meta: 封面元信息行（带日期/分隔符）不把封面判成流程页",
          meta_cover["understanding"]["evidence"] != "sequence"
          and meta_cover["composition"]["chosen"] != "linear_structure"
          and meta_cover["focus"]["element_role"] not in ("structure_map", "process_flow",
                                                          "timeline")
          and str(meta_cover["claim"].get("text") or "").startswith("发布周期"),
          str({"ev": meta_cover["understanding"]["evidence"],
               "comp": meta_cover["composition"]["chosen"],
               "focus": meta_cover["focus"]["element_role"],
               "claim": meta_cover["claim"].get("text")}))
    check("meta: 首页若是一句证词（无日期/分隔符）不被当成元信息丢掉",
          intel.judgment_view({"title": "走访", "content": "华东 6 家门店走访记录"}, 0)["content"]
          == "华东 6 家门店走访记录")
    # 19 · 标题即论断：无正文时，断言谓语/对比句式/「从 A 增至 B」仍是结论；名词标签不是
    check("claim: 无正文的论断标题（不在…在…）是结论位",
          intel.extract_claim({"title": "扩张的答案不在开店速度，在单店质量"})["source"] == "extracted")
    check("claim: 无正文的名词标签（18 个月落地节奏）不编结论",
          intel.extract_claim({"title": "18 个月落地节奏"})["source"] == "absent")
    # 20 · 场所词：有可核对数字时它是计数对象；只有证人动作让场所成为现场证词
    check("media: 「门店 120 家」这类计数句不是现场（数据页不配生活图）",
          intel.media_necessity(intel.understand("门店规模稳定，单店客单价是 86 元",
                                                 content="门店 120 家，营收 3.2 亿元"), "prove")["decision"] == "none")
    walk = intel.media_necessity(intel.understand("华东 6 家门店走访记录"), "prove")
    check("media: 有证人动作（走访）的带数字现场页：词表提示证词，声明后才出图",
          walk["decision"] == "none" and walk["source"] == "hint" and walk["suggested"] == "witness"
          and intel.media_necessity(intel.understand("华东 6 家门店走访记录"), "prove",
                                    declared="required")["decision"] == "required",
          str(walk))
    # 21 · 部分与整体需要明确分解语义；算术相加本身不是共同分母的证据
    check("part_whole: 明示其中且合计 100% 的构成候选",
          intel.understand("其中 68% 来自老客户，32% 来自新客")["evidence"] == "part_whole")
    check("structure: 三项以上并列（无数值）是结构",
          intel.understand("渠道分三类：直营、经销、电商")["evidence"] == "structure")
    # 29 · 证据形状第二轮（10 月 14 日）：标注不是内容、端点变化是形状、并列量不是路径、
    # 全角括号里的口径不产生比较、被否定的比较仍是比较、跨单位的两端仍并置
    _anno = intel.understand(
        "翻台率 3.4 次，同店销售额连续 8 个季度为正 2025 年 Q2 同店销售额同比 +12.4%"
        "（口径：可比门店，剔除新开业 90 天内门店）")
    check("shape: 口径标注里的「可比」不产生比较（标注不是内容）",
          _anno["comparison"] is False and _anno["evidence"] == "series",
          str({"cmp": _anno["comparison"], "ev": _anno["evidence"]}))
    check("shape: 端点变化（从 A 到 B，同单位）是变化的形状，不是「无证据」",
          intel.understand("单店回本周期从 19 个月压到 11 个月 样本 42 家门店")["evidence"] == "series"
          and intel.understand("毛利率从 58% 提升到 67%（来源：月报）")["evidence"] == "series",
          str(intel.understand("单店回本周期从 19 个月压到 11 个月")["evidence"]))
    check("shape: 同单位的并列量（规格行）不是序列（三个数字不是充分条件）",
          intel.understand("规格：3 次烧成、1280℃、1450℃、1320℃")["evidence"] != "series"
          and intel.understand("2023 年 38%，2024 年 47%，2025 年 61%")["evidence"] == "series")
    check("shape: 期间词（环比/同比）不是两端比较（变化不是并置）",
          intel.understand("环比 +12.4%；占比 68% 的品类在增长")["comparison"] is False)
    _neg = intel.understand("尚未超过对手 对手门店数 420 家；我们 180 家（来源：行业年报）")
    check("shape: 被否定的比较仍是比较（宾语是名词，不是界限）",
          _neg["comparison"] is True, str(_neg))
    check("shape: 界限（「不超过 5%」）不被当成两端比较（尺子没被放松）",
          intel.understand("毛利率 58%；退货率不超过 5%（来源：月报）")["comparison"] is False)
    check("shape: 「A 对 B」两端的跨单位对照仍构成比较",
          intel.understand("营收 1.2 亿元对 1,800 万元（来源：财报）")["comparison"] is True)
    check("claim: 条件结果句（只要…就能…）是结论，不是证据行",
          intel.extract_claim({"title": "配送门槛",
                               "content": "只要门店密度达到每 10 万人 1.2 家，"
                                          "配送成本就能压到 8 元以内（来源：模型）"})["source"]
          == "extracted")
    check("claim: 复合主张（并列三项 + 缺一不可）是结论，不是清单",
          intel.extract_claim({"title": "降价、扩店、换帅三件事缺一不可"})["source"] == "extracted")
    # 30 · 档案坐标全页生效：样本/口径/来源行是取证范围，不是内容、不是结论、不是主语
    _scoped = intel.understand("单店回本周期从 19 个月压到 11 个月 "
                              "样本 42 家门店；口径 2023 年 1 月—2025 年 6 月；来源 内部经营系统",
                              content="样本 42 家门店；口径 2023 年 1 月—2025 年 6 月；来源 内部经营系统")
    check("scope: 取证范围行不进证据量（样本量不冒充页面的量）",
          [n["raw"] for n in _scoped["numerals"]] == ["19 个月", "11 个月"],
          str([n["raw"] for n in _scoped["numerals"]]))
    check("scope: 取证范围行不夺结论位（样本行不是主张）",
          intel.extract_claim({"title": "原材料成本上行的传导",
                               "content": "样本 42 家门店；口径 2023 年 1 月—2025 年 6 月"})["source"]
          == "absent")
    check("scope: 语义相同的普通句子不被误剥（「客户来源是口碑」照原样参与判断）",
          [n["raw"] for n in intel.understand(
              "客户来源结构 客户来源是口碑推荐，占比 62%；来源 调研 2025",
              content="客户来源是口碑推荐，占比 62%；来源 调研 2025")["numerals"]] == ["62%"])
    # 22 · 基准的已知缺口登记在册：新增失败必须显式出现，不能被静默忽略
    import bench as bench_mod
    bench_report = bench_mod.run()
    failing = {r["id"] for r in bench_report["cases"] if not r["passed"]}
    # 已登记缺口已于 10 月 14 日的判断修复中全部关闭（N02/X02 比较判定、I01/K01 主张抽取）；
    # 集合现在必须是空的——任何新失败都要显式出现，不允许用改期望值的方式消失。
    check("bench: 全部案例通过（4 个已登记缺口 I01/K01/N02/X02 已修复，集合为空）",
          failing == set(), str(sorted(failing)))
    m02 = next(r for r in bench_report["cases"] if r["id"] == "M02")
    check("bench: M02「一位釉工…」media_source=hint 且 suggested=witness（个体语法信号，非词表）",
          m02["actual"]["media_source"] == "hint" and m02["actual"]["media_suggested"] == "witness",
          str(m02["actual"]))
    # 25 · 阈值列举 ≠ 比较：多个子句各带不同单位的数（万元 / % / 人）是并列条件；
    # 两端可比的真比较仍然是比较。实测（S05 风险页）：阈值列举被判成 comparison，
    # 焦点落到不可相减的两个量上。
    thr = intel.understand("单店月净利连续两月低于 15 万元；退货率高于 8%；核心店长流失超过 2 人。")
    check("threshold: 单位互不相容的阈值列举是结构（并列条件），不是比较",
          thr["comparison"] is False and thr["evidence"] == "structure", str(thr["evidence"]))
    check("threshold: 无阈值词的两个不同单位陈述（「开业前 60 天…缩短 4 个月」）不是结构",
          intel.understand("一位店长记录：开业前 60 天完成本地客群访谈的门店，回本周期平均缩短 4 个月。")["evidence"] != "structure")
    check("threshold: 两端同单位的「高于」仍是比较（真比较不被误伤）",
          intel.understand("华东高端线均价 480 元高于泰国 300 元")["comparison"] is True)
    check("threshold: 「是……的 1.6 倍」仍是比较",
          intel.understand("高端线售价是泰国的 1.6 倍")["comparison"] is True)
    # 31 · 判断修复回归（2026-10-15）：形状判据的边界与护栏，不是词表扩充——
    # 字母紧贴编号、口径标注、条件从句、复合分期名、与…不同句形、序列当前值、
    # 分母不作主语、日历月坐标。每条先有跨行业实证，再锁行为。
    check("numerals: 字母紧贴的编号不进证据（3C/28nm 是命名，不是量）",
          [n["raw"] for n in intel.evidence_numerals("3C 类占比 3.2%，28nm 良率 81%")]
          == ["3.2%", "81%"],
          str([n["raw"] for n in intel.evidence_numerals("3C 类占比 3.2%，28nm 良率 81%")]))
    check("numerals: 紧贴的可测量单位是量（220V/2,100W/3.2kg/48dB 是规格值不是代号）",
          [n["raw"] for n in intel.evidence_numerals(
              "额定电压 220V；峰值功率 2,100W；整机重量 3.2kg；噪音 48dB")]
          == ["220V", "2,100W", "3.2kg", "48dB"],
          str([n["raw"] for n in intel.evidence_numerals(
              "额定电压 220V；峰值功率 2,100W；整机重量 3.2kg；噪音 48dB")]))
    check("numerals: 口径标注里的数字是档案坐标（不制造单位冲突、不进证据）",
          intel.claim_units_conflict("翻台率 4.2 次/天（口径：可比门店，剔除新开业 90 天内门店）") is False
          and [n["raw"] for n in intel.evidence_numerals(
              "翻台率 4.2 次/天（口径：可比门店，剔除新开业 90 天内门店）")] == ["4.2 次"],
          str([n["raw"] for n in intel.evidence_numerals(
              "翻台率 4.2 次/天（口径：可比门店，剔除新开业 90 天内门店）")]))
    check("numerals: 裸 N 月是日历坐标（时长写 N 个月才是量）",
          [(n["raw"], n["year"]) for n in intel.numerals("1 月 120 万，2 月 135 万")]
          == [("1 月", True), ("120 万", False), ("2 月", True), ("135 万", False)],
          str([(n["raw"], n["year"]) for n in intel.numerals("1 月 120 万，2 月 135 万")]))
    _guard = intel.number_roles("预算先按 30 万元执行；若提前到 45 万元封顶则需复审",
                               "预算先按 30 万元执行")[1]
    check("role: 条件从句里的护栏值不作主语（「若…则」内的数不抢现实值）",
          _guard["raw"] == "30 万元", str(_guard))
    check("role: 分母（总体）不作主语（「12 家中 9 家」主张落在 9 家）",
          intel.number_roles("12 家试点门店中，9 家续约", "12 家试点门店中，9 家续约")[1]["raw"]
          == "9 家",
          str(intel.number_roles("12 家试点门店中，9 家续约", "12 家试点门店中，9 家续约")[1]))
    check("role: 班次/先后两句里的首句候选不被主张端点规则抢走（早班 62 人）",
          intel.number_roles("早班 62 人，晚班 12 人", "门店净增持续放缓")[1]["raw"] == "62 人",
          str(intel.number_roles("早班 62 人，晚班 12 人", "门店净增持续放缓")[1]))
    _qtxt = "2025Q1 87.2%，2025Q2 85.1%，2025Q3 82.6%，2025Q4 79.3%。"
    _qclaim = {"text": "准点率连续四个季度下滑"}
    _qanchor = intel.number_roles(_qtxt, _qclaim["text"])
    _qfocus = intel.focus_of("prove", intel.understand(_qtxt, content=_qtxt),
                            {"source": "given", "text": _qclaim["text"]},
                            number_anchor=_qanchor, page_text=_qtxt)
    check("role: 同指标序列取当前值（末值），理由是「序列的当前值」不是「变化终值」",
          _qanchor[1]["raw"] == "79.3%" and "序列的当前值" in _qfocus["why"]
          and "变化终值" not in _qfocus["why"], str(_qfocus))
    check("shape: 复合分期名（三期临床）是命名，不是序列",
          intel.understand("三期临床达到主要终点 应答率 63.4%，对照组 41.2%")["evidence"]
          != "sequence")
    check("shape: 「与…不同」的差异陈述是两端比较的形状（差异被主张）",
          intel.understand("清关时效在越南与印尼不同 越南 1.8 天，印尼 6.4 天")["comparison"]
          is True
          and intel.extract_claim({"title": "清关时效在越南与印尼不同",
                                   "content": "越南 1.8 天，印尼 6.4 天"})["source"]
          == "extracted")
    check("shape: 状态/归因词（提交、来自）不制造序列",
          intel.understand("合规整改尚未获批 整改方案已于 3 月提交监管，"
                           "目前处于征求意见阶段")["evidence"] != "sequence"
          and intel.understand("服务营收增长主要来自家电 增速 21.4%，"
                           "其余品类 6.8%")["evidence"] != "sequence")
    check("claim: 逐项编号页整体是主张（「设备 ZX-81；工位 AB-26」类内容不当序列）",
          intel.extract_claim({"title": "设备编号不是步骤",
                               "content": "设备 ZX-81；工位 AB-26。"})["source"]
          == "extracted"
          and intel.understand("设备编号不是步骤 设备 ZX-81；工位 AB-26。")["evidence"]
          != "sequence")
    # 眉标不取数量残片（「15 万元」的「万元」）；派不出主题就省略，不输出角色词
    check("眉标: 数量短语的残片（「万元」）不作眉标",
          intel.section_label({}, {"text": "单店月净利连续两月低于 15 万元"}, "compare") == "")
    check("眉标: 派不出主题时不退回角色词（「解释」「对比」）",
          intel.section_label({}, {"text": "第1页结论句"}, "explain") == "")
    # 26 · 产物行为：说明行（desc）必须真的进入 PPTX（此前 big_number_row / timeline 静默丢弃，
    # release 硬门放行）；溢出的修复包必须指到具体页（此前只报元素 id，页面无法定位）。
    from pptx import Presentation as _Prs
    def _pptx_texts(path: Path) -> list[str]:
        return [sh.text_frame.text for sl in _Prs(str(path)).slides
                for sh in sl.shapes if sh.has_text_frame]
    desc_build = tmp / "desc_build.py"
    desc_build.write_text(
        "SPEC = " + repr({
            "canvas": {"width": 1280, "height": 720, "grid_columns": 12, "grid_unit": 8},
            "theme": {"colors": {"background": "#F4EEE8", "ink": "#12120F", "muted": "#8B8883",
                                 "primary": "#12120F", "secondary": "#8B8883", "accent": "#A2563F"},
                      "fonts": {"cn": "Source Han Serif SC", "latin": "Georgia"}},
            "slides": [
                {"id": "d1", "elements": [
                    {"id": "t", "type": "text", "x": 48, "y": 96, "width": 1184, "height": 64,
                     "text": "风险可控：三个触发条件任一出现即暂停", "size": 40, "color": "ink", "role": "title"},
                    {"id": "trig", "type": "chart", "chart_kind": "big_number_row", "x": 48, "y": 236,
                     "width": 1184, "height": 380, "role": "focus",
                     "data": [{"label": "单店月净利", "value": 15, "display": "15 万元",
                               "desc": "DESC_ALPHA 连续两月低于 15 万元"},
                              {"label": "退货率", "value": 8, "display": "8%", "desc": "DESC_BETA"}]}]},
                {"id": "d2", "elements": [
                    {"id": "t", "type": "text", "x": 48, "y": 96, "width": 1184, "height": 64,
                     "text": "三步推进", "size": 40, "color": "ink", "role": "title"},
                    {"id": "tl", "type": "chart", "chart_kind": "timeline", "x": 48, "y": 200,
                     "width": 1184, "height": 440, "role": "focus",
                     "data": [{"label": "Q4 2026", "value": 3, "desc": "DESC_GAMMA 3 家试点"},
                              {"label": "2027 上半年", "value": 12, "desc": "DESC_DELTA"}]}]}]}),
        encoding="utf-8")
    # 溢出稿件单独成件：draft 会 BLOCK，不产出 PPTX；它只用来验证修复包的页面定位。
    over_build = tmp / "over_build.py"
    over_build.write_text(
        "SPEC = " + repr({
            "canvas": {"width": 1280, "height": 720, "grid_columns": 12, "grid_unit": 8},
            "theme": {"colors": {"background": "#F4EEE8", "ink": "#12120F", "muted": "#8B8883",
                                 "primary": "#12120F", "secondary": "#8B8883", "accent": "#A2563F"},
                      "fonts": {"cn": "Source Han Serif SC", "latin": "Georgia"}},
            "slides": [
                {"id": "d3", "elements": [
                    {"id": "big", "type": "text", "x": 48, "y": 200, "width": 1184, "height": 60,
                     "text": "这是一段会溢出框高的超长文字" * 6, "size": 60, "color": "ink",
                     "role": "focus"}]}]}), encoding="utf-8")
    desc_out = tmp / "desc_out.pptx"
    rd = run_vao("check", str(desc_build), str(desc_out), "--mode", "draft", "--speed", "fast")
    over_out = tmp / "over_out.pptx"
    run_vao("check", str(over_build), str(over_out), "--mode", "draft", "--speed", "fast")
    texts = " ".join(_pptx_texts(desc_out)) if desc_out.exists() else ""
    check("desc: big_number_row 与 timeline 的说明行进入 PPTX（不再静默丢弃）",
          all(k in texts for k in ("DESC_ALPHA", "DESC_BETA", "DESC_GAMMA", "DESC_DELTA")),
          rd.stdout[-300:] + rd.stderr[-300:])
    # 根因：compiler._rows 归一化丢弃 desc 键，下游 desc 分支成为死代码。文字存在不够，
    # 说明行必须落在页内（几何断言；此前只检查文字，页外的文字也会通过）。
    _prs = _Prs(str(desc_out)) if desc_out.exists() else None
    _inside = bool(_prs) and all(
        sh.left >= 0 and sh.top >= 0 and sh.left + sh.width <= _prs.slide_width
        and sh.top + sh.height <= _prs.slide_height
        for sl in _prs.slides for sh in sl.shapes
        if sh.has_text_frame and "DESC_" in sh.text_frame.text)
    check("desc: 说明行几何落在页内（不只是文字存在）", _inside)
    locs = []
    try:
        rep = json.loads((over_out.with_suffix(".repair.json")).read_text(encoding="utf-8"))
        locs = rep.get("affected_slides") or []
    except Exception as exc:  # 缺修复包即失败，不静默吞
        locs = [f"<no repair packet: {exc}>"]
    check("repair: 溢出修复包指到具体页面（affected_slides 含页 id，不只元素 id）",
          "d3" in locs, str(locs))
    # 27 · 设计代码化回归（审计 2026-10-14：三个失败案例先固化，再修复）
    # R1 语域：受众文本里的「媒体」只是受众描述，不能把整副 deck 切成深色舞台。
    def _paper(aud, occ, dec):
        return intel.think({"audience": aud, "occasion": occ, "decision": dec,
                            "slides": [{"id": "a", "title": "华东单店利润 18.4 万元",
                                        "content": "同比 +6.2%（来源：台账）"}]})["deck"]["world"]["paper"]
    check("R1 register: 受众文本中的「媒体」不切换为深色舞台（纸色稳定）",
          _paper("渠道伙伴与媒体", "", "报名参加华东渠道合作计划")
          == _paper("董事会与投资人", "", "批准预算"))
    check("R1 register: 真实的发布场景（场合写明发布会）仍切换为舞台语域",
          _paper("渠道伙伴", "新品发布会现场", "报名参加") != _paper("董事会与投资人", "", "批准预算"))
    # R3 节奏：证据为零的叙事页不因平局全部落到同一构图
    _narr = [{"id": "b1", "title": "一位店长的经验，决定了新店能否快速回本",
              "content": "一位店长记录：开业前 60 天完成客群访谈的门店，回本周期更短。"},
             {"id": "b2", "title": "她每天最早到店，最晚离开", "content": "店长李敏记录了一个月的值班日志。"},
             {"id": "b3", "title": "这套做法，是从一次失败的开店里学来的",
              "content": "第一家店关门后，团队用三周复盘了选址。"}]
    _np = intel.think({"audience": "董事会与投资人", "decision": "批准下一阶段预算",
                       "tension": "投入是否值得？", "slides": _narr})["pages"]
    _comps = [p["judgment"]["composition"]["chosen"] for p in _np]
    check("R3 rhythm: 证据为零的叙事页不因平局连续同构（人物故事 3 页不全相同）",
          len(set(_comps)) > 1, str(_comps))
    # R4 结构：话题清单（选址、排班、库存……）不是结构证据；有分类动词的才是
    _topic = "我们在过去两年里反复讨论门店选址、排班、库存与会员运营之间的关系，每一次复盘都会发现新的依赖。"
    check("R4 structure: 叙述中的话题清单不判为结构证据（无分类线索）",
          intel.understand(_topic, content=_topic)["evidence"] != "structure")
    check("R4 structure: 明示分类的列举仍是结构（渠道分三类：直营、经销、电商）",
          intel.understand("渠道分三类：直营、经销、电商")["evidence"] == "structure")
    # 28 · 设计代码化第二轮（R2 收束决定 / R5 删除候选 / 未消费字段门 / 单位落地）
    # R2：末页的决定类型决定焦点；带数字或主张已请求的末页不被覆盖
    _r2_base = [{"id": "p1", "title": "华东单店利润 18.4 万元，同比增长 6.2%",
                 "content": "2026 年 Q3 华东 42 家门店（来源：门店台账）"}]
    _r2_close = {"id": "pz", "title": "我们需要董事会在本月内给出明确答复", "content": "两种选择各有代价，答复决定下一步节奏。"}
    _r2_focus = {d: intel.think({"audience": "董事会", "decision": d, "slides": _r2_base + [_r2_close]})["pages"][-1]["judgment"]["focus"]["element_role"]
                 for d in ("批准 2027 华东门店扩张预算", "报名参加华东渠道合作计划", "评估华东扩张的方法论是否可复用")}
    check("R2 closing: 决定类型不同 → 末页焦点不同（批准/报名/评估 三类）",
          len(set(_r2_focus.values())) == 3, str(_r2_focus))
    _r2_num = intel.think({"audience": "董事会", "decision": "批准预算", "slides": _r2_base + [
        {"id": "pz", "title": "扩张 30 家门店，回本周期 14 个月", "content": "两种选择各有代价（来源：测算）"}]})["pages"][-1]["judgment"]["focus"]["element_role"]
    check("R2 closing: 带数字的末页不被决定文本覆盖", _r2_num != "decision_ask", _r2_num)
    _r2_ask = intel.think({"audience": "董事会", "decision": "批准预算", "slides": _r2_base + [
        {"id": "pz", "title": "请确认评审名单", "content": "请于本周五前确认"}]})["pages"][-1]["judgment"]["focus"]["element_role"]
    check("R2 closing: 主张已点名请求（请确认…）时焦点就是主张，不另换", _r2_ask != "decision_ask", _r2_ask)
    # R5：高密度页给删除候选（按贡献排序、不含主张数字的句子），硬禁令保持为空
    _dense = {"id": "d1", "title": "华东单店利润 18.4 万元，同比增长 6.2%",
              "content": ("2026 年 Q3 华东 42 家门店参与统计；门店类型包括直营 18 家、加盟 24 家；"
                          "我们在过去两年反复复盘选址、排班、库存与会员运营之间的依赖关系；"
                          "其中 3 个月的数据来自 6 次专项走访，样本覆盖 9 个城市；"
                          "这些问题的讨论记录超过 120 页，团队形成了一套跨部门协作的方式。")}
    _dj = intel.think({"audience": "董事会", "decision": "批准预算", "slides": [_dense]})["pages"][0]["judgment"]
    check("R5 delete: 高密度页给出删除候选（delete_candidates 非空）",
          len(_dj["information_weight"]["delete_candidates"]) >= 1, str(_dj["information_weight"]))
    check("R5 delete: 候选不是硬禁令（delete 与 must_not 不被候选填充）",
          _dj["information_weight"]["delete"] == [] and _dj["production"]["must_not"] == [])
    check("R5 delete: 携带主张数字的句子永不成为候选",
          not any("42 家" in c or "18.4" in c for c in _dj["information_weight"]["delete_candidates"]))
    _sparse = intel.think({"audience": "董事会", "decision": "批准预算", "slides": [
        {"id": "s1", "title": "华东单店利润 18.4 万元", "content": "同比 +6.2%（来源：台账）"}]})["pages"][0]["judgment"]
    check("R5 delete: 低密度页不产生候选（不为凑数删内容）",
          _sparse["information_weight"]["delete_candidates"] == [])
    # 未消费字段门：单位必须落到可见标签；行上的 unit/period 等键必须被 release 硬门拦截
    _unit_dict = ({
        "canvas": {"width": 1280, "height": 720, "grid_columns": 12, "grid_unit": 8},
        "theme": {"colors": {"background": "#F4EEE8", "ink": "#12120F", "muted": "#8B8883",
                             "primary": "#12120F", "secondary": "#8B8883", "accent": "#A2563F"},
                  "fonts": {"cn": "Source Han Serif SC", "latin": "Georgia"}},
        "slides": [{"id": "u1", "elements": [
            {"id": "t", "type": "text", "x": 48, "y": 96, "width": 1184, "height": 64,
             "text": "华东高端线售价", "size": 40, "color": "ink", "role": "title"},
            {"id": "c", "type": "chart", "chart_kind": "comparison_bar", "x": 48, "y": 200,
             "width": 1184, "height": 440, "role": "focus", "unit": "元",
             "source": "来源：门店监测", "period": "2026 年 9 月", "basis": "同口径",
             "data": [{"label": "华东", "value": 480}, {"label": "泰国", "value": 300}]}]}]})
    _unit_spec = "SPEC = " + repr(_unit_dict)
    _u_build = tmp / "unit_build.py"
    _u_build.write_text(_unit_spec, encoding="utf-8")
    _u_out = tmp / "unit_out.pptx"
    run_vao("check", str(_u_build), str(_u_out), "--mode", "draft", "--speed", "fast")
    import zipfile as _zf
    _chart_xml = ""
    if _u_out.exists():
        with _zf.ZipFile(_u_out) as _z:
            _chart_xml = " ".join(_z.read(n).decode("utf-8", "ignore") for n in _z.namelist()
                                  if n.startswith("ppt/charts/") and n.endswith(".xml"))
    check("unit: 元素级单位落到数据标签（General\"元\"），不再静默丢弃",
          'General&quot;元&quot;' in _chart_xml or 'General"元"' in _chart_xml)
    check("unit: 数值轴上限为最大值的 1.15 倍（外侧标签留白，XML 层确认；ghost 不读取该值）",
          re.search(r'<c:max val="(\d+(?:\.\d+)?)"', _chart_xml) is not None
          and abs(float(re.search(r'<c:max val="(\d+(?:\.\d+)?)"', _chart_xml).group(1)) - 480 * 1.15) < 0.01)
    import copy as _copy
    _bd = _copy.deepcopy(_unit_dict)
    _bd["slides"][0]["elements"][1]["data"][0]["unit"] = "元"   # 行级 unit：渲染器不消费 → 必须拦截
    _bad = "SPEC = " + repr(_bd)
    _bad_build = tmp / "unit_bad.py"
    _bad_build.write_text(_bad, encoding="utf-8")
    _rb = run_vao("check", str(_bad_build), str(tmp / "unit_bad.pptx"), "--mode", "release", "--speed", "fast")
    check("unit gate: release 硬门拦截图表行上的未消费字段 unit",
          "BLOCKED" in _rb.stdout, _rb.stdout[-300:] + _rb.stderr[-200:])
    _ok = run_vao("check", str(_u_build), str(tmp / "unit_ok.pptx"), "--mode", "release", "--speed", "fast")
    check("unit gate: 元素级 unit 的正向样例 release 仍 PASS", "PASS" in _ok.stdout, _ok.stdout[-200:])
    # R6 盲评通道：一致性计算的边界；评审不足三位时不宣称共识（盲评本身仍未执行）
    import bench as _bench
    check("blind: ICC(2,1) 完全一致时为 1，完全随机时低于共识线",
          abs(_bench._icc_2_1([[1, 1, 1], [3, 3, 3], [5, 5, 5]]) - 1.0) < 1e-9
          and (_bench._icc_2_1([[1, 5, 3], [5, 1, 3], [3, 3, 5], [1, 3, 1]]) or 0) < 0.5)
    _bsheet = tmp / "blind_sheet.csv"
    _bsheet.write_text("reviewer,item,narrative,focus,composition,typography,completeness,reason\n"
                       "r1,V01,4,4,4,4,4,\nr1,V02,3,3,3,3,3,\n", encoding="utf-8")
    _bkey = tmp / "blind.key.json"
    _bkey.write_text('{"V01": "a.pptx", "V02": "b.pptx"}', encoding="utf-8")
    check("blind: 评审不足三位时输出 NOT_ENOUGH_REVIEWERS（不宣称共识）",
          _bench.blind_score([str(_bsheet)], str(_bkey))["consensus"] == "NOT_ENOUGH_REVIEWERS")
    # 盲评链路审计钉（bench 轮）：文献对拍 / 匿名净化 / 溯源 / 量表 / 统计诚实性
    check("blind: ICC(2,1) 对拍 Shrout & Fleiss (1979) 文献值 0.290",
          abs(_bench._icc_2_1([[9, 2, 5, 8], [6, 1, 3, 2], [8, 4, 6, 8],
                               [7, 1, 2, 6], [10, 5, 6, 9], [6, 2, 4, 7]]) - 0.290) < 5e-4)
    from pptx import Presentation as _Prs
    import zipfile as _zf
    _prs = _Prs()
    _prs.core_properties.author = "审计作者"
    _vp = tmp / "ident.pptx"
    _prs.save(str(_vp))
    _pk = _bench.blind_pack([str(_vp)], tmp / "blindpack", seed=11)
    _core = _zf.ZipFile(tmp / "blindpack" / "V01.pptx").read("docProps/core.xml").decode("utf-8")
    _key2 = json.loads(Path(_pk["key"]).read_text(encoding="utf-8"))
    check("blind: 打包抹去 pptx 作者身份并记录前后哈希（材料可查调包）",
          "审计作者" not in _core
          and _key2["sha256"]["V01"]["source_sha256"] != _key2["sha256"]["V01"]["packed_sha256"]
          and bool(_key2.get("packed_at_utc")) and _key2.get("seed") == 11)
    import shutil as _shutil
    import io as _io
    import csv as _csv
    _vd = tmp / "deckdir"
    (_vd / "out_preview").mkdir(parents=True)
    (_vd / "out.manifest.json").write_text("{}", encoding="utf-8")
    (_vd / "out_preview" / "ghost.meta.json").write_text("{}", encoding="utf-8")
    _shutil.copyfile(_vp, _vd / "deck.pptx")
    _bench.blind_pack([str(_vd)], tmp / "blinddir", seed=1)
    _copied = [q.name for q in (tmp / "blinddir" / "V01").rglob("*")]
    check("blind: 目录源不携带流程工件（manifest/meta 留在评审目录外）",
          "out.manifest.json" not in _copied and "ghost.meta.json" not in _copied
          and "deck.pptx" in _copied)
    _k3 = tmp / "k3.json"
    _k3.write_text(json.dumps({"V01": "a.pptx", "V02": "b.pptx", "V03": "c.pptx"}), encoding="utf-8")

    def _wsheet(rows, name):
        stream = _io.StringIO(newline="")
        writer = _csv.writer(stream)
        writer.writerow(["reviewer", "item", *_bench.BLIND_DIMS, "reason"])
        for r_ in rows:
            writer.writerow(r_)
        p_ = tmp / name
        p_.write_text(stream.getvalue(), encoding="utf-8")
        return p_

    _full = lambda r_, i_, v_: [r_, i_, *([str(v_)] * 5), ""]
    try:
        _bench.blind_score([str(_wsheet([_full("r1", "V01", 9), _full("r1", "V02", 1),
                                         _full("r1", "V03", 1)], "oob.csv"))], str(_k3))
        _rej = False
    except ValueError:
        _rej = True
    check("blind: 量表外分数拒绝（1–7 之外不入统计）", _rej)
    _rep_same = _bench.blind_score(
        [str(_wsheet([_full(r_, i_, 4) for r_ in ("r1", "r2", "r3")
                      for i_ in ("V01", "V02", "V03")], "same.csv"))], str(_k3))
    check("blind: 完全一致矩阵按「一致」处理并注明零方差（不再误判 NO_CONSENSUS）",
          _rep_same["consensus"] == "CONSENSUS_SUPPORTED" and "零方差" in _rep_same["dimensions"]["叙事"]["note"])
    _rep2 = _bench.blind_score(
        [str(_wsheet([_full("r1", i_, v_) for i_, v_ in (("V01", 5), ("V02", 3), ("V03", 4))]
                     + [_full("r2", i_, v_) for i_, v_ in (("V01", 4), ("V02", 4), ("V03", 5))],
                     "two.csv"))], str(_k3))
    check("blind: 评审不足三位仍报告均值与缺口（不丢弃已收数据）",
          _rep2["consensus"] == "NOT_ENOUGH_REVIEWERS" and _rep2["dimensions"]["叙事"]["mean"]
          and _rep2["missing_items_by_reviewer"] == {})
    _vec = lambda r_: [_full(r_, "V01", 5), _full(r_, "V02", 3), _full(r_, "V03", 4)]
    _rep_dup = _bench.blind_score(
        [str(_wsheet(_vec("r1") + _vec("r2")
                     + [_full("r3", i_, v_) for i_, v_ in (("V01", 2), ("V02", 5), ("V03", 3))],
                     "dup.csv"))], str(_k3))
    check("blind: 完全相同向量双评审被标记（保留数据交组织者复核，不静默平均）",
          _rep_dup.get("suspect_identical_reviewers") == [["r1", "r2"]])
    _rep_ci = _bench.blind_score(
        [str(_wsheet([_full("r1", "V01", 5), _full("r1", "V02", 3), _full("r1", "V03", 4),
                      _full("r2", "V01", 4), _full("r2", "V02", 4), _full("r2", "V03", 5),
                      _full("r3", "V01", 5), _full("r3", "V02", 3), _full("r3", "V03", 3)],
                     "ci.csv"))], str(_k3))
    _dim = _rep_ci["dimensions"]["叙事"]
    check("blind: ICC 带自助法区间与小样本注记（不假装点估计很准）",
          isinstance(_dim.get("icc_ci95"), list) and _dim["n_items"] == 3 and "版本少" in _dim["note"]
          and _dim["icc_ci95"][0] <= _dim["icc"] <= _dim["icc_ci95"][1])
    check("blind: 汇总报告写明一致性与质量的边界（读法不靠约定俗成）",
          "不衡量设计质量" in _rep_ci["how_to_read"] and _rep_ci["identified_report"] is True)
    # R7 文档：SKILL 保留逐页五问，且五问不被删成空话
    check("docs: SKILL §3 保留逐页五问（受众 / 主张 / 证据 / 焦点 / 构图理由）",
          all(k in (ROOT / "SKILL.md").read_text(encoding="utf-8")
              for k in ("每页先答五问", "唯一主张", "证据关系", "视觉焦点", "构图理由")))
    # 24 · 个体语法信号：一位具体的人是证词线索；计数（12 位客户）不是个体；新职业不靠词表
    check("individual: 「一位釉工」为个体（不依赖职业词表）",
          intel.understand("一位釉工用了三十年时间调配这一抹青")["individual"] is True)
    check("individual: 「12 位企业客户」是计数，不是个体",
          intel.understand("来自 12 位企业客户访谈")["individual"] is False)
    check("individual: 新职业「一名窑工」同样为个体",
          intel.understand("一名窑工靠手感判断温度")["individual"] is True)
    check("individual: 「位/名」的非人复合词（名称/名单/位置）不误判为个体",
          not any(intel.understand(x)["individual"] for x in ("这名称由团队统一确定", "该名单每周更新", "某位置的库存已调整")))
    check("individual: 个体只在证词角色下出 hint；compare 角色不把人物当证词",
          intel.media_necessity(intel.understand("一位釉工用了三十年时间调配这一抹青"), "persuade")["source"] == "hint"
          and intel.media_necessity(intel.understand("一位釉工用了三十年时间调配这一抹青"), "compare")["decision"] == "none")
    check("bench: 基准严重缺陷率为 0（critical 案例全部通过）",
          bench_report["metrics"]["critical_defect_rate"] == 0.0,
          str(bench_report["metrics"]))
    # 23 · 作者声明的角色与证据原样生效；非法声明给 warning，不静默吞掉
    decl_w: list = []
    decl_role = intel.page_judgment(
        {"id": "d1", "title": "净利率 18.4%", "role": "persuade"}, index=3, total=8,
        brief={"audience": "评审", "decision": "批准"},
        world=intel.derive_world({"audience": "评审"}, [intel.understand("x")]), warnings=decl_w)
    check("declare: 作者声明 role 原样生效（source=author），压过内容判断",
          decl_role["visual_role"]["role"] == "persuade" and decl_role["visual_role"]["source"] == "author",
          str(decl_role["visual_role"]))
    bad_w: list = []
    bad = intel.page_judgment(
        {"id": "d2", "title": "净利率 18.4%", "role": "hero", "evidence": "magic"}, index=3, total=8,
        brief={"audience": "评审", "decision": "批准"},
        world=intel.derive_world({"audience": "评审"}, [intel.understand("x")]), warnings=bad_w)
    rules = {w["rule"] for w in bad_w}
    check("declare: 非法 role/evidence 不静默生效：给出 warning 并按内容判断",
          bad["visual_role"]["role"] == "prove" and bad["visual_role"].get("source") == "judgment"
          and {"declared_role_invalid", "declared_evidence_invalid"} <= rules, str(rules))
    check("hint: 词表命中的人物/场所不出图，并把问题交还作者（open_questions 中可见）",
          any("声明 asset" in q for q in (intel.page_judgment(
              {"id": "h1", "title": "华东 6 家门店走访记录"}, index=3, total=8,
              brief={"audience": "评审", "decision": "批准"},
              world=intel.derive_world({"audience": "评审"}, [intel.understand("x")]),
              warnings=[]).get("open_questions") or [])))


def test_cjk_rhythm_regression(tmp: Path):
    """构图排版（中文）与跨页节奏的本轮回归：每条钉住一个可复现的误判现场。

    覆盖：长主张断行建议、口径句不当主张、KPI 面板角色、双指标当前值、
    眉标断词碎片、相邻页数字复现、节奏方向语义、正文请求句与超长主张提示。
    """
    import intelligence as intel

    # 1 · 断行建议此前只认标点/系词/「的」：无标点的长主张（中文标题最常见形态）
    #     拿不到任何断口（实测「…制把单均成本压到 5.4 元」返回 None）。
    check("break: 处置式「把」字句是谓语起点，给出语义断口",
          intel.suggest_break("同行领先者用区域枢纽制把单均成本压到 5.4 元")
          == "同行领先者用区域枢纽制｜把单均成本压到 5.4 元",
          str(intel.suggest_break("同行领先者用区域枢纽制把单均成本压到 5.4 元")))
    check("break: 结果连词「带动」领起新谓语；时间单位不再劈开「全年经常性收入」",
          intel.suggest_break("以用户留存为核心重构会员体系带动全年经常性收入增长")
          == "以用户留存为核心重构会员体系｜带动全年经常性收入增长",
          str(intel.suggest_break("以用户留存为核心重构会员体系带动全年经常性收入增长")))
    check("break: 事件状语「后」之后是合法断口",
          intel.suggest_break("试点落地后单店模型跑通复制速度取决于店长供给")
          == "试点落地后｜单店模型跑通复制速度取决于店长供给",
          str(intel.suggest_break("试点落地后单店模型跑通复制速度取决于店长供给")))
    # 2 · 行长上限：旧 cap 随句长放宽（len//2+4），41–48 字的主张被建议 21–26 字
    #     的一行——超出主张字阶的可读域。两行装不下就如实说装不下。
    check("break: 超出两行预算的断口不再给出（宁无可改写提示，不给超宽行）",
          intel.suggest_break("长三角三城联动与区域仓网重构之后配送成本下降四成但履约时长仍在持续上升爬坡")
          is None,
          str(intel.suggest_break("长三角三城联动与区域仓网重构之后配送成本下降四成但履约时长仍在持续上升爬坡")))
    long_claim = "投产后单均履约成本将从当前的 7.2 元一路回落到 5.8 元并且在两年之内收回全部 3.2 亿元的改造投资"
    judged_long = intel.page_judgment(
        {"id": "q0", "title": long_claim, "content": "预计节约 2.6 亿元"},
        index=1, total=3, brief={"audience": "委员会", "decision": "批准"},
        world=intel.derive_world({"audience": "委员会"}, [intel.understand("x")]), warnings=[])
    check("typography: 两行装不下的主张改提「先改写」，不硬给断点",
          judged_long["typography"]["scale"].get("lead_rewrite") is not None
          and judged_long["typography"]["scale"].get("lead_break") is None,
          str(judged_long["typography"]["scale"]))

    # 3 · 口径句被抽成主张：「数据周期为 2026 年 4 至 9 月」曾凭「短句＋数字」
    #     赢下结论位（实测）。期间/范围声明与「样本 42 家」同池出局。
    scope_claim = intel.extract_claim(
        {"title": "试点县城的经营数据",
         "content": "本轮试点覆盖 12 个县城：单均利润 1.9 元，月订单 8.4 万单。数据周期为 2026 年 4 至 9 月"})
    check("claim: 口径/期间声明句不当主张（ 宁可 absent，不拿档案坐标当结论）",
          scope_claim["text"] is None and scope_claim["source"] == "absent",
          str(scope_claim))

    # 4 · KPI 面板页：多指标并列（无单一形状）曾判 persuade（观点页），焦点落到
    #     statement_text——这一页靠一组可核验的量成立，是证据页。主张自己带着
    #     数字的不在此列：数字是被讨论的对象，页面仍在下判断。
    world = intel.derive_world({"audience": "管理层"}, [intel.understand("x")])
    panel = intel.page_judgment(
        {"id": "p1", "title": "试点县城单均利润为正但薄",
         "content": "试点 12 个县城，单均利润 1.9 元，月订单 8.4 万单，履约成本 6.2 元/单，客单价 58 元"},
        index=2, total=6, brief={"audience": "管理层", "decision": "批准"},
        world=world, warnings=[])
    specimen = intel.page_judgment(
        {"id": "p2", "title": "相加为100不是共同分母",
         "content": "60% 缓解率，40% 不良反应发生率，同一患者可同时满足两者"},
        index=2, total=6, brief={"audience": "管理层", "decision": "批准"},
        world=world, warnings=[])
    check("role: 多指标并列面板是证据页（prove），不再判成观点页",
          panel["visual_role"]["role"] == "prove",
          str(panel["visual_role"]["role"]))
    check("role: 主张自带数字的标本页仍是观点页（数字是被讨论的对象）",
          specimen["visual_role"]["role"] == "persuade",
          str(specimen["visual_role"]["role"]))

    # 5 · 双指标系列的主语：主张点名两个指标时全部取值同分、单位不同挡住末值
    #     回退——max 落在第一个数上，2023 年的基期值成了页面主语（实测）。
    _, primary = intel.number_roles(
        "2023年订单量 1,840 万单，履约成本 11.2 亿元；2024年订单量 2,150 万单，"
        "履约成本 13.8 亿元；2025年订单量 2,490 万单，履约成本 16.1 亿元",
        "去年订单量增长与履约成本同步上升")
    check("focus: 双指标系列按「指标内末值＋主张主语位」取当前量，不回基期",
          primary is not None and primary["raw"] == "2,490 万",
          str(primary and primary["raw"]))

    # 6 · 眉标断词：单字边界切在词内，碎片通过长度与首尾检查成了页眉
    #     （「率不足」「保留现」均为实测）。完整主语短语（7 字）照常派生。
    for text, want, why in (
            ("利用率不足的担忧通过分期投放来化解", "利用率不足", "「用」不再劈开「利用率」"),
            ("建议建设两个区域枢纽仓并保留现有干线", "保留现有干线", "「有」不再劈开「现有」"),
            ("去年订单量增长与履约成本同步上升", "去年订单量增长", "7 字完整主语好过缺省"),
            ("付费会员的核心价值是把一次性交易变成持续关系", "", "切口落在两字词内部时宁可缺省")):
        got = intel.section_label({}, {"text": text}, "prove")
        check(f"eyebrow: {why}", got == want, repr(got))

    # 7 · 相邻页信息增量：后页数字完全复现前页（无一新增）= 重复铺陈候选，
    #     交作者裁决；末页回收量（summarize）与比较页引用两端不在其列。
    dup = intel.think({"audience": "管理层", "decision": "批准", "tension": "成本", "slides": [
        {"id": "r1", "title": "试点数据",
         "content": "试点 12 个县城，单均利润 1.9 元，月订单 8.4 万单，履约成本 6.2 元/单，客单价 58 元"},
        {"id": "r2", "title": "试点经营情况",
         "content": "本轮试点覆盖 12 个县城：单均利润 1.9 元，月订单 8.4 万单，履约成本 6.2 元/单，客单价 58 元"},
        {"id": "r3", "title": "谢谢", "content": "供应链部"}]})
    check("coherence: 相邻页数字完全复现进 notes（确认增量，不自动改页）",
          dup["deck"]["coherence"]["notes"]
          and "r2 与 r1 数字完全复现" in dup["deck"]["coherence"]["notes"][0],
          str(dup["deck"]["coherence"]["notes"]))
    exempt = intel.think({"audience": "管理层", "decision": "批准", "tension": "成本", "slides": [
        {"id": "e1", "title": "试点数据",
         "content": "试点 12 个县城，单均利润 1.9 元，月订单 8.4 万单，履约成本 6.2 元/单，客单价 58 元"},
        {"id": "e2", "title": "谢谢",
         "content": "试点 12 个县城，单均利润 1.9 元，月订单 8.4 万单"}]})
    check("coherence: 收束页回收已建立的量不算重复",
          exempt["deck"]["coherence"]["notes"] is None,
          str(exempt["deck"]["coherence"]["notes"]))

    # 8 · 节奏方向：密度抬升是加压不是喘息——两个方向共用 breathe 会把
    #     「负荷增加」说成「让出阅读时间」。
    tempo = intel.think({"audience": "管理层", "decision": "批准", "tension": "成本", "slides": [
        {"id": "a", "title": "开场", "content": "我们只做一件事"},
        {"id": "b", "title": "面板",
         "content": "华东营收 3.2 亿元，华南营收 2.8 亿元，西南营收 1.1 亿元，华北 0.9 亿元，海外 0.3 亿元"},
        {"id": "c", "title": "谢谢", "content": "供应链部"}]})
    check("rhythm: 密度抬升记 build（负荷上升），密度回落才是 breathe",
          [p["judgment"]["rhythm"]["move"] for p in tempo["pages"]]
          == ["open", "build", "close"],
          str([p["judgment"]["rhythm"]["move"] for p in tempo["pages"]]))

    # 9 · 请求句藏在正文里而主张没点名它：数字焦点合法占了第一落点，但请求被
    #     挤成一行正文＝决定被推迟（实测「请委员会批准…」被 5.8 元压掉）。
    ask_page = intel.page_judgment(
        {"id": "q1", "title": "投产后单均成本可回到 5.8 元，两年收回投资",
         "content": "预计节约 2.6 亿元。请委员会批准方案并同意 11 月启动一期建设"},
        index=1, total=3, brief={"audience": "委员会", "decision": "批准"},
        world=world, warnings=[])
    check("open_questions: 正文含请求句时提醒作者把它落到收束位",
          any("请求" in q for q in ask_page["open_questions"] or []),
          str(ask_page["open_questions"]))


def test_ghost_observations(tmp: Path):
    """ghost 专项审计回归：预览的证据诚实性与口径对齐（每条对应一个实测缺陷）。

    覆盖：max_lines 截断可见（页面标记 + 台账）、段距与产物同义、混排间隙
    与 compiler/guard 同一份文本、联络表混尺寸页图对齐格子、证据边界随图走。
    这些是「结构预览」对自身边界说话的纪律，不是把预览当第二个渲染器。
    """
    import ghost as g
    from ghost import ghost_page, ghost_deck, text_truncations, make_contact_sheet

    theme = {"colors": {"background": "#FFFFFF", "ink": "#111111", "muted": "#888888"}}

    def spec(elements):
        return {"canvas": {"width": 1280, "height": 720}, "theme": theme,
                "slides": [{"id": "s1", "elements": elements}]}

    def slide_of(elements):
        return {"id": "s1", "elements": elements}

    # 1 · max_lines 截断不再静默：产物（PowerPoint 原生文本）从不截断，预览按
    #     声明收笔时必须留下「还有 N 行未显示」的可见标记与台账记录。
    #     截断机制与字面无关：拉丁文本在真字体/回退字体下走同一套代码，先钉机制
    #     本身（CI 的 Linux 无 CJK 字面，正好覆盖回退路径）；中文度量再按字面
    #     环境分支——有真 CJK 字面验度量，没有则验回落记账，不假装有字体。
    from ghost import _wrap_lines, _font
    latin_text = ("Define the layout decision in one short repeatable sentence "
                  "for every single slide review meeting.")
    g._TRUNCATIONS.clear()
    latin_page = ghost_page(slide_of([{"id": "en", "type": "text", "x": 100, "y": 100,
                                       "width": 150, "height": 80, "text": latin_text,
                                       "size": 16, "color": "ink", "max_lines": 2}]),
                            spec([]), scale=1.0)
    px = latin_page.load()
    corner = [(x, y) for y in range(156, 186) for x in range(110, 252)
              if sum(px[x, y][:3]) < 500]
    expected_dropped = max(0, len(_wrap_lines(latin_text, _font(16, cjk=False), 150)) - 2)
    check("ghost: max_lines 截断在框右下角留下可见标记",
          expected_dropped > 0 and len(corner) > 30,
          f"dropped={expected_dropped} corner={len(corner)}")
    check("ghost: 截断台账按「页/元素」记录被丢弃的行数",
          text_truncations().get("s1/en") == expected_dropped, str(text_truncations()))
    long_text = "这是一段相当长的中文文本，它的折行数会超过作者声明的最大行数，观察预览如何标记截断。"
    has_cjk_metrics = _font(18, cjk=True).getlength("第") >= 12
    if has_cjk_metrics:
        g._TRUNCATIONS.clear()
        ghost_page(slide_of([{"id": "t1", "type": "text", "x": 100, "y": 100,
                              "width": 240, "height": 70, "text": long_text,
                              "size": 18, "color": "ink", "max_lines": 2}]),
                   spec([]), scale=1.0)
        check("ghost: 中文文本按真实 CJK 度量折行截断（13 字/行 × 4 行 → 丢 2 行）",
              text_truncations().get("s1/t1") == 2, str(text_truncations()))
    else:
        g._TRUNCATIONS.clear()
        g._FONT_FALLBACKS.clear()
        g._FONT_CACHE.clear()   # 回落只在真实加载时记账；清缓存强制重走加载路径
        ghost_page(slide_of([{"id": "t1", "type": "text", "x": 100, "y": 100,
                              "width": 240, "height": 70, "text": long_text,
                              "size": 18, "color": "ink", "max_lines": 2}]),
                   spec([]), scale=1.0)
        check("ghost: 本机无 CJK 字面时回落记账（缺字面必须可见，不静默降级）",
              g.font_fallbacks() != {}, str(g.font_fallbacks()))
    g._TRUNCATIONS.clear()
    ghost_page(slide_of([{"id": "t1", "type": "text", "x": 100, "y": 100,
                      "width": 240, "height": 70, "text": latin_text,
                      "size": 18, "color": "ink"}]), spec([]), scale=1.0)
    check("ghost: 未声明 max_lines 时不产生截断记录（溢出画出框外，与产物同向）",
          text_truncations() == {}, str(text_truncations()))

    # 2 · 段距与产物同义（compiler 写 p.space_before/space_after）：两段之间的
    #     空隙必须随 space_before/after 拉开，预览不再把多段落排得比产物紧。
    def bands(elem):
        im = ghost_page(slide_of([elem]), spec([]), scale=1.0)
        p = im.load()
        rows = [y for y in range(40, 500)
                if sum(1 for x in range(90, 720) if sum(p[x, y][:3]) < 600) > 2]
        out, start = [], rows[0]
        for a, b in zip(rows, rows[1:]):
            if b - a > 3:
                out.append((start, a))
                start = b
        out.append((start, rows[-1]))
        return out

    two_para = {"id": "p", "type": "text", "x": 100, "y": 50, "width": 600,
                "height": 400, "text": "第一段\n第二段", "size": 20, "color": "ink"}
    tight = bands(two_para)
    spaced = bands(dict(two_para, space_before=40, space_after=20))
    base_gap = tight[1][0] - tight[0][1]
    grew_gap = spaced[1][0] - spaced[0][1]
    check("ghost: 段前/段后距计入段间隙（与 compiler 段落间距同义）",
          grew_gap - base_gap == 60,
          f"base={base_gap} grew={grew_gap}")

    # 3 · 混排间隙总是插入：compiler 对所有文本 run 插脚本间隙、guard 按间隙后
    #     文本量宽，预览曾只在显式字距时插入——同一行文本三个口径。
    import primitives as prim
    calls = {"n": 0}
    real = prim.insert_script_gaps

    def spy(t):
        calls["n"] += 1
        return real(t)

    prim.insert_script_gaps = spy
    try:
        ghost_page(slide_of([{"id": "m", "type": "text", "x": 100, "y": 100,
                          "width": 600, "height": 60, "text": "订单量 1,840 万单",
                          "size": 20, "color": "ink"}]), spec([]), scale=1.0)
    finally:
        prim.insert_script_gaps = real
    check("ghost: 字距缺省时同样插入混排间隙（与产物/guard 同一份文本）",
          calls["n"] >= 1, str(calls))

    # 4 · 联络表：独立调用可能给出不同尺寸的页图，网格按第一张算——
    #     不缩放会把错误的结构画进证据（大图压相邻格 / 小图周围露底）。
    from PIL import Image as _Img
    d = tmp / "sheet"
    d.mkdir(parents=True, exist_ok=True)
    big, small = d / "ghost-01.png", d / "ghost-02.png"
    _Img.new("RGB", (640, 360), (200, 60, 60)).save(big)
    _Img.new("RGB", (320, 180), (60, 60, 200)).save(small)
    sheet = make_contact_sheet([big, small], d / "sheet.png", columns=2)
    center = _Img.open(sheet).getpixel((640 + 24 + 160, 24 + 180))
    check("ghost: 联络表混尺寸页图对齐格子尺寸（不把错误结构画进证据）",
          center == (60, 60, 200), str(center))

    # 5 · 证据边界随图走：联络表是最常被单独传播的工件，标签带（表格 chrome，
    #     不碰页面像素）必须说清「结构预览 · 非目标软件渲染」。
    mark = _Img.open(sheet).crop((400, 384, 664, 414))
    dark = sum(1 for pxl in mark.getdata() if sum(pxl[:3]) < 620)
    check("ghost: 联络表标签带带证据边界声明（结构预览 ≠ 目标软件渲染）",
          dark > 40, str(dark))

    # 6 · bbox 完全出画布的元素记账（部分出血不算）：观众在预览和产物里都看不到它，
    #     静默裁掉会让「看不见」混进「不存在」。
    g._TRUNCATIONS.clear()
    g._OUT_OF_CANVAS.clear()
    ghost_page(slide_of([
        {"id": "bleed", "type": "shape", "x": 1200, "y": 100, "width": 200,
         "height": 100, "fill": "ink"},
        {"id": "off", "type": "shape", "x": 1400, "y": 100, "width": 200,
         "height": 100, "fill": "ink"},
        {"id": "neg", "type": "shape", "x": -300, "y": 100, "width": 200,
         "height": 100, "fill": "ink"},
    ]), spec([]), scale=1.0)
    check("ghost: bbox 完全出画布的元素记账（部分出血不算）",
          g.out_of_canvas() == {"s1": ["off", "neg"]}, str(g.out_of_canvas()))

    # 7 · ghost_deck 是台账的轮次边界：上一轮的截断记录留在缓存 PNG 与上轮
    #     meta 里，新的一轮从零开始记账。
    g._TRUNCATIONS.clear()
    g._TRUNCATIONS["stale/x"] = 9
    deck_spec = spec([{"id": "t1", "type": "text", "x": 100, "y": 100, "width": 150,
                       "height": 80, "text": latin_text, "size": 16, "color": "ink",
                       "max_lines": 2}])
    ghost_deck(deck_spec, tmp / "deck_preview", scale=0.5)
    deck_dropped = max(0, len(_wrap_lines(latin_text, _font(8, cjk=False), 75)) - 2)
    check("ghost: ghost_deck 开新账（上轮截断记录不串轮）",
          text_truncations() == {"s1/t1": deck_dropped}, str(text_truncations()))


def test_audit_boundaries(tmp):
    """独立关系不变式 + 字体/文件缓存身份；不把结构证据升级为视觉终验。"""
    import os
    import ghost as g
    import bench
    import intelligence
    from unittest.mock import patch
    from primitives import file_digest, json_read_cached
    corpus = json.loads((SCRIPTS / "json" / "judgment_cases.json").read_text(encoding="utf-8"))
    groups = {name: [case for case in corpus["cases"] if case.get("suite") == name]
              for name in ("cases", "holdout", "transfer", "semantic_transfer")}
    check("bench: 精简后四组判断案例未丢失",
          {name: len(items) for name, items in groups.items()} ==
          {"cases": 26, "holdout": 12, "transfer": 6, "semantic_transfer": 13})
    for row in bench.run()["cases"]:
        if row["suite"] != "cases":
            check("judgment regression: " + row["id"], row["passed"], str(row["checks"]))
    acceptance = corpus["acceptance_brief"]
    pages = {p["id"]: p["judgment"] for p in intelligence.think(acceptance)["pages"]}
    check("plan: 泛称不是证词、并列职责不是时间轴、证据不能顶替标题",
          pages["s02"]["visual_role"]["role"] == "persuade"
          and pages["s03"]["understanding"]["evidence"] == "structure"
          and pages["s03"]["composition"]["chosen"] == "linear_structure"
          and len(pages["s03"]["understanding"].get("relations") or []) == 2
          and pages["s05"]["claim"]["text"] == acceptance["slides"][4]["title"]
          and pages["s06"]["claim"]["text"] == acceptance["slides"][5]["title"])
    probe = subprocess.run([sys.executable, "-c",
        f"import sys; sys.path.insert(0, {str(SCRIPTS)!r}); import compiler; assert 'xlsxwriter' not in sys.modules"],
        capture_output=True, encoding="utf-8")
    check("import: 无原生图表任务不预载工作簿写入器", probe.returncode == 0, probe.stderr[-300:])
    # 公开原语不得靠调用方的守卫才不炸：split_runs("") 曾抛 IndexError，
    # 编译器只是恰好用 `if not line: continue` 避开了它。
    from primitives import split_runs
    check("primitives: split_runs 空串返回空列表（不在公开原语里埋 IndexError）",
          split_runs("") == [] and split_runs("中A") == [("中", True), ("A", False)])
    # 包级后处理写盘中断：不许在交付物旁边留下半截 .tmp（实测损坏条目会中途抛错）。
    import io as _io, zipfile as _zip
    from compiler import _postprocess_package, compile_deck as _compile_deck
    _healthy = tmp / "postprocess-healthy.pptx"
    _compile_deck({"canvas": {"width": 320, "height": 180},
                   "theme": {"colors": {"background": "#FFFFFF", "ink": "#111111"}},
                   "slides": [{"id": "s1", "elements": [
                       {"type": "text", "id": "t", "x": 10, "y": 10, "width": 200,
                        "height": 40, "size": 18, "text": "hello", "color": "ink"}]}]},
                  _healthy)
    _raw = bytearray(_healthy.read_bytes())
    with _zip.ZipFile(_io.BytesIO(bytes(_raw))) as _z:
        _info = [i for i in _z.infolist() if i.filename.startswith("ppt/theme/")][0]
        _off = _info.header_offset + 30 + len(_info.filename) + len(_info.extra)
    _raw[_off + 5] ^= 0xFF                      # 打坏一个条目的压缩流
    _broken = tmp / "postprocess-broken.pptx"
    _broken.write_bytes(bytes(_raw))
    try:
        _postprocess_package(_broken)
    except Exception:
        pass
    check("compiler: 后处理中断不留半截包（临时文件泄漏）",
          not _broken.with_suffix(_broken.suffix + ".tmp").exists())
    # 缓存写盘失败必须可见（否则「缓存坏了」与「冷启动」长得一样，每轮白重编）。
    _cache_probe = tmp / "cacheprobe"
    _cache_probe.mkdir(parents=True, exist_ok=True)
    (tmp / "cacheprobe_vao").write_text("block", encoding="utf-8")   # 同名文件占位
    import io as _io2, contextlib as _cl, json as _json2
    _spec_f = tmp / "cacheprobe.py"
    _spec_f.write_text("SPEC = " + _json2.dumps(
        {"canvas": {"width": 320, "height": 180},
         "theme": {"colors": {"background": "#FFFFFF", "ink": "#111111"}},
         "slides": [{"id": "s1", "elements": [
             {"type": "text", "id": "t", "x": 10, "y": 10, "width": 200, "height": 40,
              "size": 18, "text": "x", "color": "ink"}]}]}), encoding="utf-8")
    from vao import run_check as _run_check
    _err = _io2.StringIO()
    with _cl.redirect_stderr(_err), _cl.redirect_stdout(_io2.StringIO()):
        # 输出名决定缓存目录名（output.stem + "_vao"）；占位文件必须正对那一个。
        _res, _ = _run_check(str(_spec_f), str(tmp / "cacheprobe.pptx"),
                             mode="draft", json_output=True, deadline=5)
    check("cache: 缓存写盘失败出声（不把缓存故障伪装成冷启动）",
          bool((_res.get("performance") or {}).get("cache_write_error"))
          and "warning" in _err.getvalue(),
          str((_res.get("performance") or {}).get("cache_write_error"))[:80])
    # 基准 CLI 是质量门，不能打印 FAIL 后仍退 0。
    bad_cases = tmp / "bad-cases.json"
    bad_cases.write_text(json.dumps({"cases": [{"id": "fail", "category": "test",
        "severity": "critical", "index": 0, "total": 1, "title": "话题",
        "expect": {"evidence": ["impossible"]}}]}), encoding="utf-8")
    proc = subprocess.run([sys.executable, str(SCRIPTS / "bench.py"), "--cases", str(bad_cases)],
                          capture_output=True, encoding="utf-8")
    check("bench: 判断失败必须非零退出", proc.returncode == 1 and "FAIL" in proc.stdout,
          f"returncode={proc.returncode}; stdout={proc.stdout[-250:]}; stderr={proc.stderr[-250:]}")
    # 真实 Windows 回归：英文 runner 的 ANSI 管道常为 cp1252；命令行须仍输出
    # UTF-8 的 FAIL 详情并保持非零退出，不能只因 UnicodeEncodeError 退出 1。
    legacy_env = {**os.environ, "PYTHONIOENCODING": "cp1252"}
    legacy = subprocess.run([sys.executable, str(SCRIPTS / "bench.py"), "--cases", str(bad_cases)],
                            capture_output=True, encoding="utf-8", errors="replace", env=legacy_env)
    check("bench: ANSI 管道下保留负例与 UTF-8 诊断",
          legacy.returncode == 1 and "FAIL" in legacy.stdout
          and "impossible" in legacy.stdout and "UnicodeEncodeError" not in legacy.stderr,
          f"returncode={legacy.returncode}; stdout={legacy.stdout[-200:]}; stderr={legacy.stderr[-300:]}")
    # 保持 size/mtime 的外部写入不能绕过同进程 JSON 与摘要缓存。
    f = tmp / "identity.json"
    f.write_text('{"v": 1}', encoding="utf-8")
    first = file_digest(f); json_read_cached(f); st = f.stat()
    f.write_text('{"v": 2}', encoding="utf-8")
    os.utime(f, ns=(st.st_atime_ns, st.st_mtime_ns))
    check("identity: 同尺寸恢复 mtime 的覆盖仍失效",
          file_digest(f) != first and json_read_cached(f) == {"v": 2})
    st = f.stat(); first = file_digest(f)
    other = tmp / "replacement.json"; other.write_text('{"v": 3}', encoding="utf-8")
    os.utime(other, ns=(st.st_atime_ns, st.st_mtime_ns)); os.replace(other, f)
    check("identity: 同尺寸同 mtime 的原子替换仍失效",
          file_digest(f) != first and json_read_cached(f) == {"v": 3})
    # 同图跨页与不同 QC 区域：一个事实一个主条目，不让重复 ID 阻断整稿。
    import assets
    from PIL import Image, ImageDraw
    image = tmp / "reuse.png"
    pic = Image.new("RGB", (1280, 720), (240, 236, 228))
    ImageDraw.Draw(pic).ellipse((700, 200, 1000, 500), fill=(120, 90, 60))
    pic.save(image)
    common = dict(title="图像绑定", content="原创测试插图，不是证词", asset="required",
                  asset_role="background", asset_source=dict(kind="original", path=str(image), source="synthetic fixture"))
    brief = dict(BRIEF, slides=[dict(common, id="a"), dict(common, id="b")])
    bundle = intelligence.think(brief)
    manifest = assets.build_manifest(brief, bundle)
    entries = assets.asset_entries(manifest)
    check("asset: 同图同 QC 契约合并主条目和页绑定",
          len(entries) == 1 and entries[0]["slide_ids"] == ["a", "b"])
    brief["slides"][1]["safe_area"] = dict(x=0.5, y=0.1, width=0.3, height=0.7)
    split = assets.asset_entries(assets.build_manifest(brief, intelligence.think(brief)))
    check("asset: 同图不同安全区不合并 QC", len(split) == 2 and split[0]["asset_id"] != split[1]["asset_id"])
    # 完整入口复验两页复用：真实像素 QC / 绑定 / 编译 / 全页取证都必须通过。
    brief["slides"][1].pop("safe_area")
    brief_path = tmp / "reuse-brief.json"
    brief_path.write_text(json.dumps(brief), encoding="utf-8")
    plan_path, manifest_path = tmp / "reuse-plan.json", tmp / "reuse-assets.json"
    plan_run = run_vao("plan", str(brief_path), "--out", str(plan_path), "--assets-out", str(manifest_path), "--json")
    planned = json.loads(manifest_path.read_text(encoding="utf-8"))
    aid = planned["assets"][0]["asset_id"]
    deck = dict(canvas=dict(width=1280, height=720), theme={"colors": {"background": "#FFFFFF"}}, slides=[dict(id=sid,
        elements=[dict(id="image", type="image", asset_id=aid, x=80, y=120, width=640, height=360, fit="cover")]) for sid in ("a", "b")])
    build = tmp / "reuse-spec.json"; build.write_text(json.dumps(deck), encoding="utf-8")
    production = run_vao("check", str(build), str(tmp / "reuse.pptx"), "--assets-manifest", str(manifest_path), "--mode", "release", "--speed", "strict", "--json")
    check("asset: 两页同图复用端到端 release strict 通过", plan_run.returncode == 0 and production.returncode == 0, production.stdout[-500:])
    import primitives as font_facts
    saved = (font_facts._PREVIEW_FONT_STAMP, dict(g._FONT_CACHE), dict(g._FONT_FALLBACKS), g._FONT_IDENTITY)
    try:
        # mock 字体对象而非依赖宿主是否装 Noto：无候选/坏候选均精确计数。
        with patch.object(Path, "exists", return_value=False), \
                patch.object(g.ImageFont, "load_default", wraps=g.ImageFont.load_default) as default:
            g._FONT_CACHE.clear(); g._FONT_FALLBACKS.clear()
            g._font(20, cjk=True); g._font(20, cjk=True)
            g._font(28, cjk=True); g._font(20, cjk=True, bold=True)
            check("font: 回退按字号字重缓存项累计，命中不重复记账",
                  default.call_count == 3 and g.font_fallbacks() == {"cjk": 3})
        with patch.object(Path, "exists", return_value=True), \
                patch.object(g.ImageFont, "truetype", side_effect=OSError("bad font")), \
                patch.object(g.ImageFont, "load_default", return_value=object()) as default:
            g._FONT_CACHE.clear(); g._FONT_FALLBACKS.clear()
            g._font(20, cjk=True, bold=True)
            check("font: 所有候选损坏仍回退并记账", default.call_count == 1 and g.font_fallbacks() == {"cjk": 1})
        build = make_build(tmp, chart=False)
        from vao import load_spec, _ghost
        spec = load_spec(build)[0]
        target = tmp / "fonts-preview"
        initial = _ghost(spec, target, speed="strict", output_sha="test")
        warm = _ghost(spec, target, speed="strict", output_sha="test")
        code = (f"import sys; sys.path.insert(0, {str(SCRIPTS)!r}); "
                "import primitives as f; "
                f"f.PREVIEW_FONT_CANDIDATES.update({g._FONT_CANDIDATES!r}); "
                f"f.PREVIEW_FONT_BOLD_CANDIDATES.update({g._FONT_BOLD_CANDIDATES!r}); "
                "from vao import load_spec, _ghost; "
                f"r = _ghost(load_spec({str(build)!r})[0], {str(target)!r}, speed='strict', output_sha='test'); "
                "assert r['reused']; assert not {'ghost', 'PIL', 'compiler', 'pptx'} & set(sys.modules)")
        probe = subprocess.run([sys.executable, "-c", code], capture_output=True, encoding="utf-8")
        check("font: 暖预览仍核对字体身份且零 PIL/ghost/Office 导入", probe.returncode == 0, probe.stderr[-300:])
        with patch.dict(g._FONT_BOLD_CANDIDATES, {"cjk_serif": (str(tmp / "changed-bold.ttf"),)}):
            changed = _ghost(spec, target, speed="strict", output_sha="test")
        check("font: 未变整套预览复用，粗体候选变化使整套和页缓存失效",
              warm["reused"] and not changed["reused"]
              and changed["fonts"] != initial["fonts"] and changed["pages_drawn"] == 5)
    finally:
        font_facts._PREVIEW_FONT_STAMP = saved[0]
        g._FONT_IDENTITY = saved[3]
        g._FONT_CACHE.clear(); g._FONT_CACHE.update(saved[1])
        g._FONT_FALLBACKS.clear(); g._FONT_FALLBACKS.update(saved[2])


def test_delivery_safety(tmp: Path):
    """真实故障的负例；路径冲突一律在写任何修复包/清单之前拒绝。"""
    import copy
    import io
    import os
    import contextlib
    from unittest.mock import patch
    import vao
    import compiler
    import ghost
    from intelligence import think
    from verify import check_spec
    from primitives import json_write
    from pptx import Presentation

    build = make_build(tmp, chart=False)
    spec, _ = vao.load_spec(build)

    def rejects(call):
        try:
            call()
        except (ValueError, FileExistsError):
            return True
        return False

    with contextlib.redirect_stdout(io.StringIO()) as captured:
        original = build.read_bytes()
        rejected = rejects(lambda: vao.run_check(str(build), str(build), json_output=True))
        check("safety: output 不能覆盖 build 输入", rejected and build.read_bytes() == original)
        build.write_bytes(original)
        rejected = rejects(lambda: vao.run_check(str(build), str(tmp / 'safe.pptx'),
                                                packet=str(build), json_output=True))
        check("safety: packet 不能覆盖 build 输入", rejected and build.read_bytes() == original)
        build.write_bytes(original)
        linked = tmp / 'build-hardlink.py'
        os.link(build, linked)
        rejected = rejects(lambda: vao.run_check(str(build), str(linked), json_output=True))
        check("safety: 硬链接别名也不得覆盖输入", rejected and build.read_bytes() == original)
        linked.unlink()
        build.write_bytes(original)
        brief = tmp / 'brief.json'
        json_write(brief, BRIEF)
        original_brief = brief.read_bytes()
        rejected = rejects(lambda: vao.run_plan(str(brief), str(brief)))
        check("safety: plan 不能覆盖 brief", rejected and brief.read_bytes() == original_brief)
        brief.write_bytes(original_brief)
        plan = tmp / 'plan.json'
        rejected = rejects(lambda: vao.run_plan(str(brief), str(plan), skeleton=str(plan)))
        check("safety: plan 与 skeleton 不能写到同一路径", rejected and not plan.exists())
        # 持久 QC 的派生路径也属于写集合，不能撞到 spec。
        man = tmp / 'asset_manifest.json'
        json_write(man, {"schema": "vao-assets-v2", "assets": []})
        qbuild = tmp / 'asset_manifest.qc.json'
        qspec = copy.deepcopy(spec)
        qspec['slides'][0]['elements'].append(
            {"type": "image", "id": "im", "src": "missing.png", "asset_id": "x",
             "x": 0, "y": 0, "width": 10, "height": 10})
        json_write(qbuild, qspec)
        before = qbuild.read_bytes()
        rejected = rejects(lambda: vao.run_check(str(qbuild), str(tmp / 'image.pptx'),
                                                assets_manifest=str(man), json_output=True))
        check("safety: QC 派生报告不能覆盖 spec", rejected and qbuild.read_bytes() == before)

    print("\n".join(line for line in captured.getvalue().splitlines() if line.startswith("[")))

    # 混合坏项以前从 drawable 列表里消失，spec 档会错误通过。
    bad = copy.deepcopy(spec)
    bad['slides'][0]['elements'].append(42)
    report = check_spec(bad)
    check("guard: elements 中的非对象项必须阻断", not report['passed'])
    for value in (0, -1, float('nan'), float('inf'), True, 'bad'):
        bad = copy.deepcopy(spec)
        bad['canvas']['width'] = value
        check(f"guard: 显式非法 canvas.width={value!r} 必须阻断", not check_spec(bad)['passed'])

    duplicate = copy.deepcopy(BRIEF)
    duplicate['slides'][1]['id'] = duplicate['slides'][0]['id']
    check("plan: 重复页面 id 在判断/资产映射前拒绝", rejects(lambda: think(duplicate)))

    kpi = {"type": "chart", "chart_kind": "kpi", "id": "zero", "value": 0,
           "label": "零事故", "x": 48, "y": 200, "width": 500, "height": 240}
    ks = copy.deepcopy(spec)
    ks['slides'] = [{"id": "s1", "elements": [kpi]}]
    check("guard: 0 是合法 KPI，不是缺失值", check_spec(ks)['passed'])
    kp = tmp / 'zero.pptx'
    kr = compiler.compile_deck(ks, kp)
    text = {s.name: s.text for s in Presentation(kp).slides[0].shapes if s.has_text_frame}
    check("compiler: 0 KPI 原样保留为可编辑文字", kr['passed'] and text.get('zero__value') == '0')
    with patch.object(ghost.ImageDraw.ImageDraw, 'text', autospec=True) as draw:
        ghost.ghost_page(ks['slides'][0], ks, scale=0.5)
    check("ghost: 0 KPI 在结构预览中不消失", any('0' == str(c.args[2]) for c in draw.call_args_list))
    for kind in ('pie', 'donut', 'donut_composition'):
        bad = copy.deepcopy(ks)
        bad['slides'][0]['elements'][0].update(chart_kind=kind,
            data=[{"label": "正", "value": 10}, {"label": "负", "value": -3}])
        check(f"guard: {kind} 不接受无定义的负构成", not check_spec(bad)['passed'])

    # 已有 PASS 后改成含图无清单：早退必须撤销旧发布资格。
    with contextlib.redirect_stdout(io.StringIO()) as captured:
        out = tmp / 'stale.pptx'
        _, code = vao.run_check(str(build), str(out), mode='release', speed='strict',
                                json_output=True)
        check("release: stale 负例的健康前置条件", code == 0)
        broken = copy.deepcopy(spec)
        broken['slides'][0]['elements'].append({"type": "image", "id": "im",
            "src": "absent.png", "x": 0, "y": 0, "width": 10, "height": 10})
        js = tmp / 'broken.json'
        json_write(js, broken)
        _, code = vao.run_check(str(js), str(out), mode='release', json_output=True)
        manifest = json.loads(out.with_suffix('.manifest.json').read_text(encoding='utf-8'))
        check("release: 资产早退必须替换旧 PASS 清单", code == 2
              and manifest['status'] == 'BLOCKED' and not manifest['release_eligible'])

    print("\n".join(line for line in captured.getvalue().splitlines() if line.startswith("[")))

    # 编译/后处理失败不能破坏上一次的完整包；固定 .tmp 不得成为符号链接写入面。
    target = tmp / 'atomic.pptx'
    compiler.compile_deck(spec, target)
    healthy = target.read_bytes()
    failed = False
    try:
        with patch.object(compiler, '_postprocess_package', side_effect=OSError('disk full')):
            compiler.compile_deck(spec, target)
    except OSError:
        failed = True
    check("compiler: 后处理失败保留旧完整产物", failed and target.read_bytes() == healthy)
    victim = tmp / 'victim.txt'
    victim.write_text('keep', encoding='utf-8')
    trap = target.with_suffix('.pptx.tmp')
    # Windows CI 未授权 symlink 时仍测试同名普通文件不能被覆盖或删除。
    try:
        trap.symlink_to(victim)
    except OSError:
        trap.write_text('keep', encoding='utf-8')
    compiler._postprocess_package(target)
    check("compiler: 不触碰固定临时路径（含符号链接）", victim.read_bytes() == b'keep'
          and trap.exists() and trap.read_bytes() == b'keep')
    check("compiler: 成败路径均清理唯一临时文件", not list(tmp.glob('.atomic.pptx-*')))

    # 快档也不能把 size/mtime 当 QC 字节证明：等长 BMP 换成黑图并恢复 mtime。
    import assets
    from PIL import Image
    image = tmp / 'same-stat.bmp'
    Image.new('RGB', (320, 180), (240, 240, 240)).save(image)
    mpath = tmp / 'same-stat.json'
    entry = {"asset_id": "same", "decision": "existing", "slide_ids": ["s1"],
             "origin": {"kind": "original", "path": str(image), "source": "test"},
             "safe_area": {}, "background_color": "#FFFFFF", "retry_budget": 0}
    source_brief = tmp / "source-brief.json"
    json_write(source_brief, {"slides": [{"id": "s1"}]})
    source_workflow = {"schema": "vao-asset-chain-v2", "brief_path": str(source_brief),
                       "brief_file_sha256": assets.file_digest(source_brief)}
    json_write(mpath, {"schema": "vao-assets-v2",
                      "workflow": source_workflow, "assets": [entry]})
    good, _ = assets.qc_report(mpath, speed='fast', phase='release')
    check("assets: 同 stat 换图负例的健康前置条件", good['status'] == 'PASS')
    stat = image.stat()
    Image.new('RGB', (320, 180), (0, 0, 0)).save(image)
    os.utime(image, ns=(stat.st_atime_ns, stat.st_mtime_ns))
    shots, hashes = {}, {}
    changed, _ = assets.qc_report(mpath, speed='fast', phase='release', snapshots=shots, digests=hashes)
    check("assets: fast 同 size/mtime 换图重测并阻断", changed['status'] == 'BLOCKED'
          and not changed['results'][0]['qc'].get('reused'))
    check("assets: QC 暖路径也交接当轮字节与真实 SHA", shots.get(str(image.resolve())) == image.read_bytes()
          and hashes.get(str(image.resolve())) == assets.file_digest(image))
    # 同一文件不同 QC 契约只读取一次，不把不同契约错误合并。
    both = copy.deepcopy(entry)
    both.update(asset_id='second', safe_area={"x": 0, "y": 0, "width": 0.2, "height": 0.5})
    json_write(mpath, {"schema": "vao-assets-v2", "workflow": source_workflow,
                      "assets": [entry, both]})
    reader = Path.read_bytes
    calls = []
    def track_read(path):
        if path.resolve() == image.resolve():
            calls.append(path)
        return reader(path)
    with patch.object(Path, 'read_bytes', track_read):
        assets.qc_report(mpath, speed='strict', persist=False)
    check("assets: 同源不同 QC 契约共享单轮读图", len(calls) == 1, str(len(calls)))


    # 路径检查需覆盖资产源，不止 CLI 的显式输入文件。
    with contextlib.redirect_stdout(io.StringIO()) as captured:
        js = tmp / 'asset-source.json'
        source_spec = copy.deepcopy(spec)
        source_spec['slides'] = [{"id": "s1", "elements": [{"type": "image", "id": "i",
            "asset_id": "same", "x": 0, "y": 0, "width": 100, "height": 100}]}]
        json_write(js, source_spec)
        before = image.read_bytes()
        rejected = rejects(lambda: vao.run_check(str(js), str(image), assets_manifest=str(mpath),
                                                json_output=True))
        check("safety: 输出不能覆盖清单绑定的原图", rejected and image.read_bytes() == before)
        b = copy.deepcopy(BRIEF)
        b['slides'][0]['asset_source'] = entry['origin']
        bf = tmp / 'with-source.json'
        json_write(bf, b)
        rejected = rejects(lambda: vao.run_plan(str(bf), str(image)))
        check("safety: plan 不能覆盖 brief 引用的原图", rejected and image.read_bytes() == before)
    print("\n".join(line for line in captured.getvalue().splitlines() if line.startswith("[")))

    import bench
    baseline = {"cases": [{"id": "passed", "passed": True}, {"id": "failed", "passed": False}]}
    reg = bench.regression({"cases": []}, baseline)
    check("bench: 缺失案例不是修复、也不能免除回归", reg['regressed_cases'] == ['passed']
          and reg['fixed_cases'] == [] and reg.get('missing_cases') == ['failed', 'passed'])
    cases = json.loads((SCRIPTS / 'json' / 'judgment_cases.json').read_text(encoding='utf-8'))
    cases['cases'] = cases['cases'][:1]
    cp = tmp / 'reduced-cases.json'
    json_write(cp, cases)
    bp = tmp / 'baseline.json'
    full = bench.run(cp)
    full['cases'].append({"id": "removed", "passed": True})
    json_write(bp, full)
    probe = subprocess.run([sys.executable, str(SCRIPTS / 'bench.py'), '--cases', str(cp),
                            '--baseline', str(bp)], capture_output=True, encoding='utf-8')
    check("bench: baseline 案例消失时 CLI 非零", probe.returncode == 1)


    module = tmp / 'dataclass-build.py'
    module.write_text("from __future__ import annotations\nfrom dataclasses import dataclass\n"
                      "@dataclass\nclass Label:\n    text: str\n"
                      "SPEC = {'slides': [], 'label': Label('first').text}\n", encoding='utf-8')
    ok = False
    try:
        loaded_module, _ = vao._load_module(module, name='vao_dataclass_probe')
        ok = loaded_module.SPEC.get('label') == 'first'
    except Exception:
        pass
    finally:
        sys.modules.pop('vao_dataclass_probe', None)
    check("loader: 可信 build 支持 dataclass 与延迟注解", ok)
    module.write_text("SPEC = {'slides': [], 'label': 'second'}\n", encoding='utf-8')
    loaded, _ = vao.load_spec(module)
    check("loader: 连续加载替换模块身份，不保留旧 SPEC", sys.modules['vao_build'].SPEC == loaded)
    prior = sys.modules['vao_build']
    module.write_text("raise RuntimeError('author failure')\n", encoding='utf-8')
    try:
        vao.load_spec(module)
    except RuntimeError:
        pass
    check("loader: 作者模块异常回滚注册状态", sys.modules.get('vao_build') is prior)

    broken_png = tmp / 'truncated.png'
    Image.effect_noise((128, 128), 60).convert('RGB').save(broken_png)
    raw = broken_png.read_bytes()
    broken_png.write_bytes(raw[:len(raw) // 2])
    result = None
    try:
        result = assets.image_qc(str(broken_png))
    except OSError:
        pass
    check("assets: 已识别头部但解码失败的图片返回可诊断 error", isinstance(result, dict)
          and result.get('status') == 'error' and bool(result.get('issue')))


    # 同一生成视觉需求不能吞掉第二页的压字/底色/角色契约。
    b = {"audience": "评审", "decision": "发布", "tension": "可读性", "slides": []}
    for i, safe in enumerate(({}, {"x": 0.1, "y": 0.1, "width": 0.3, "height": 0.4})):
        b['slides'].append({"id": f's{i}', "title": "同一图像", "content": "保留证据",
                            "asset": "required", "asset_subject": "soft paper",
                            "asset_function": "context", "asset_role": "background", "safe_area": safe})
    check("assets: 生成去重不得静默吞掉跨页 QC 契约", rejects(lambda: assets.build_manifest(b, think(b))))


    from intelligence import load_brief
    module.write_text("from __future__ import annotations\nfrom dataclasses import dataclass\n"
                      "@dataclass\nclass Label:\n    text: str\n"
                      "BRIEF = {'subject': Label('brief').text, 'slides': []}\n", encoding='utf-8')
    loaded_brief = None
    try:
        loaded_brief = load_brief(module)
    except Exception:
        pass
    check("loader: Python brief 与 build 共用安全模块注册语义", loaded_brief is not None
          and loaded_brief.get('subject') == 'brief')


    for kind in ('pie', 'donut', 'donut_composition'):
        for values in ([10, -3], [0, 0]):
            bad = copy.deepcopy(ks)
            bad['slides'][0]['elements'][0].update(chart_kind=kind, categories=['A', 'B'],
                                                 series=[{'name': 'one', 'values': values}])
            check(f"guard: {kind} series 载荷也拒绝负构成/全零 {values}", not check_spec(bad)['passed'])
    bad['slides'][0]['elements'][0].update(chart_kind='pie',
        series=[{'name': 'one', 'values': [10, 20]}, {'name': 'two', 'values': [20, 10]}])
    check("guard: pie 不得悄悄隐藏第二条序列", not check_spec(bad)['passed'])



def test_material_direction(tmp: Path):
    """锁定作者判断的执行保真，不评分素材美感或增加自动配图。"""
    import copy
    import assets as a
    import intelligence as intel

    def planned(**overrides):
        slide = {"id": "diagram", "title": "剖面说明传动关系", "content": "输入轴驱动输出机构。",
                 "role": "explain", "asset": "required", "asset_role": "illustration",
                 "asset_subject": "a mechanical cutaway showing an input shaft and an output linkage",
                 "safe_area": {}, **overrides}
        brief = {"audience": "技术评审", "decision": "理解机构关系", "tension": "不清楚传动关系",
                 "slides": [slide]}
        bundle = intel.think(brief)
        bundle["deck"]["world"].update(
            visual_world="chinese ink-wash stage with paper grain",
            light="warm theatrical spotlights from the rear", texture_keys=["paper", "stage"])
        return brief, bundle

    def card_for(brief, bundle):
        return a.asset_card(bundle["pages"][0], brief, bundle["deck"], "material-test")

    brief, bundle = planned()
    card, page = card_for(brief, bundle)
    result = a.build_asset_prompt(card, page)
    check("material: 未声明介质的对象不被强制为 3D/软材质/透明剪影",
          not any(x in result["prompt"] for x in
                  ("3D minimal", "soft material", "pure transparent background")))
    check("material: 对象比例句不再误称背景",
          "16:9 presentation illustration" in result["prompt"]
          and "presentation background" not in result["prompt"])
    check("material: 对象不继承世界纸纹/舞台灯光/融合，品牌色仍保留",
          card["material_source"] == "fallback" and card["lighting_source"] == "fallback"
          and not card["texture"] and not card["fusion"] and bool(card["color"])
          and not any(x in result["prompt"] for x in ("paper grain", "spotlights", "melts into")))

    icon = {**card, "asset_type": "icon", "asset_role": "", "medium": "solid pictogram",
            "subject": ["one solid lock symbol"], "style": ["filled geometric symbol"],
            "lighting": ["flat graphic rendering without shading"], "lighting_source": "declared"}
    icon.pop("resolved", None)
    icon_result = a.build_asset_prompt(icon, page, ratio="1:1")
    check("material: 图标保留实心声明，不强加细线/空心/透明底",
          "filled geometric symbol" in icon_result["prompt"]
          and "1:1 presentation icon" in icon_result["prompt"]
          and not any(x in icon_result["prompt"] for x in
                      ("thin stroke", "no fill", "single line icon", "pure transparent background"))
          and icon_result["meta"]["fusion_enabled"] is False)

    reason = "看清不可从外观照片辨认的内部传动；不充当实测证据。"
    composition = "tight cutaway at right; retain the full linkage and its connection to the input shaft"
    textures = ["fine hatching on the cut face", "stippling on the bearing", "smooth line on the outer shell"]
    brief, bundle = planned(asset_reason=reason, asset_composition=composition,
                            medium="scientific cutaway illustration", material="precise schematic linework",
                            lighting="uniform flat studio lighting without directional shading",
                            texture=textures, energy="high", negative=["decorative dividers"])
    card, page = card_for(brief, bundle)
    result = a.build_asset_prompt(card, page)
    check("material: 素材理由进入判断与骨架而不进入出图主体",
          bundle["pages"][0]["judgment"]["media"]["necessity"] == reason
          and reason in intel.build_skeleton(bundle) and reason not in result["prompt"])
    check("material: 素材内部构图原样到提示词且不改 QC 几何",
          card["composition"] == [composition] and composition in result["prompt"]
          and page["safe_area"] == {})
    check("material: 作者三条纹理均抵达实际提示词，不被截为一两条",
          all(x in result["prompt"] for x in textures))
    check("material: 作者光照不被高能量/舞台光接管，禁项实际透传",
          result["prompt"].count("uniform flat studio lighting without directional shading") == 1
          and card["resolved"]["light"]["from"] == "declared"
          and not card["resolved"]["light"]["energy_applied"]
          and "no decorative dividers" in result["negative"])
    changed = copy.deepcopy(bundle)
    changed["pages"][0]["declarations"]["asset_composition"] = "wide view preserving the entire mechanism"
    changed_card, changed_page = card_for(brief, changed)
    check("material: 资产构图改变进入视觉身份，避免复用旧构图提示词",
          a.asset_fingerprint(card, page) != a.asset_fingerprint(changed_card, changed_page))

    brief, bundle = planned(asset_role="background", medium="photography", texture=[],
                            fusion_enabled=False, light_direction="right")
    card, page = card_for(brief, bundle)
    result = a.build_asset_prompt(card, page)
    check("material: 背景仍继承空间/光，显式空纹理和关闭融合可生效",
          card["material_source"] == "world" and card["lighting_source"] == "world"
          and not card["texture"] and not card["fusion"]
          and result["meta"]["fusion_enabled"] is False and page["light_direction"] == "right")
    check("material: 摄影尊重世界光，不强加自然单光/胶片/空中透视",
          a.photo_gate_active(card) and card["resolved"]["light"]["from"] == "world"
          and "warm theatrical spotlights from the rear" in result["prompt"]
          and not any(x in result["prompt"] for x in
                      ("single natural light", "medium format film", "soft micro grain", "faint air between planes")))
    photo = {**card, "asset_type": "illustration", "medium": "photography",
             "subject": ["a photographic reproduction of a chinese ink-wash painting"],
             "lighting": ["uniform flat studio lighting"], "lighting_source": "declared"}
    photo.pop("resolved", None)
    photo_result = a.build_asset_prompt(photo, page)
    check("material: 摄影对象中的水墨题材不会反向接管显式摄影介质",
          not a.ink_gate_active(photo) and a.photo_gate_active(photo)
          and "uniform flat studio lighting" in photo_result["prompt"]
          and "no uniform flat studio lighting" not in photo_result["negative"]
          and "ink-wash marks and pigment behavior" not in photo_result["prompt"])
    ink = {**photo, "medium": "chinese ink-wash painting", "composition": ["centered subject"]}
    ink_result = a.build_asset_prompt(ink, page)
    check("material: 水墨不强加半幅留白/固定色阶或禁止居中",
          a.ink_gate_active(ink) and not a.photo_gate_active(ink)
          and not any(x in ink_result["prompt"] for x in
                      ("more than half", "three to four values"))
          and "no symmetric centered composition" not in ink_result["negative"])
    people = {**photo, "subject": ["people shaking hands with a humanoid robot under neon glow"]}
    people_result = a.build_asset_prompt(people, page)
    check("material: 不按人物词或风格名附加与主体矛盾的禁项",
          not any(x in people_result["negative"] for x in
                  ("no humanoid robot", "no corporate handshake", "no neon glow", "no cyberpunk")))

    no_image = intel.think({"slides": [{"id": "text", "title": "关系由文字说明",
                                       "asset_reason": "暂存设计考虑，不是出图指令"}]})
    none = intel.think({"slides": [{"id": "none", "title": "直接陈述", "asset": "none",
                                   "asset_subject": "a product photograph"}]})
    check("material: 仅写理由不触发出图，显式 none 仍覆盖已有主体",
          no_image["assets_hint"]["generate"] == [] and none["assets_hint"]["generate"] == [])
    u = intel.understand("输入轴驱动输出机构")
    check("material: 图表与必要对象可由作者协调共存，不自动加图",
          intel.media_necessity(u, "explain", has_chart=True)["decision"] == "none"
          and intel.media_necessity(u, "explain", has_chart=True,
                                   declared="required")["decision"] == "required")

    brief, bundle = planned()
    fresh = a.build_manifest(brief, bundle)
    entry = next(x for x in fresh["assets"] if x["decision"] == "generate")
    cache = tmp / "old-prompt-cache.json"
    cache.write_text(json.dumps({"schema": "vao-asset-prompt-cache-v2", "entries": {
        entry["fingerprint"]: {"prompt": "outdated forced 3D minimal illustration",
                               "negative": "no flat graphics", "meta": {}}}}), encoding="utf-8")
    rebuilt = a.build_manifest(brief, bundle, cache_path=cache)
    warm = a.build_manifest(brief, bundle, cache_path=cache)
    check("material: 旧 v2 提示词缓存重建，v3 暖缓存保留确定性复用",
          rebuilt["prompt_cache"]["hits"] == 0 and warm["prompt_cache"]["hits"] == 1
          and json.loads(cache.read_text(encoding="utf-8"))["schema"] == "vao-asset-prompt-cache-v3"
          and rebuilt["assets"][0]["prompt"] == warm["assets"][0]["prompt"]
          and "outdated forced" not in rebuilt["assets"][0]["prompt"])


def _portable_relpath(target: Path, base: Path) -> str:
    """尽量给出相对路径，做不到就退回绝对路径。

    `os.path.relpath` 在 Windows **跨盘符**时无解——CI 上工作区在 `D:`、临时目录在
    `C:`，于是 `ValueError: path is on mount 'C:', start on mount 'D:'` 直接掀掉整个
    回归套件（本轮审计拿到的 windows-latest 失败日志即此）。需要「相对优先」的地方
    统一走这里，别让测试在地形上崩掉而不是在断言上失败。
    """
    try:
        return os.path.relpath(target, base)
    except ValueError:                      # Windows：两个路径没有共同根
        return str(target)


def test_win_portability(tmp: Path):
    """平台地形不该掀掉回归套件：同盘相对 vs 跨盘绝对必须分开验。

    背景（本轮拿到的 CI 日志）：`selftest · windows-latest · py3.12` 在
    `test_deep_release` 里 `os.path.relpath(blind_src, Path.cwd())` 抛
    `ValueError: path is on mount 'C:', start on mount 'D:'` —— 工作区在 D:、
    临时目录在 C:，两棵根之间没有相对路径，测试直接崩在 setup 而不是断言上。
    同一台机器上没有 POSIX 权限位——那条判据在 test_atomic_publish（覆盖前后
    权限不变，数值相等只在 POSIX 追加），本测试只验路径分支，都在本机以
    **平台无关的方式**复现，不需要 Windows。
    """
    import ntpath
    from unittest.mock import patch

    # 1) 复现 CI 的确切条件：ntpath 在跨盘符时就是抛 ValueError（本机可验证）
    message = ""
    cross_drive = False
    try:
        ntpath.relpath(r"C:\Users\runner\AppData\Local\Temp\blind-src", r"D:\a\repo")
    except ValueError as exc:
        cross_drive = isinstance(exc, ValueError)
        message = str(exc)
    check("win: 跨盘符 relpath 确实无解（CI 崩溃的真实条件，本机用 ntpath 复现）",
          cross_drive and "mount" in message, message[:80])

    # 2) 同盘（同根）对走相对路径分支。前提修正：不能拿 tmp 与 Path.cwd() 配对
    #    代表「同盘」——windows-latest 的 tmp 在 C:、工作区在 D:，那对本身就是跨盘，
    #    恰好落进兜底分支，断言与前提相反。同盘对必须两头同根：Windows 语义用
    #    ntpath 直接验，POSIX 语义用同一棵树里的目录对验。
    work = tmp / "winrt"
    work.mkdir()
    check("win: 同盘符 relpath 保持相对（Windows 语义，ntpath 同盘对）",
          ntpath.relpath(r"D:\a\repo\review\blind-src", r"D:\a\repo") == "review\\blind-src")
    relative_ok = _portable_relpath(work, tmp)
    check("win: 同盘符仍给出相对路径（不因兜底丢掉本来的覆盖）",
          not Path(relative_ok).is_absolute() and (tmp / relative_ok).resolve() == work.resolve(),
          relative_ok)

    # 3) 跨盘（mock 复现 CI 的 ValueError）时退回绝对路径而不是抛错
    with patch("os.path.relpath", side_effect=ValueError(
            "path is on mount 'C:', start on mount 'D:'")):
        fallback = _portable_relpath(work, Path.cwd())
    check("win: 跨盘符时退回绝对路径而不是抛错",
          Path(fallback).is_absolute() and fallback == str(work))

    # 4) blind_pack 必须接受这个绝对路径并原样记录来源标识（覆盖 CI 上会走的那条分支）
    import bench
    (work / "blind-src").mkdir()
    (work / "blind-src" / "evidence.txt").write_bytes(b"evidence")
    with patch("os.path.relpath", side_effect=ValueError("mount")):
        source = _portable_relpath(work / "blind-src", Path.cwd())
    pack = bench.blind_pack([source], work / "out", seed=3)
    key = json.loads(Path(pack["key"]).read_text(encoding="utf-8"))
    check("win: 跨盘符分支下 blind pack 仍可用且来源标识不被改写",
          (work / "out" / "V01" / "evidence.txt").read_bytes() == b"evidence"
          and key["items"] == {"V01": source} and key["seed"] == 3, str(key))


def test_deep_release(tmp: Path):
    """本轮已复现缺陷：辅助写入、未完成发布、数据载荷与像素表达。"""
    import contextlib
    import copy
    import io
    import os
    import subprocess
    import zipfile
    from unittest.mock import patch
    from PIL import Image
    import assets
    import compiler
    import ghost
    import vao
    import verify
    from primitives import json_write
    base = {"canvas": {"width": 1280, "height": 720},
            "theme": {"colors": {"background": "#FFFFFF", "ink": "#111111", "primary": "#123456",
                                  "secondary": "#789ABC", "accent": "#C05030", "muted": "#555555"}},
            "slides": [{"id": "one", "elements": [{"id": "t", "type": "text", "x": 64, "y": 80,
                        "width": 1000, "height": 100, "text": "Engineering fixture", "size": 32, "color": "ink"}]}]}

    def quiet(build, output, **kw):
        with contextlib.redirect_stdout(io.StringIO()):
            return vao.run_check(str(build), str(output), **kw)

    # 专用辅助目录不仅要保护固定 marker，还要覆盖页图、清理范围和链接目标。
    for label, kind in (("预览覆盖 PPTX", "output"), ("预览 marker 覆盖输入", "marker"),
                        ("编译缓存覆盖输入", "cache"), ("packet 覆盖预览页图", "packet")):
        d = tmp / kind; d.mkdir()
        build = d / "input.json"; output = d / "out.pptx"; kw = {}
        if kind == "cache": build = d / "out_vao" / "compile_cache.json"
        elif kind == "marker": build = d / "ghost.meta.json"; kw["preview"] = str(d)
        elif kind == "output": output = d / "ghost-01.png"; kw["preview"] = str(d)
        else:
            kw.update(preview=str(d / "preview"), packet=str(d / "preview" / "ghost-01.png"))
        json_write(build, base); before = build.read_bytes(); rejected = False
        try: quiet(build, output, mode="release", speed="strict", **kw)
        except ValueError: rejected = True
        check("deep: " + label + " 在写入前拒绝", rejected and build.read_bytes() == before and not output.exists())

    raw_image = tmp / "raw.png"; Image.new("RGB", (320, 180), "gray").save(raw_image)
    raw_spec = copy.deepcopy(base)
    raw_spec["slides"][0]["elements"] = [{"id": "i", "type": "image", "src": str(raw_image),
                                        "x": 64, "y": 80, "width": 200, "height": 100}]
    raw_build = tmp / "raw.json"; json_write(raw_build, raw_spec); before = raw_image.read_bytes()
    rejected = False
    try: quiet(raw_build, tmp / "raw.pptx", packet=str(raw_image), mode="release")
    except ValueError: rejected = True
    check("deep: 缺清单的失败诊断也不许覆盖 raw src", rejected and raw_image.read_bytes() == before)

    for link_kind in ("symlink", "hardlink"):
        preview = tmp / link_kind; preview.mkdir()
        victim = tmp / (link_kind + "-victim.txt"); victim.write_bytes(b"precious input")
        trap = preview / "ghost-01.png"
        try:
            trap.symlink_to(victim) if link_kind == "symlink" else os.link(victim, trap)
        except OSError:
            print(f"[SKIP] deep: 本平台不能创建 {link_kind}")
            continue
        build = tmp / (link_kind + ".json"); json_write(build, base); rejected = False
        try: quiet(build, tmp / (link_kind + ".pptx"), preview=str(preview), mode="release")
        except ValueError: rejected = True
        check(f"deep: 辅助目录 {link_kind} 不得覆盖外部文件", rejected and victim.read_bytes() == b"precious input")

    build = tmp / "state.json"; output = tmp / "state.pptx"; json_write(build, base)
    initial, rc = quiet(build, output, mode="release", speed="strict")
    check("deep: 发布负例的真实健康前置条件", rc == 0 and initial["release_eligible"])
    observed = []
    def interrupted_compile(*args, **kwargs):
        observed.append(json.loads(output.with_suffix(".manifest.json").read_text(encoding="utf-8"))["status"])
        raise RuntimeError("injected compile interruption")
    with patch.object(vao, "_compile_step", interrupted_compile):
        _, rc = quiet(build, output, mode="release", speed="strict")
    check("deep: 编译进行中已撤销旧 PASS", observed == ["BLOCKED"] and rc == 2)
    quiet(build, output, mode="release", speed="strict")
    interrupted = False
    with patch.object(vao, "_ghost", side_effect=KeyboardInterrupt()):
        try: quiet(build, output, mode="release", speed="strict")
        except KeyboardInterrupt: interrupted = True
    state = json.loads(output.with_suffix(".manifest.json").read_text(encoding="utf-8"))
    check("deep: KeyboardInterrupt 后无旧发布资格", interrupted and state["status"] == "BLOCKED" and not state["release_eligible"])

    for command, code in (("check", "raise SystemExit(0)"), ("plan", "raise SystemExit(0)"),
                           ("check", "raise KeyboardInterrupt()")):
        script = tmp / (command + ("-ctrlc" if "Keyboard" in code else "-exit") + ".py")
        script.write_text(code, encoding="utf-8")
        args = [command, str(script), str(output), "--mode", "release", "--json"] if command == "check" else [
            "plan", str(script), "--out", str(tmp / "abort-plan.json")]
        run = subprocess.run([sys.executable, str(Path(vao.__file__)), *args], capture_output=True, text=True,
                             encoding="utf-8")
        expected = 130 if "Keyboard" in code else 2
        check(f"deep: CLI {command} / {code} 不伪报成功", run.returncode == expected, str(run.returncode))

    import bench
    blind_src = tmp / "blind-src"; blind_src.mkdir(); (blind_src / "evidence.txt").write_bytes(b"evidence")
    rejected = False
    with patch("shutil.copytree", side_effect=RuntimeError("unsafe copier reached")):
        try: bench.blind_pack([str(blind_src)], blind_src / "review")
        except (ValueError, RuntimeError) as exc: rejected = isinstance(exc, ValueError)
    check("deep: blind pack 在复制前拒绝递归输出", rejected and not (blind_src / "review").exists())
    blind_output = tmp / "blind-output"; blind_key = tmp / "blind-output.key.json"
    blind_key.write_bytes(b"precious source"); original_key = blind_key.read_bytes()
    run = subprocess.run([sys.executable, str(Path(bench.__file__)), "blind", "pack", str(blind_output),
                          str(blind_key)], capture_output=True, text=True, encoding="utf-8")
    check("deep: blind pack 的 key 不得覆盖待评原件", run.returncode == 2 and blind_key.read_bytes() == original_key)
    relative_source = _portable_relpath(blind_src, Path.cwd())
    healthy_pack = bench.blind_pack([relative_source], tmp / "healthy-blind", seed=1)
    check("deep: blind pack 正常目录及评分表保持可用", (tmp / "healthy-blind/V01/evidence.txt").read_bytes() == b"evidence"
          and Path(healthy_pack["sheet"]).is_file()
          and json.loads(Path(healthy_pack["key"]).read_text(encoding="utf-8"))["items"] == {"V01": relative_source})

    chart = {"id": "c", "type": "chart", "chart_kind": "ranked_bar", "x": 64, "y": 80,
             "width": 1000, "height": 400, "data": [{"label": "Visible", "value": 1}],
             "categories": ["Hidden"], "series": [{"name": "Hidden series", "values": [9]}],
             "source": "fixture", "unit": "count", "period": "test", "basis": "test"}
    for kind in ("ranked_bar", "column"):
        spec = copy.deepcopy(base); spec["slides"][0]["elements"] = [dict(chart, chart_kind=kind)]
        check(f"deep: {kind} 不可静默选择一份双载荷", not verify.check_spec(spec)["passed"])
        report = compiler.compile_deck(spec, tmp / (kind + ".pptx"))
        check(f"deep: 独立编译 {kind} 双载荷也不能给 passed", not report["passed"])
    valid = copy.deepcopy(base); e = dict(chart, chart_kind="column"); e.pop("data")
    valid["slides"][0]["elements"] = [e]
    check("deep: 合法原生 categories + series 路径不退化", verify.check_spec(valid)["passed"]
          and compiler.compile_deck(valid, tmp / "valid-series.pptx")["passed"])

    brief = tmp / "brief.json"; json_write(brief, {"slides": [{"id": "one"}]})
    wf = {"schema": "vao-asset-chain-v2", "brief_path": str(brief), "brief_file_sha256": assets.file_digest(brief)}
    man = {"schema": "vao-assets-v2", "workflow": wf, "assets": []}
    check("deep: 完整资产需求凭证可验证", not assets.verify_sources(man))
    for label, changes in (("缺 SHA", {"brief_file_sha256": None}), ("坏 SHA", {"brief_file_sha256": "x" * 64}),
                            ("缺 brief", {"brief_path": str(tmp / "missing.json")}),
                            ("伪链版本", {"schema": "vao-asset-chain-invented"})):
        check("deep: 资产来源 " + label + " 不再静默通过", bool(assets.verify_sources({**man, "workflow": {**wf, **changes}})))

    # 像素暖复用不等于工作流状态不变：brief 改变时 QC 磁盘报告也须撤销 PASS。
    qc_image = tmp / "warm-qc.png"; Image.new("RGB", (320, 180), (240, 240, 240)).save(qc_image)
    qc_man = {**man, "assets": [{"asset_id": "warm", "slide_ids": ["one"], "decision": "existing",
        "origin": {"kind": "original", "path": str(qc_image), "source": "fixture"},
        "safe_area": {}, "asset_role": "background", "asset_role_source": "declared"}]}
    qc_manifest = tmp / "warm-manifest.json"; json_write(qc_manifest, qc_man)
    healthy, _ = assets.qc_report(qc_manifest, speed="strict", phase="release")
    original_brief = brief.read_bytes(); brief.write_bytes(original_brief + b" ")
    blocked, _ = assets.qc_report(qc_manifest, speed="strict", phase="release")
    persisted = json.loads(Path(blocked["report_path"]).read_text(encoding="utf-8"))
    check("deep: QC 像素暖命中但需求改变时磁盘撤销旧 PASS", healthy["status"] == "PASS"
          and blocked["status"] == "BLOCKED" and persisted["status"] == "BLOCKED")
    brief.write_bytes(original_brief)
    recovered, _ = assets.qc_report(qc_manifest, speed="strict", phase="release")
    check("deep: QC 需求恢复后的状态也同步落盘", recovered["status"] == "PASS"
          and json.loads(Path(recovered["report_path"]).read_text(encoding="utf-8"))["status"] == "PASS")

    alpha = tmp / "alpha.png"; image = Image.new("RGBA", (100, 100), (255, 0, 0, 0))
    image.paste((0, 255, 0, 255), (25, 25, 75, 75)); image.save(alpha)
    spec = copy.deepcopy(base); spec["slides"][0]["background"] = "#171310"
    spec["slides"][0]["elements"] = [{"id": "i", "type": "image", "x": 64, "y": 80,
                                     "width": 200, "height": 100, "src": str(alpha), "fit": "contain"}]
    out = tmp / "alpha.pptx"; report = compiler.compile_deck(spec, out)
    with zipfile.ZipFile(out) as z:
        embedded = next(n for n in z.namelist() if n.startswith("ppt/media/"))
        with Image.open(io.BytesIO(z.read(embedded))) as pixels:
            check("deep: 透明符号与 contain 留白不压平为 deck 白底", report["passed"]
                  and pixels.mode == "RGBA" and pixels.getpixel((0, 0))[3] == 0)

    split = Image.new("RGB", (100, 100), "red"); split.paste((0, 255, 0), (50, 0, 100, 100))
    source = tmp / "split.png"; split.save(source)
    spec["slides"][0]["elements"][0].update(src=str(source), crop=[0.5, 0, 0, 0], fit="cover")
    preview = ghost.ghost_page(spec["slides"][0], spec)
    check("deep: ghost 实际消费图片 crop", preview.getpixel((84, 100))[:3] == (0, 255, 0))
    second = copy.deepcopy(spec["slides"][0]); second["id"] = "two"
    second["elements"][0]["crop"] = [0, 0, 0.5, 0]; spec["slides"].append(second)
    rendered = []; ghost.ghost_deck(spec, tmp / "crops", scale=1, store=False, images_out=rendered)
    check("deep: 同源不同裁切不共用错误的预览适配缓存", rendered[0].getpixel((84, 100))[:3] == (0, 255, 0)
          and rendered[1].getpixel((84, 100))[:3] == (255, 0, 0))
    for crop in ([-1, -1, 0, 0], [float("nan"), 0, 0, 0], [0, 0, 1, 0], [0, 0]):
        bad = copy.deepcopy(spec); bad["slides"][0]["elements"][0]["crop"] = crop
        check("deep: 非物理 crop 在硬门拒绝 " + str(crop), not verify.check_spec(bad)["passed"])

    folder = tmp / "sheet"; paths = ghost.ghost_deck(base, folder)
    cache = folder / "cache"; cache.mkdir(); target = folder / "sheet.png"
    ghost.make_contact_sheet(paths, target, cache_dir=cache)
    cached = next(cache.glob("sheet-*.png")); cached.write_bytes(b"broken PNG")
    ghost.make_contact_sheet(paths, target, cache_dir=cache)
    try:
        with Image.open(target) as decoded: decoded.verify()
        valid_png = True
    except OSError:
        valid_png = False
    check("deep: 损坏联络表缓存重建并修复而非发布坏字节", valid_png
          and target.read_bytes() != b"broken PNG" and cached.read_bytes() != b"broken PNG")
    if not valid_png:
        target.write_bytes(paths[0].read_bytes())  # 下一项中断测试仍以完整 PNG 为前置
    previous = target.read_bytes()
    def partial_save(img, destination, *args, **kw):
        if hasattr(destination, "write"): destination.write(b"partial PNG")
        else: Path(destination).write_bytes(b"partial PNG")
        raise KeyboardInterrupt()
    with patch.object(Image.Image, "save", partial_save):
        try:
            if hasattr(ghost, "_write_png"): ghost._write_png(target, split)
            else: split.save(target, "PNG")
        except KeyboardInterrupt: pass
    check("deep: PNG 写入中断保留原完整图且清理暂存", target.read_bytes() == previous
          and not list(target.parent.glob("." + target.name + "-*")))
    quiet(build, output, mode="release", speed="strict")
    release = json.loads(output.with_suffix(".manifest.json").read_text(encoding="utf-8"))
    preview_claim = copy.deepcopy(release["ghost_preview"])
    preview_claim["pages"][0] = str(tmp / "absent-preview.png")
    check("deep: 有声明但无实际页图不算发布预览", bool(verify.preview_issues(preview_claim, ["one"])))

    # 离线评测导出也不能成功后毁掉自己的案例/基线文件。
    import bench
    scripts = Path(vao.__file__).parent
    for script, source in (("bench.py", "judgment_cases.json"),
                           ("unseen_bench.py", "generalization_cases.json")):
        cases = tmp / source; cases.write_bytes((scripts / "json" / source).read_bytes()); before = cases.read_bytes()
        run = subprocess.run([sys.executable, str(scripts / script), "--cases", str(cases), "--out", str(cases)],
                             capture_output=True, text=True, encoding="utf-8")
        check(f"deep: {script} 报告不能覆盖案例", run.returncode == 2 and cases.read_bytes() == before)
    baseline = tmp / "baseline.json"; json_write(baseline, bench.run(scripts / "json" / "judgment_cases.json")); before = baseline.read_bytes()
    run = subprocess.run([sys.executable, str(scripts / "bench.py"), "--baseline", str(baseline), "--out", str(baseline)],
                         capture_output=True, text=True, encoding="utf-8")
    check("deep: bench 报告不能覆盖回归基线", run.returncode == 2 and baseline.read_bytes() == before)


def test_integrity_followup(tmp: Path):
    """失败发布、行丢失、坏缓存及盲评输入的真实失败边界。"""
    import copy
    import csv
    import io
    import contextlib
    import bench
    import compiler
    import assets
    import vao
    import verify
    from primitives import CACHE_NAME, json_write
    base, _ = vao.load_spec(make_build(tmp))

    def deck(element):
        spec = copy.deepcopy(base)
        spec["slides"] = [spec["slides"][2]]
        spec["slides"][0]["elements"] = [element]
        return spec

    chart = {"id": "c", "type": "chart", "chart_kind": "process_flow",
             "x": 48, "y": 232, "width": 1184, "height": 400}
    for bad in (None, {}, {"label": None}, {"label": ""}, {"label": "  "}):
        spec = deck(dict(chart, data=[{"label": "保留内容"}, bad]))
        check(f"integrity: 非数值图不丢行 {bad!r}", not verify.check_spec(spec)["passed"])
        target = tmp / "invalid-rows.pptx"
        report = compiler.compile_deck(spec, target)
        check(f"integrity: 独立编译不静默丢行 {bad!r}", not report["passed"])
    multi = dict(chart, chart_kind="column", categories=["A"],
                 series=[{"name": "Series", "values": [1]}])
    for bad in ([], {}, "", False, None):
        spec = deck(dict(multi, data=bad))
        check(f"integrity: 显式双载荷 {bad!r}", not verify.check_spec(spec)["passed"])
        check(f"integrity: compiler 双载荷 {bad!r}",
              not compiler.compile_deck(spec, tmp / "dual.pptx")["passed"])

    target = tmp / "preserved.pptx"
    target.write_bytes(b"previous complete delivery")
    spec = deck({"id": "bad", "type": "unknown", "x": 0, "y": 0, "width": 10, "height": 10})
    report = compiler.compile_deck(spec, target)
    check("integrity: 编译失败保留旧完整产物", not report["passed"]
          and target.read_bytes() == b"previous complete delivery")
    target.unlink()
    report = compiler.compile_deck(spec, target)
    check("integrity: 编译失败不发布部分新产物", not report["passed"] and not target.exists())

    minimal = {"slides": [{"id": "minimal", "elements": [
        {"id": "t", "type": "text", "x": 10, "y": 10, "width": 400,
         "height": 80, "size": 24, "text": "中文默认背景"}]}]}
    check("integrity: 未声明主题背景不是编译失败",
          compiler.compile_deck(minimal, tmp / "minimal.pptx")["passed"])
    hidden = dict(multi, height=120, categories=list("abcdef"),
                  series=[{"name": "Series", "values": [1, 2, 3, 4, 5, 6]}],
                  show_values=True, label_collision_policy="hide_redundant")
    check("integrity: 作者授权隐藏冗余标签不是编译失败",
          compiler.compile_deck(deck(hidden), tmp / "authorized-hide.pptx")["passed"])

    # Guard 放行但绘制层失败：CLI 不得覆盖旧 PPTX、签署其摘要或污染编译缓存。
    build = tmp / "runtime.json"; json_write(build, base)
    output = tmp / "runtime.pptx"
    with contextlib.redirect_stdout(io.StringIO()):
        _, code = vao.run_check(str(build), str(output), mode="release", speed="strict")
    old = output.read_bytes()
    broken = copy.deepcopy(base)
    broken["slides"][0]["elements"].append({"id": "unsupported-shape", "type": "shape",
        "shape": "not-a-shape", "x": 64, "y": 300, "width": 100, "height": 100})
    json_write(build, broken)
    with contextlib.redirect_stdout(io.StringIO()):
        result, failed_code = vao.run_check(str(build), str(output), mode="release", speed="strict")
    manifest = json.loads(output.with_suffix(".manifest.json").read_text(encoding="utf-8"))
    check("integrity: release 绘制失败保留旧 PPTX 且撤销资格", code == 0
          and failed_code == 2 and old == output.read_bytes()
          and manifest["status"] == "BLOCKED" and not manifest["release_eligible"])
    check("integrity: 不为旧 PPTX 签署本轮失败编译摘要", not result["compile"].get("output_sha256"))
    cache = json.loads((tmp / "runtime_vao" / CACHE_NAME).read_text(encoding="utf-8"))
    check("integrity: 失败编译不写入缓存", cache["compile"]["report"]["passed"])

    # 合法 JSON 但错误顶层类型仍是坏缓存，而不是生产输入错误。
    good = deck(dict(chart, data=[{"label": "A"}, {"label": "B"}]))
    preview = tmp / "preview"
    preview.mkdir()
    json_write(preview / "ghost.meta.json", ["corrupt cache"])
    try:
        result = vao._ghost(good, preview, speed="strict", output_sha="a" * 64)
        ok = result["count"] == 1 and not result["reused"]
    except (AttributeError, TypeError, ValueError):
        ok = False
    check("integrity: 非对象 ghost marker 重建", ok)
    brief = {"audience": "评审", "decision": "批准", "tension": "证据",
             "slides": [{"id": "s01", "title": "检查素材", "asset": "required",
                         "asset_subject": "a tea bowl"}]}
    import intelligence
    bundle = intelligence.think(brief)
    cache = tmp / "prompt-cache.json"
    json_write(cache, {"schema": "vao-asset-prompt-cache-v3", "entries": ["bad"]})
    try:
        manifest = assets.build_manifest(brief, bundle, cache_path=cache)
        ok = len(assets.asset_entries(manifest)) == 1
    except (AttributeError, TypeError, ValueError):
        ok = False
    check("integrity: 错型 prompt cache 重建", ok)
    brief_file = tmp / "brief.json"; json_write(brief_file, brief)
    manifest_file = tmp / "assets.json"
    json_write(manifest_file, {"schema": "vao-assets-v2", "assets": [],
        "workflow": {"schema": "vao-asset-chain-v2", "brief_path": str(brief_file),
                     "brief_file_sha256": assets.file_digest(brief_file)}})
    json_write(tmp / "assets.qc.json", ["bad"])
    try:
        rep, code = assets.qc_report(manifest_file)
        ok = code == 0 and rep["status"] == "PASS"
    except (AttributeError, TypeError, ValueError):
        ok = False
    check("integrity: 非对象 QC cache 重建", ok)

    key = tmp / "key.json"; json_write(key, {"V01": "a.pptx", "V02": "b.pptx", "V03": "c.pptx"})
    sheet = tmp / "scores.csv"

    def scores(items, value=None, duplicate=False):
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow(["reviewer", "item", *bench.BLIND_DIMS, "reason"])
        for r in range(3):
            for i, item in enumerate(items):
                writer.writerow([f"r{r}", item, *([value or str(i + 1)] * len(bench.BLIND_DIMS)), "多行\n理由"])
        if duplicate:
            writer.writerow(["r0", items[0], *(["3"] * len(bench.BLIND_DIMS)), "duplicate"])
        sheet.write_text(stream.getvalue(), encoding="utf-8")

    for val in ("nan", "inf", "-inf"):
        scores(["V01", "V02", "V03"], val)
        try:
            bench.blind_score([str(sheet)], str(key)); rejected = False
        except ValueError:
            rejected = True
        check(f"integrity: 非有限盲评分拒绝 {val}", rejected)
    scores(["V01", "V02"])
    rep = bench.blind_score([str(sheet)], str(key))
    check("integrity: 全体漏评同一版本不能宣称共识", rep["consensus"] == "NOT_ENOUGH_REVIEWERS"
          and rep["complete_reviewers"] == 0)
    for items, duplicate in ((["V01", "UNKNOWN"], False), (["V01", "V02", "V03"], True)):
        scores(items, duplicate=duplicate)
        try:
            bench.blind_score([str(sheet)], str(key)); rejected = False
        except ValueError:
            rejected = True
        check(f"integrity: 未知版本/重复评分拒绝 {duplicate}", rejected)
    scores(["V01", "V02", "V03"])
    check("integrity: 合法独立评分仍形成统计一致", bench.blind_score([str(sheet)], str(key))["consensus"]
          == "CONSENSUS_SUPPORTED")
    scores(["V01", "V02", "V03"], "nan")
    with contextlib.redirect_stderr(io.StringIO()), contextlib.redirect_stdout(io.StringIO()):
        try:
            code = bench.main(["blind", "score", str(key), str(sheet)])
        except Exception:
            code = None
    check("integrity: 盲评非法输入 CLI 返回 2", code == 2)
    src = tmp / "source"; src.mkdir(); (src / "content.txt").write_text("new", encoding="utf-8")
    out = tmp / "blind-output"
    bench.blind_pack([str(src)], out, seed=1)
    stale = out / "V01" / "stale.txt"; stale.write_text("private old version", encoding="utf-8")
    try:
        bench.blind_pack([str(src)], out, seed=1); rejected = False
    except (ValueError, FileExistsError):
        rejected = True
    check("integrity: 重复盲评导出拒绝污染且保留旧工作", rejected
          and stale.read_text(encoding="utf-8") == "private old version")


def test_atomic_publish(tmp: Path):
    """原子发布的权限口径：mkstemp 的 0600 不得泄漏成产物权限。

    背景（本轮审计实测复现）：mkstemp 以 0600 建暂存文件，os.replace 把权限一起带到
    产物上，于是 out.pptx / manifest / QC / 缓存 / 预览 PNG 在共享目录、CI 收集、另一
    服务账号读取时全部 Permission denied；覆盖已存在的 0644 文件还会静默收窄权限。
    纪律与 open(path,"w") 一致：目标存在 → 保持原权限；不存在 → umask 默认（此处
    断言「同 umask 下与普通 open 写出的文件同权限」，在任意 umask 下都成立）。
    """
    import os
    from primitives import json_write
    from compiler import compile_deck

    probes = tmp / "publish"
    probes.mkdir()

    # 1) 新建 JSON（即 manifest/QC/缓存/plan 走的那条路）：权限与普通写入一致
    fresh = probes / "fresh.json"
    json_write(fresh, {"ok": True})
    plain = probes / "plain.txt"
    plain.write_text("x", encoding="utf-8")
    same = (fresh.stat().st_mode & 0o777) == (plain.stat().st_mode & 0o777)
    check("publish: 新建 JSON 权限与普通写入一致（非 mkstemp 0600）", same)

    # 2) 覆盖既有文件不得收窄权限。
    #    Windows 没有 POSIX 权限位（stat 只反映只读标志：chmod(0o644/0664/0600) 之后
    #    一律报 0o666），所以判据写成跨平台可证的形式——「覆盖前后权限不变」，
    #    只有在 POSIX 上才追加数值相等。CI 的 Windows 腿因此不会把 0644/0664 误判成缺陷。
    posix = os.name == "posix"
    os.chmod(fresh, 0o644)
    before = fresh.stat().st_mode & 0o777
    json_write(fresh, {"ok": True, "round": 2})
    after = fresh.stat().st_mode & 0o777
    check("publish: 覆盖既有文件保持原权限（POSIX 下即 0644）",
          after == before and (not posix or after == 0o644))
    os.chmod(fresh, 0o664)
    before = fresh.stat().st_mode & 0o777
    json_write(fresh, {"ok": True, "round": 3})
    after = fresh.stat().st_mode & 0o777
    check("publish: 覆盖既有 0664 不收窄（组可写位保持）",
          after == before and (not posix or after == 0o664))

    # 3) 真正的 PPTX 产物（compiler 的 mkstemp→save→replace 路径）
    pptx = probes / "deck.pptx"
    spec = {"canvas": {"width": 1280, "height": 720}, "theme": {"colors": {}},
            "slides": [{"id": "s01", "elements": []}]}
    compile_deck(spec, str(pptx), speed="fast")
    check("publish: PPTX 产物权限与普通写入一致（非 0600）",
          (pptx.stat().st_mode & 0o777) == (plain.stat().st_mode & 0o777))
    os.chmod(pptx, 0o600)
    before = pptx.stat().st_mode & 0o777
    compile_deck(spec, str(pptx), speed="fast")
    after = pptx.stat().st_mode & 0o777
    check("publish: 覆盖既有 PPTX 保持原权限（POSIX 下即 0600，不越权放宽）",
          after == before and (not posix or after == 0o600))


def test_root_order(tmp: Path):
    """同名相对底图在两处各有一张时，预览/编译/见证必须解析同一张。

    背景（本轮审计实测复现）：编译侧按 [输出目录, 稿件目录] 解析相对 src，
    vao 传给预览的根却是 [稿件目录, cwd]——`out/images/hero.png` 与
    `src/images/hero.png` 不同色时，产物贴的是输出目录那张，预览画的是稿件目录
    那张，两边都自认正确。修复：vao 把编译侧的两个根（同序）显式传给预览，
    ghost 以 `spec["_image_roots"]` 优先，独立调用 ghost 的老行为不变。
    """
    import ghost as ghost_mod
    from compiler import _resolve_src
    from PIL import Image
    from vao import _ghost

    root = tmp / "rootorder"
    build_dir, out_dir = root / "src", root / "out"
    (build_dir / "images").mkdir(parents=True)
    (out_dir / "images").mkdir(parents=True)
    red, blue = (220, 30, 30), (30, 30, 220)
    Image.new("RGB", (80, 80), red).save(build_dir / "images" / "hero.png")
    Image.new("RGB", (80, 80), blue).save(out_dir / "images" / "hero.png")

    element = {"type": "image", "id": "hero", "x": 0, "y": 0, "width": 1280,
               "height": 720, "src": "images/hero.png", "role": "focus"}
    spec = {"canvas": {"width": 1280, "height": 720},
            "theme": {"colors": {"background": "#FFFFFF", "ink": "#111111"}},
            "slides": [{"id": "s01", "elements": [element]}]}

    compiled = _resolve_src(element, str(out_dir / "deck.pptx"), str(build_dir / "build.py"))
    check("root-order: 编译侧取输出目录那张（基准口径）",
          Path(compiled).resolve() == (out_dir / "images" / "hero.png").resolve())

    info = _ghost(spec, out_dir / "preview", base_path=build_dir,
                  image_roots=[out_dir, build_dir], speed="strict", limit=1)
    page = Path(info["pages"][0])
    check("root-order: 预览产出页图", page.is_file())
    with Image.open(page) as im:
        pixel = im.convert("RGB").getpixel((im.width // 2, im.height // 2))
    preview_is_blue = pixel[2] > pixel[0] and pixel[2] > 150
    check("root-order: 预览画的是输出目录那张（与产物同源）", preview_is_blue,
          f"中心像素={pixel}（红=稿件目录那张，说明分叉）")

    # 独立调用 ghost（无 _image_roots）保持旧口径：base 优先
    legacy = dict(spec)
    legacy["_base_path"] = str(build_dir)
    check("root-order: 独立 ghost 调用仍按 base 优先（兼容）",
          ghost_mod._resolve_image("images/hero.png", ghost_mod._image_roots(legacy))
          == (build_dir / "images" / "hero.png").resolve())


def _header_only_png(path: Path, w: int, h: int) -> Path:
    """合法签名 + 合法 IHDR（含 CRC）+ 空 IDAT：`Image.open` 读得到 size，
    `load()` 必然失败。用来证明预算判断发生在**解码之前**且不花内存。"""
    import struct
    import zlib

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xffffffff))

    path.write_bytes(b"\x89PNG\r\n\x1a\n"
                     + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(b"")) + chunk(b"IEND", b""))
    return path


def test_prompt_cache_key(tmp: Path):
    """提示词缓存键必须覆盖提示词的**全部输入**。

    实测（本轮审计）：身份指纹刻意不含 safe_area / text_color / ratio（校准压字区
    不该让已出图改名），但 `build_asset_prompt` 恰恰读它们。按身份查缓存时，
    作者把压字区从左侧改到右侧、再跑一次 plan，拿回的是旧提示词，而清单里的
    `safe_area` 已是新值——清单描述与出图指令互相矛盾，且**逐字相同**的提示词
    让人完全看不出这一步没生效。
    """
    from assets import build_manifest, prompt_cache_key
    from intelligence import load_brief, think

    work = tmp / "prompt-cache"
    work.mkdir()
    brief_file = work / "brief.yml"
    brief_file.write_text(
        "goal: 二期产线改造路演材料\n"
        "audience: 制造业客户与战略伙伴\n"
        "decision: 批准二期产线改造预算\n"
        "slides:\n"
        "  - id: s01\n"
        "    title: \"二期产线改造：11 个月收回投资\"\n"
        "    content: \"改造预算 4,800 万元；按 2025 年实际节拍测算，月均节省 436 万元。\"\n"
        "    asset: required\n"
        "    asset_subject: \"a production line in a factory\"\n"
        "    asset_role: illustration\n", encoding="utf-8")
    brief = load_brief(str(brief_file))
    bundle = think(brief)
    cache = work / "asset_prompt_cache.json"
    area = {"x": 0.06, "y": 0.08, "width": 0.34, "height": 0.78}

    def plan(extra):
        for page in bundle["pages"]:
            page["declarations"] = dict(page.get("declarations") or {}, **extra)
        manifest = build_manifest(brief, bundle, cache_path=cache)
        return manifest["assets"][0], manifest["prompt_cache"]

    dark, _ = plan({"safe_area": area, "text_color": "#111111", "asset_ratio": "16:9"})
    check("prompt-cache: 压字区从深色字改浅色字 → 提示词必须重算",
          "suitable for dark type" in dark["prompt"])
    light, _ = plan({"safe_area": area, "text_color": "light", "asset_ratio": "16:9"})
    check("prompt-cache: 明暗指令随 text_color 改变（不是复用旧提示词）",
          light["prompt"] != dark["prompt"]
          and "suitable for light type" in light["prompt"])
    check("prompt-cache: 资产 id 不因压字区校准改名（身份仍稳定）",
          light["asset_id"] == dark["asset_id"])

    no_area, _ = plan({"safe_area": {}, "text_color": "#111111", "asset_ratio": "16:9"})
    check("prompt-cache: 显式无安全区 → 留白指令消失（提示词随输入变化）",
          "suitable for dark type" not in no_area["prompt"]
          and no_area["prompt"] != dark["prompt"])

    again, stats = plan({"safe_area": area, "text_color": "#111111", "asset_ratio": "16:9"})
    check("prompt-cache: 输入完全相同时仍然命中缓存（不是每次重算）",
          stats["hits"] == 1 and again["prompt"] == dark["prompt"])

    key_a = prompt_cache_key("asset-x", {}, {"safe_area": area, "text_color": "#111111"})
    key_b = prompt_cache_key("asset-x", {}, {"safe_area": area, "text_color": "light"})
    check("prompt-cache: 缓存键把提示词输入混进身份（键本身可解释）",
          key_a != key_b and key_a.startswith("asset-x|"))


def test_chart_target(tmp: Path):
    """目标线：声明了就必须画得出来，画不出来必须出声。

    实测（本轮审计）：`target` 只有 ranked_bar 分支消费，而编译器把转换异常
    `except (TypeError, ValueError): pass` 吞掉——`target: "abc"` 的稿件 release
    仍 PASS，产物里一根参考线都没有，作者却以为「门槛已经标上」。
    反向的错觉同样要拦：progress_bar 上写 target 根本不渲染，必须留痕。
    """
    import zipfile

    from compiler import compile_deck
    from verify import check_spec

    work = tmp / "target"
    work.mkdir()
    theme = {"colors": {"background": "#FFFFFF", "ink": "#111111", "muted": "#777777",
                        "primary": "#222222", "secondary": "#333333", "accent": "#AA0000"},
             "fonts": {"cn": "Source Han Sans SC", "latin": "Arial"}}

    def spec(kind, **extra):
        element = {"type": "chart", "id": "c", "x": 100, "y": 100, "width": 900,
                   "height": 400, "chart_kind": kind, "metric": "完成率", "unit": "%",
                   "source": "内部", "period": "2025", "basis": "同口径",
                   "data": [{"label": "甲线", "value": 62}, {"label": "乙线", "value": 48}]}
        element.update(extra)
        return {"canvas": {"width": 1280, "height": 720}, "theme": theme,
                "slides": [{"id": "s01", "elements": [element]}]}

    office = work / "ranked.pptx"
    rep = compile_deck(spec("ranked_bar", target=50, target_label="目标 50%"),
                       str(office), speed="strict")
    with zipfile.ZipFile(office) as z:
        xml = z.read("ppt/slides/slide1.xml").decode("utf-8")
    check("target: ranked_bar 的数字目标线落进产物（线 + 标签）",
          bool(rep.get("passed")) and 'name="c__target"' in xml
          and 'name="c__target_label"' in xml)

    bad = check_spec(spec("ranked_bar", target="abc"))
    msgs = [c["msg"] for c in bad["checks"] if c["level"] == "error"]
    check("target: 非数字目标被 guard 拒绝（不再静默消失）",
          not bad["passed"] and any("target" in m for m in msgs), str(msgs[:1])[:110])

    direct = compile_deck(spec("ranked_bar", target="abc"), str(work / "bad.pptx"), speed="strict")
    check("target: 编译器兜底（直接编译也不静默改稿）",
          not direct.get("passed")
          and any("target" in str(w) for w in direct.get("warnings") or []))

    loose = check_spec(spec("progress_bar", max=100, target=80))
    warns = [(c["rule"], c["msg"]) for c in loose["checks"] if c["level"] == "warn"]
    check("target: 不被消费的 kind 出告警（写了不生效必须留痕）",
          loose["passed"] and any(r == "chart_field_unconsumed" and "target" in m
                                  for r, m in warns),
          str(warns[:1])[:110])


def test_char_spacing(tmp: Path):
    """字距越界会写出非法 OOXML：`a:rPr/@spc` 是 ST_TextPoint（±4000pt）。

    实测（本轮审计）：`char_spacing: 1000000`（pt）一路 PASS 并把
    `spc="100000000"` 写进 slide XML——超出 ECMA-376 的值域，PowerPoint 打开会进
    修复流程。guard 前置拒绝（可修诊断），编译器按同一常量夹紧兜底。
    """
    import re
    import zipfile

    from compiler import compile_deck
    from primitives import CHAR_SPACING_MAX_PT
    from verify import check_spec

    work = tmp / "spacing"
    work.mkdir()
    theme = {"colors": {"background": "#FFFFFF", "ink": "#111111",
                        "muted": "#777777", "primary": "#222222",
                        "secondary": "#333333", "accent": "#AA0000"},
             "fonts": {"cn": "Source Han Sans SC", "latin": "Georgia"}}

    def spec(spacing):
        return {"canvas": {"width": 1280, "height": 720}, "theme": theme,
                "slides": [{"id": "s01", "elements": [
                    {"type": "text", "id": "t", "x": 100, "y": 100, "width": 800,
                     "height": 200, "text": "自有内容占比 61%", "size": 21,
                     "color": "ink", "role": "title", "char_spacing": spacing}]}]}

    legal = check_spec(spec(2))
    check("spacing: 正常字距（2pt）不受影响", legal["passed"])

    wild = check_spec(spec(10 ** 6))
    msgs = [c["msg"] for c in wild["checks"] if c["level"] == "error"]
    check("spacing: 越界字距被 guard 拒绝（可修诊断）",
          not wild["passed"] and any("ST_TextPoint" in m for m in msgs), str(msgs[:1])[:110])

    for value, expect in ((2, 200), (4000, int(CHAR_SPACING_MAX_PT * 100)),
                          (10 ** 6, int(CHAR_SPACING_MAX_PT * 100))):
        office = work / f"spc-{value}.pptx"
        rep = compile_deck(spec(value), str(office), speed="strict")
        with zipfile.ZipFile(office) as z:
            written = re.findall(r'spc="(-?\d+)"',
                                 z.read("ppt/slides/slide1.xml").decode("utf-8"))
        check(f"spacing: 编译写入 spc={expect}（声明 {value}pt）",
              bool(rep.get("passed")) and written and all(int(w) == expect for w in written),
              f"写入 {written}")

    bounded = compile_deck(spec(-10 ** 6), str(work / "spc-neg.pptx"), speed="strict")
    with zipfile.ZipFile(work / "spc-neg.pptx") as z:
        neg = [int(w) for w in re.findall(r'spc="(-?\d+)"',
                                         z.read("ppt/slides/slide1.xml").decode("utf-8"))]
    check("spacing: 负方向同样夹在可表达区间内（不写非法 OOXML）",
          bool(bounded.get("passed")) and neg
          and all(abs(v) <= CHAR_SPACING_MAX_PT * 100 for v in neg), f"写入 {neg}")


def test_asset_budget(tmp: Path):
    """不可信资产的资源边界：超预算的图必须在**解码前**被拦成可修诊断。

    背景（本轮审计实测）：12000×12000（144MP）的图在 2GB 机器上把整轮 check 打到
    OOM 被 SIGKILL——作者拿到的是「进程没了」。Pillow 自己的炸弹线是 178.9MP，
    89.5–178.9M 那段只警告不拦；QC 的 fast 档只缩小统计工作域，原图仍要完整解码。
    所以设一条流水线共用的解码预算（`assets.QC_MAX_PIXELS`），越线即 BLOCK。
    """
    import time

    import numpy as np
    from PIL import Image

    import assets as assets_mod

    probe = tmp / "budget"
    probe.mkdir()
    rng = np.random.default_rng(11)
    small = probe / "small.png"
    Image.fromarray((rng.normal(205, 12, (656, 640, 3)).clip(0, 255)).astype("uint8")).save(small)

    plain = assets_mod.image_qc(str(small), max_side=None)
    check("budget: 预算内的正常图不受影响",
          plain.get("status") == "ok" and plain.get("dimensions") == [640, 656])

    saved = assets_mod.QC_MAX_PIXELS
    try:
        assets_mod.QC_MAX_PIXELS = 100_000
        tight = assets_mod.image_qc(str(small), max_side=None)
        check("budget: 越线即 BLOCK，判词含像素数与预算",
              tight.get("status") == "error" and "解码预算" in str(tight.get("issue"))
              and "100,000" in str(tight.get("issue")))
    finally:
        assets_mod.QC_MAX_PIXELS = saved

    # 真实预算下的一张 144MP「只有头」的图：必须由**我们的**预算判词拦下，
    # 且在一瞬间完成——真去解码这个文件只会报 truncated。
    huge = _header_only_png(probe / "huge.png", 12000, 12000)
    started = time.perf_counter()
    blocked = assets_mod.image_qc(str(huge), max_side=None)
    elapsed = time.perf_counter() - started
    check("budget: 144MP 在解码前被拦（不是 truncated、不是 OOM）",
          blocked.get("status") == "error" and "解码预算" in str(blocked.get("issue"))
          and "truncated" not in str(blocked.get("issue")) and elapsed < 1.0,
          f"{elapsed*1000:.0f}ms · {str(blocked.get('issue'))[:70]}")

    check("budget: 预算写进 QC 档案（可审计）",
          assets_mod.qc_profile("strict")["max_pixels"] == saved
          and assets_mod.qc_profile("fast")["max_pixels"] == saved)


def test_asset_diagnosis(tmp: Path):
    """资产读不出来时，QC 的原始判词必须直达修复包。

    此前这种资产进 `pending_assets`，链上只报「无接受该图片的 QC 结果」；真正的原因
    （损坏 / 截断 / 超预算）只写在 QC 报告 results[].qc.issue 里，作者得再开一个文件
    才知道哪里坏了——与「硬门阻断必须是可修的诊断」不符。
    """
    from assets import qc_report, verify_chain

    work = tmp / "diagnosis"
    work.mkdir()
    broken = work / "broken.png"
    broken.write_bytes(b"this is not an image at all\n" * 40)      # 伪装成 png 的文本
    manifest = work / "asset_manifest.json"
    manifest.write_text(json.dumps({
        "schema": "vao-assets-v2",
        "assets": [{"asset_id": "a1", "decision": "generate",
                    "expected_filename": "broken.png", "prompt": "p", "negative": ["n"],
                    "slide_ids": ["s01"], "asset_role": "illustration"}]}), encoding="utf-8")
    report, code = qc_report(str(manifest), str(work), phase="release", speed="strict")
    check("asset-diag: 坏图 QC 判 BLOCKED 且记下原始判词",
          code == 2 and report["status"] == "BLOCKED"
          and report["results"][0].get("missing") is True
          and bool(str((report["results"][0].get("qc") or {}).get("issue") or "").strip()))

    spec = {"slides": [{"id": "s01", "elements": [
        {"type": "image", "id": "hero", "asset_id": "a1", "src": str(broken),
         "x": 640, "y": 0, "width": 640, "height": 656}]}]}
    chain = verify_chain(spec, str(manifest), str(report["report_path"]), str(work))
    text = " | ".join(chain["issues"])
    check("asset-diag: 链上判词带出 QC 的原始原因（不是只有「无接受结果」）",
          chain["status"] == "BLOCKED" and "图片不可用（" in text,
          text[:120])


def test_guard_diagnostics(tmp: Path):
    """硬门阻断必须给出可修诊断：坐标的 bool、画布的物理窗口。

    两处都是本轮审计实测：`x: True` 一路 PASS（float(True)=1px 静默当坐标），
    与 canvas.width:True 被 `_positive_num` 拒是两套标准；canvas 5377px 则让
    python-pptx 抛裸范围错误（`value must be in range(914400, 51206400)`），
    作者读不出是画布尺寸的问题。
    """
    from verify import check_spec
    from primitives import SLIDE_MAX_PX

    def base(element):
        return {"canvas": {"width": 1280, "height": 720},
                "theme": {"colors": {"background": "#FFFFFF", "ink": "#111111"}},
                "slides": [{"id": "s01", "elements": [element]}]}

    text = {"type": "text", "id": "t", "x": 100, "y": 200, "width": 600,
            "height": 80, "text": "结论句", "size": 24, "color": "ink",
            "role": "title"}
    ok = check_spec(base(text))
    check("guard-diag: 基准合规稿无错误", ok["passed"])

    bad = check_spec(base(dict(text, x=True)))
    msgs = [c["msg"] for c in bad["checks"] if c["level"] == "error"]
    check("guard-diag: 坐标 bool 被拒（不再是静默 1px）",
          not bad["passed"] and any("geometry" in m for m in msgs), str(msgs[:2]))

    sub = check_spec(base(dict(text, width=5e-5)))
    sub_msgs = [c["msg"] for c in sub["checks"] if c["level"] == "error"]
    check("guard-diag: 亚 EMU 尺寸被判退化（落盘 0 EMU = 不可见）",
          not sub["passed"] and any("退化" in m for m in sub_msgs), str(sub_msgs[:2]))

    wide = check_spec({"canvas": {"width": SLIDE_MAX_PX + 1, "height": 720},
                       "theme": {"colors": {"background": "#FFFFFF", "ink": "#111111"}},
                       "slides": [{"id": "s01", "elements": [text]}]})
    msgs2 = [c["msg"] for c in wide["checks"] if c["level"] == "error"]
    check("guard-diag: 超 56 英寸画布给可修诊断（不是裸 ValueError）",
          not wide["passed"] and any("1–56 英寸" in m for m in msgs2), str(msgs2[:2]))

    legal = check_spec({"canvas": {"width": SLIDE_MAX_PX, "height": 720},
                        "theme": {"colors": {"background": "#FFFFFF", "ink": "#111111"}},
                        "slides": [{"id": "s01", "elements": [text]}]})
    check("guard-diag: 56 英寸边界本身合法",
          all("英寸" not in c["msg"] for c in legal["checks"]))


def test_unit_visibility(tmp: Path):
    """unit × show_values：release 出处硬门要求 unit，而值标签关闭时 unit 不可见。

    此前编译器把这个组合直接升格为编译失败，而 release（-5 秒 deadline 除外）
    又强制数值图表写 unit：作者只有「打开数值标签」一条路，修复包只报
    COMPILE_FAIL，说不出这是两个门的冲突。修复：编译器不再拦截（渲染职责），
    guard 出可执行告警 `chart_unit_invisible`（进发布清单 notes，不阻断）。
    """
    from verify import check_spec
    from compiler import compile_deck

    spec = {"canvas": {"width": 1280, "height": 720},
            "theme": {"colors": {"background": "#FFFFFF", "ink": "#111111",
                                 "muted": "#777777", "primary": "#222222",
                                 "secondary": "#333333", "accent": "#AA0000"}},
            "slides": [{"id": "s01", "elements": [
                {"type": "chart", "id": "c", "x": 140, "y": 100, "width": 1000,
                 "height": 520, "chart_kind": "pie", "unit": "件",
                 "source": "内部测算", "period": "2025", "basis": "同型号",
                 "data": [{"label": "甲", "value": 30}, {"label": "乙", "value": 45},
                          {"label": "丙", "value": 25}]}]}]}

    report = check_spec(spec)
    rules = [(c.get("rule"), c.get("level")) for c in report.get("checks", [])]
    check("unit-visibility: guard 出可执行告警（不阻断）",
          ("chart_unit_invisible", "warn") in rules
          and all(level != "error" for _, level in rules))

    out = tmp / "unit" / "deck.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    rep = compile_deck(spec, str(out), speed="strict")
    check("unit-visibility: pie+unit 不再编译失败（渲染职责与设计决定分离）",
          bool(rep.get("passed")) and out.is_file())

    shown = json.loads(json.dumps(spec))
    shown["slides"][0]["elements"][0]["show_values"] = True
    report2 = check_spec(shown)
    check("unit-visibility: 值可见时不再告警",
          all(c.get("rule") != "chart_unit_invisible"
              for c in report2.get("checks", [])))


def test_art_direction_linkage(tmp: Path):
    """艺术指导链路钉：asset_function 词表同源、出图主体缺口、视觉分量与高潮传递。"""
    import assets as a
    import intelligence as intel

    # P1 · asset_function 词表与文档契约同源（hero|proof|emotion|context|frame|separate）：
    # 审计实测旧提示词表缺 hero/proof/emotion/context（写了静默丢失），
    # 另有 direct/contextualize/immersive 三个死键（全库零消费者）。
    check("fn 词表: 与文档六值同源，无死键",
          set(a.ASSET_FUNCTION_PHRASES) == {"hero", "proof", "emotion", "context",
                                            "frame", "separate"},
          str(sorted(a.ASSET_FUNCTION_PHRASES)))
    card = {"asset_type": "background", "subject": ["a workshop"], "color": ["muted green"],
            "material": ["stone"], "lighting": ["even light"], "composition": ["left subject"],
            "resolved": {"light": {}}}
    for fn in a.ASSET_FUNCTION_PHRASES:
        r = a.build_asset_prompt(card, {}, asset_function=fn, include_qc=False)
        check(f"fn 词表: {fn} 的用途句进入提示词（不再静默丢失）",
              a.ASSET_FUNCTION_PHRASES[fn] in r["prompt"], r["prompt"][:120])
    bad = {**card, "asset_function": "protagonist", "apc": "APC-TEST"}
    check("fn 词表: 未知 asset_function 在卡片自检里被点名",
          any("未知 asset_function" in x for x in a.validate_asset_card(bad)),
          str(a.validate_asset_card(bad)))
    check("fn 词表: 合法值不误报",
          not any("未知 asset_function" in x
                  for x in a.validate_asset_card({**bad, "asset_function": "hero"})))
    brief_fn = {"audience": "评审", "decision": "批准", "tension": "风险",
                "slides": [{"id": "s1", "title": "现场记录", "content": "分拣现场记录（来源：日报）",
                            "asset": "reuse", "asset_subject": "a sorting line",
                            "asset_source": {"kind": "provided", "path": "img.png",
                                             "source": "用户提供的现场照片"}}]}
    bundle_fn = intel.think(brief_fn)
    entry_fn = next(x for x in a.build_manifest(brief_fn, bundle_fn)["assets"]
                    if x.get("decision") == "existing")
    check("fn 词表: 既有素材缺省与生成路径同源（frame，不再两套缺省）",
          entry_fn["asset_function"] == "frame", str(entry_fn.get("asset_function")))

    # P2 · 出图主体缺口：声明出图但未声明画面主体 → 判断层必须开口（旧条件挂在
    # 永不触发的「judgment 来源 required」死分支上，plan 层静默，观点句当画面回落）。
    def oq(slide, **kw):
        b = intel.think({"audience": "评审", "decision": "批准预算", "tension": "成本",
                         "slides": [{"id": "p", "title": "把资源从买量转向内容资产",
                                     "content": "买量 CAC 涨 38%。（来源：台账）", **slide, **kw}]})
        return b["pages"][0]["judgment"]["open_questions"] or []
    check("subject 缺口: required 未声明主体 → 判断层提问（不再静默回落）",
          any("图像主题由作者确认" in q for q in oq({"asset": "required"})),
          str(oq({"asset": "required"})))
    check("subject 缺口: 声明 asset_subject 后不再提问",
          not any("图像主题由作者确认" in q
                  for q in oq({"asset": "required", "asset_subject": "a workshop bench"})))
    check("subject 缺口: 既有素材登记不需要画面主体（不误报）",
          not any("图像主题由作者确认" in q
                  for q in oq({"asset": "reuse",
                               "asset_source": {"kind": "provided", "path": "p.png",
                                                "source": "用户照片"}})))

    # P3 · 视觉分量与高潮：内容高潮（请求/决定句 → 主视觉素材 → 建立主张）推导，
    # 进判断卡与骨架，让 SKILL「视觉高潮必须对应内容高潮」在 BUILD 侧有落点。
    def deck(slides):
        return intel.think({"audience": "投资机构", "decision": "推进本轮融资",
                            "tension": "增长是否可持续", "slides": slides})
    pitch = deck([
        {"id": "s01", "title": "全年经常性收入增长 42%", "content": "ARR 增长 42%。（来源：财报）"},
        {"id": "s02", "title": "增量主要来自存量扩容", "content": "扩容贡献 61%。（来源：经营分析）"},
        {"id": "s03", "title": "请批准本轮融资 2.5 亿元", "content": "用于产能与海外交付网络。"},
    ])
    check("视觉分量: 请求/决定句所在页是高潮位（结论权重）",
          pitch["deck"]["coherence"]["climax"]["page"] == "s03"
          and pitch["pages"][2]["judgment"]["emphasis"]["level"] == "climax",
          str(pitch["deck"]["coherence"]["climax"]))
    hero = deck([
        {"id": "s01", "title": "青瓷茶碗立住这一季的器物语言",
         "content": "釉色与器型成为系列基准。", "asset": "required",
         "asset_subject": "a celadon tea bowl on a matte surface"},
        {"id": "s02", "title": "系列复购带动客单提升", "content": "客单 +18%。（来源：台账）"},
    ])
    check("视觉分量: 无请求句时主视觉素材页承接高潮（图像承担最强分量）",
          hero["deck"]["coherence"]["climax"]["page"] == "s01"
          and hero["pages"][0]["judgment"]["emphasis"]["level"] == "climax"
          and hero["pages"][1]["judgment"]["emphasis"]["level"] == "support",
          str([(p["id"], p["judgment"]["emphasis"]["level"]) for p in hero["pages"]]))
    framed = deck([
        {"id": "s01", "title": "谢谢观看", "content": "2026 年度复盘（汇报人：设计组）"},
        {"id": "s02", "title": "请批准把客服排班改为两班制", "content": "两班制覆盖 8:00–22:00。"},
    ])
    check("视觉分量: 凭证/框架页安静（quiet，不与主张页争注意力）",
          framed["pages"][0]["judgment"]["emphasis"]["level"] == "quiet"
          and framed["deck"]["coherence"]["climax"]["page"] == "s02",
          str([(p["id"], p["judgment"]["emphasis"]["level"]) for p in framed["pages"]]))
    sk = intel.build_skeleton(pitch)
    check("视觉分量: 骨架携带高潮行与逐页视觉分量（传递到 BUILD）",
          "视觉高潮（内容推导的候选）: s03" in sk and "视觉分量[climax]" in sk
          and "视觉分量[support]" in sk and "否决" not in sk, sk[:400])
    check("视觉分量: 无内容高潮时不硬指一页（缺主张/请求/素材就承认没有）",
          deck([{"id": "c1", "title": "谢谢观看", "content": "2026 年度复盘（汇报人：设计组）"},
                {"id": "c2", "title": "问答交流", "content": "欢迎提问。"}])
          ["deck"]["coherence"]["climax"]["page"] is None)

    # P5 · 比较结论谓语补全：「X 不到 Y 的三分之一」曾被判话题/无证据（「低于」
    # 在断言与比较两表而「不到」两表都缺；名词化形状「…的…」再踩一脚）。
    cmp_title = "追加人力成本不到流失挽回的三分之一"
    cmp_content = "月度人力成本 +9.6 万元；夜间客诉导致的月度流失挽回约 32 万元。（来源：财务测算，2026-09）"
    u_cmp = intel.understand(f"{cmp_title} {cmp_content}")
    cl_cmp = intel.extract_claim({"title": cmp_title, "content": cmp_content})
    check("比较谓语: 「不到」与「低于/不足」同族入断言与比较词表",
          cl_cmp["source"] == "extracted" and cl_cmp.get("text") == cmp_title
          and u_cmp["evidence"] == "comparison", str(cl_cmp))
    check("比较谓语: 名词化形状不再否决带断言谓语的标题",
          "作者标题为断言" in (cl_cmp.get("why") or ""), str(cl_cmp.get("why")))
    check("比较谓语: 纯名词化话题（无断言谓语）仍不顶替正文",
          intel.extract_claim({"title": "门店的晨会节奏",
                               "content": "15 分钟站立会。（来源：观察）"})["text"]
          != "门店的晨会节奏")


def test_ghost_chart_color_single_source(tmp: Path):
    """P6 · ghost 单序列条/柱基础色必须与产物同源（chart_primary_color）。"""
    import ghost as _ghost
    deck = {
        "canvas": {"width": 1280, "height": 720, "grid_columns": 12, "grid_unit": 8},
        "theme": {"colors": {"background": "#FFFFFF", "ink": "#12120F",
                             "muted": "#8A8984", "primary": "#12120F",
                             "secondary": "#8A8984", "accent": "#6E6A63"},
                  "fonts": {"cn": "Source Han Sans SC", "latin": "Arial"}},
        "slides": [{"id": "c1", "elements": [
            {"id": "ch", "type": "chart", "role": "chart",
             "x": 100, "y": 100, "width": 600, "height": 360,
             "chart_kind": "comparison_bar",
             "data": [{"label": "前", "value": 47}, {"label": "后", "value": 8}],
             "highlight": "后", "primary_color": "muted",
             "unit": "分钟", "source": "s", "period": "p", "basis": "b",
             "show_values": True}]}],
    }
    img = _ghost.ghost_page(deck["slides"][0], deck, scale=1, supersample=1).convert("RGB")
    muted_hits = ink_hits = 0
    for c in img.getdata():
        if abs(c[0] - 144) <= 10 and abs(c[1] - 143) <= 10 and abs(c[2] - 138) <= 10:
            muted_hits += 1
        if abs(c[0] - 70) <= 12 and abs(c[1] - 70) <= 12 and abs(c[2] - 67) <= 12:
            ink_hits += 1
    check("ghost 单序列条基础色读 primary_color（预览=产物 muted 条）",
          muted_hits > 3000, f"muted_px={muted_hits}")
    check("ghost 非高亮条不再硬编码墨色（无大块 ink 条）",
          ink_hits < 3000, f"ink_px={ink_hits}")


def test_case004_regressions(root: Path):
    """PPT_CASE_004 实测回流回归：瀑布图 subtotal/show_values/同源色 + 环形图逐扇区反色。"""
    import primitives, compiler, ghost
    from pptx import Presentation
    from pptx.oxml.ns import qn

    spec = {
        "canvas": {"width": 1280, "height": 720, "grid_columns": 12, "grid_unit": 8},
        "theme": {
            "colors": {
                "background": "#F4F0E8",
                "surface": "#EAE4D8",
                "ink": "#1C1A19",
                "muted": "#6E685F",
                "primary": "#5C232C",
                "secondary": "#D8D0C2",
                "accent": "#5C232C",
            },
            "fonts": {"cn": "Source Han Serif SC", "latin": "Georgia"},
        },
        "slides": [
            {
                "id": "s_wf",
                "elements": [
                    {
                        "id": "wf1",
                        "type": "chart",
                        "chart_kind": "waterfall",
                        "x": 64, "y": 120, "width": 760, "height": 420,
                        "primary_color": "#5C232C",
                        "secondary_color": "#6E685F",
                        "show_values": True,
                        "unit": "€",
                        "data": [
                            {"label": "单价", "value": 240},
                            {"label": "酿造", "value": -37},
                            {"label": "毛利", "value": 203, "subtotal": True},
                            {"label": "履约", "value": -62},
                            {"label": "净利", "value": 141, "total": True},
                        ],
                    }
                ],
            },
            {
                "id": "s_donut",
                "elements": [
                    {
                        "id": "dn1",
                        "type": "chart",
                        "chart_kind": "donut",
                        "x": 680, "y": 160, "width": 480, "height": 400,
                        "highlight": "旗舰",
                        "unit": "%",
                        "center_value": "8.6万瓶",
                        "data": [
                            {"label": "旗舰", "value": 46},
                            {"label": "正牌", "value": 39},
                            {"label": "序章", "value": 15},
                        ],
                    }
                ],
            },
        ],
    }

    rows = primitives.chart_rows(spec["slides"][0]["elements"][0])
    check("waterfall: chart_rows 保留 subtotal/is_total/total（小计从零轴起且不跨桥接）",
          [r["subtotal"] for r in rows] == [False, False, True, False, True],
          str([r.get("subtotal") for r in rows]))

    out_pptx = root / "case004_reg.pptx"
    rep = compiler.compile_deck(spec, str(out_pptx))
    prs = Presentation(str(out_pptx))
    wf_shapes = {s.name: s for s in prs.slides[0].shapes}
    bar0 = wf_shapes["wf1__bar_0"]
    bar2 = wf_shapes["wf1__bar_2"]
    # 小计段（203）与首段（240）底边均落在同一零轴上（而非叠到 406 高度）
    zero_diff = abs((bar0.top + bar0.height) - (bar2.top + bar2.height))
    has_val_boxes = all(f"wf1__v_{i}" in wf_shapes for i in range(5))
    check("waterfall: show_values 在产物落柱顶值（含 unit）且小计柱底对齐零轴",
          bool(rep.get("passed")) and zero_diff < 5000 and has_val_boxes
          and "€" in wf_shapes["wf1__v_0"].text
          and "wf1__conn_1" not in wf_shapes and "wf1__conn_2" not in wf_shapes,
          f"zero_diff={zero_diff} boxes={has_val_boxes}")

    gdir = root / "ghost_wf"
    ghost.ghost_deck(spec, str(gdir))
    from PIL import Image
    wf_img = Image.open(gdir / "ghost-01.png").convert("RGB")
    oxblood_px = sum(1 for r, g, b in wf_img.getdata()
                     if abs(r - 0x5C) <= 10 and abs(g - 0x23) <= 10 and abs(b - 0x2C) <= 10)
    check("waterfall: 预览正向柱色与产物 primary_color 同源（不硬编码墨色）",
          oxblood_px > 500, f"oxblood_px={oxblood_px}")

    dn_chart = [s for s in prs.slides[1].shapes if s.has_chart][0].chart
    pts = list(dn_chart.series[0].points)
    dl_colors = [str(pt.data_label.font.color.rgb) for pt in pts]
    has_ea = all(bool(pt.data_label.font._element.xpath("a:ea")) for pt in pts)
    # 至少同时存在深墨标签与浅反白标签（跨明度扇区自动反色），且保留 a:ea 中文字体
    lums = [sum(int(c[j:j + 2], 16) for j in (0, 2, 4)) / 765.0 for c in dl_colors]
    check("donut: 扇区标签按实色自动取反色且保留 a:ea 中文字体",
          has_ea and min(lums) < 0.25 and max(lums) > 0.75,
          f"colors={dl_colors} ea={has_ea}")

    ctx = compiler.RenderContext(spec["theme"])
    pal = ctx.series_palette(3, 0)
    expected = [str(ctx.auto_text_for(c)) for c in pal]
    check("donut: 预览与产物扇区文字明暗同源 auto_text_for",
          dl_colors == expected, f"pptx={dl_colors} expected={expected}")


def test_docs():
    import re
    package = (ROOT / "pyproject.toml").read_text(encoding="utf-8")
    changelog = (ROOT / "CHANGELOG.md").read_text(encoding="utf-8")
    version = re.search(r'^version = "([^"]+)"', package, re.M)
    release = re.search(r'^## (\d+\.\d+\.\d+)', changelog, re.M)
    check("docs: 包版本与最新变更记录一致",
          bool(version and release and version[1] == release[1]))
    skill = (ROOT / "SKILL.md").read_text(encoding="utf-8")
    refs = sorted(p.name for p in (ROOT / "references").glob("*.md"))
    check("docs: SKILL.md 只保留六节（身份/哲学/判断/Context/执行/QA）",
          all(k in skill for k in ("身份", "设计哲学", "判断", "Context", "执行", "QA")),
          str(refs))
    check("docs: references 只保留执行/资产契约，不维护第二套判断",
          len(refs) <= 2 and "judgment.md" not in refs, str(refs))
    check("docs: 文档里没有退役模块名与旧概念",
          all(b not in skill for b in ("guard.py", "normalize.py", "route.py", "qa.py",
                                       "design_direction")))
    _front = skill.split("---", 2)[1] if skill.startswith("---") else ""
    check("docs: SKILL 描述可被演示任务触发（不止框架口号）",
          all(k in _front for k in ("PPTX", "pitch deck", "路演", "汇报", "中文排版",
                                     "PowerPoint"))
          and len(_front) < 1200, _front[:120])


def main():
    # 测试自己的输出也是 UTF-8：Windows 上 `python scripts/selftest.py` 被管道捕获时
    # Python 用系统 ANSI 代码页写 stdout，第一行 `[PASS] …中文…` 就抛
    # UnicodeEncodeError，整轮测试以 1 结束——CI 的 Windows 作业曾经全红在这上面。
    for _stream in (sys.stdout, sys.stderr):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError, OSError):
            pass
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        for name, fn in (("intelligence", test_intelligence), ("verify", test_verify),
                         ("production", test_production), ("cli", test_cli), ("delivery", test_delivery), ("plan_boundaries", test_plan_boundaries),
                         ("release_evidence", test_release_evidence),
                         ("design_judgment", test_design_judgment),
                         ("cjk_rhythm_regression", test_cjk_rhythm_regression),
                         ("boundary_typography", test_boundary_typography),
                         ("ghost_observations", test_ghost_observations),
                         ("audit_boundaries", test_audit_boundaries),
                         ("delivery_safety", test_delivery_safety),
                         ("material_direction", test_material_direction),
                         ("deep_release", test_deep_release),
                         ("integrity_followup", test_integrity_followup),
                         ("atomic_publish", test_atomic_publish),
                         ("win_portability", test_win_portability),
                         ("prompt_cache", test_prompt_cache_key),
                         ("chart_target", test_chart_target),
                         ("char_spacing", test_char_spacing),
                         ("asset_budget", test_asset_budget),
                         ("asset_diagnosis", test_asset_diagnosis),
                         ("guard_diagnostics", test_guard_diagnostics),
                         ("root_order", test_root_order),
                         ("unit_visibility", test_unit_visibility),
                         ("art_direction_linkage", test_art_direction_linkage),
                         ("ghost_chart_color", test_ghost_chart_color_single_source),
                         ("case004_regressions", test_case004_regressions),
                         ("audit_fixes", test_audit_fixes),
                         ("design_intelligence_r6", test_design_intelligence_r6)):
            (root / name).mkdir(parents=True, exist_ok=True)
            fn(root / name)
        test_docs()
    print(f"\n{PASS}/{PASS + FAIL} passed")
    return 0 if FAIL == 0 else 1


def test_audit_fixes(tmp: Path):
    """第四轮审计：五处「判断/证据/产物」分叉的回归钉。

    每条钉都对应一个已实测的失效模式，不是为了覆盖率的形状测试：
      1. 文字墨迹按锚点落位——底部对齐的大字号数字曾压穿下一行说明而 guard 放行；
      2. 预览墨迹与 guard 同源——预览曾把同一段字画低 ~48px，两套证据互相打脸；
      3. 眉标不输出半截词/连词——曾产出「成本结构里真」「如果」；
      4. 作者标题占主张位——数据列举行曾顶掉作者写下的结论；
      5. 单序列数值默认可读——横向条把柱图旋转 90° 就丢数字；
      6. 声明主体后提示词不自相矛盾——「not a picture of a thing」曾与主体声明打架。
    """
    import intelligence as intel
    from primitives import ink_span, text_ink_height, default_show_values
    from verify import check_spec

    # ── 1 · 墨迹跨度：行距只在行间生效，末行按字形高度计 ──────────────
    one = text_ink_height(1, 150, 1.0)
    check("ink: 单行墨迹不因行距 1.0 缩到一字高（150px 字 ≥ 180px 墨迹）",
          one >= 150 * 1.2 - 0.01, f"{one}")
    two = text_ink_height(2, 50, 1.0)
    check("ink: 两行墨迹 = 行距 + 末行字形高",
          abs(two - (50 * 1.0 + 50 * 1.2)) < 0.01, f"{two}")
    top_span = ink_span(100, 200, 0, 180, "top")
    bottom_span = ink_span(100, 200, 0, 180, "bottom")
    check("ink: 底对齐把墨迹压到框底（top 锚点则从框顶起算）",
          abs(bottom_span[1] - 300) < 0.01 and abs(top_span[0] - 100) < 0.01,
          f"{top_span} {bottom_span}")

    # guard 必须按同一语义判碰撞：底对齐大字号数字压到下方说明行 → 必须拦。
    def _spec_with_text(els):
        return {"canvas": {"width": 1280, "height": 720},
                "theme": {"colors": {"background": "#FFFFFF", "ink": "#111111",
                                     "muted": "#777777", "primary": "#222222",
                                     "secondary": "#333333", "accent": "#AA0000"}},
                "slides": [{"id": "s1", "elements": els}]}

    big = {"id": "big", "type": "text", "role": "kpi_main", "x": 100, "y": 300,
           "width": 600, "height": 200, "size": 150, "line_height": 1.0,
           "anchor": "bottom", "max_lines": 1, "text": "2.4 亿", "color": "ink"}
    caption = {"id": "cap", "type": "text", "role": "evidence", "x": 100, "y": 486,
               "width": 600, "height": 30, "size": 16, "max_lines": 1,
               "text": "分三年投入 · 每期验收后再放款", "color": "muted"}
    hit = check_spec(_spec_with_text([dict(big), dict(caption)]))
    check("guard: 底对齐大字号压住下方说明行 → OVERLAP（旧模型按框高缩系数放行）",
          any(c["rule"] == "overlap" for c in hit["checks"]),
          str([c["rule"] for c in hit["checks"]])[:200])
    # 同样的几何改成顶对齐：墨迹在框内，不构成本例的重叠。
    safe = dict(big, anchor="top")
    ok = check_spec(_spec_with_text([safe, dict(caption)]))
    check("guard: 顶对齐同几何不误报（碰撞来自锚点，不是框重叠）",
          not [c for c in ok["checks"] if c["rule"] == "overlap"],
          str([c for c in ok["checks"] if c["rule"] == "overlap"])[:200])

    # ── 2 · 预览与 guard 同一墨迹口径 ────────────────────────────────
    import ghost
    from PIL import Image
    ctx = ghost.RenderContext({"theme": {"colors": {}},
                              "canvas": {"width": 1280, "height": 720}})
    for anchor in ("middle", "bottom"):
        el = dict(big, anchor=anchor, color="#F1EFEA")
        img = Image.new("RGB", (1280, 720), (23, 19, 16))
        ghost._draw_text(img, el, ctx, 1.0)
        rows = [y for y in range(img.height)
                if any(img.getpixel((x, y))[0] > 140 for x in range(96, 700, 2))]
        span = ink_span(300, 200, 0, text_ink_height(1, 150, 1.0), anchor)
        check(f"ghost: {anchor} 锚点绘制起点与 guard 墨迹跨度一致（Δ≤3px）",
              bool(rows) and abs(rows[0] - span[0]) <= 3,
              f"ghost={rows[:1]} guard={span[0]:.0f}")

    # ── 3 · 眉标：不输出半截词、悬空修饰语、条件连词 ─────────────────
    for text, want, why in (
            ("成本结构里真正失控的是二次搬运", "成本结构", "「真正」被「正」劈开不再产出「成本结构里真」"),
            ("单仓人效三年提升 38%，但增幅正在收窄", "", "数量残片「…三」「…升」不留在眉标"),
            ("只要门店密度达到每 10 万人 1.2 家，配送成本就能压到 8 元以内", "门店密度",
             "充分条件的主题在从句里，剥连词继续派生"),
            ("如果只改一半，回报会被摊薄到 5.8 年", "", "条件从句的连词不是主题"),
            ("请批准 2027 年 2.4 亿改造预算", "", "请求句没有可定位的名词短语"),
            ("改造一期只动了三个可复制的仓", "改造一期", "干净主语照常派生")):
        got = intel.section_label({}, {"text": text}, "prove")
        check(f"eyebrow: {why}", got == want, repr(got))

    # ── 4 · 作者标题占主张位（正文数据列举不夺位）────────────────────
    j = intel.extract_claim({"title": "改造一期只动了三个可复制的仓",
                             "content": "华东 11 仓中先改 3 仓，改造后单均成本下降 0.42 元，其余 8 仓未改造"})
    check("claim: 作者标题（动词+了）不被正文数据列举顶掉",
          j.get("text") == "改造一期只动了三个可复制的仓", str(j.get("text")))
    j2 = intel.extract_claim({"title": "门店的晨会节奏",
                              "content": "15 分钟站立会。（来源：观察）"})
    check("claim: 纯名词化话题仍让位给正文结论（不反向放宽）",
          j2.get("text") != "门店的晨会节奏", str(j2.get("text")))

    # ── 5 · 数值标签缺省口径唯一 ─────────────────────────────────────
    check("chart: 单序列横向条默认显示数值（与 column 不再分叉）",
          default_show_values("horizontal_bar",
                              {"categories": ["a", "b"], "series": [{"values": [1, 2]}]}))
    check("chart: 单序列折线默认显示数值",
          default_show_values("line", {"categories": ["1"], "series": [{"values": [3]}]}))
    check("chart: 多序列默认关闭（标签互相遮挡）",
          not default_show_values("column", {"categories": ["a"],
                                             "series": [{"values": [1]}, {"values": [2]}]}))
    check("chart: 声明 value_size 即视为要求看见数值",
          default_show_values("column", {"value_size": 14, "categories": ["a"],
                                         "series": [{"values": [1]}, {"values": [2]}]}))
    check("chart: 类别过多默认关闭",
          not default_show_values("bar", {"data": [{"label": str(i), "value": i}
                                                   for i in range(9)]}))

    # ── 6 · 声明主体后提示词不自相矛盾 ───────────────────────────────
    import assets
    card = assets.resolve_asset_card({"asset_id": "x", "asset_type": "background",
                                      "subject": "an automated sorting line at right",
                                      "safe_area": {"x": .06, "y": .08, "width": .34, "height": .78}})
    prompt = assets.build_asset_prompt(card, {"safe_area": card.get("safe_area")})["prompt"]
    check("asset: 声明主体的背景提示词不再要求「没有主角/不是一张东西的照片」",
          "not a picture of a thing" not in prompt
          and "no single dominant protagonist" not in prompt, prompt[:160])
    bare = assets.resolve_asset_card({"asset_id": "y", "asset_type": "background"})
    bare_prompt = assets.build_asset_prompt(bare, {})["prompt"]
    check("asset: 未声明主体时仍保留空场纪律（不丢判据）",
          "not a picture of a thing" in bare_prompt or "no dominant" in bare_prompt,
          bare_prompt[:160])


def test_design_intelligence_r6(tmp: Path):
    """第六轮审计：判断主数、收束落点与 progress 刻度的回归钉。

    每条都对应一份真实 brief 的实测失效（仓网投决 6 页稿）：
      1. 序列主数落在基期——「一路走低」不在变化词表里，形状判据本已成立；
      2. 构成主数被弱片段顶掉——页面主题词碰上指标名，把最大部分写成分母；
      3. 收束请求页被分期词画成流程——焦点、构图、权重、眉标四处互相打架；
      4. progress 的 max 只在产物生效——行级键在归一化被丢、预览硬编码 100。
    """
    import intelligence as intel
    from primitives import chart_rows, progress_ratio
    from verify import check_spec

    # ── 1 · 序列当前值靠时间形状，不靠主张措辞是否进了词表 ──────────
    _, low = intel.number_roles("2023 年 31%，2024 年 27%，2025 年 22%。",
                                "三年复购率一路走低")
    check("r6 series: 「一路走低」的当前值是末值 22%，不是基期 31%",
          low and low["raw"] == "22%", str(low))
    _, high = intel.number_roles("2023 年 68 元，2024 年 76 元，2025 年 83 元。",
                                 "客单价一路走高")
    check("r6 series: 「一路走高」的当前值是末值 83 元",
          high and high["raw"] == "83 元", str(high))

    # ── 2 · 算术分解：焦点是最大部分，总计是分母（与书写顺序无关）────
    for text in ("设备 3200 万元，产线改造 1100 万元，数字化 500 万元，合计 4800 万元。",
                 "数字化 500 万元，设备 3200 万元，产线改造 1100 万元，合计 4800 万元。"):
        _, main_part = intel.number_roles(text, "改造钱花在哪")
        check("r6 part_whole: 分解页主数是最大部分 3200 万元，不是弱片段命中的 1100",
              main_part and main_part["raw"] == "3200 万元", str(main_part))
    _nested = intel.number_roles(
        "采后损耗占 47%，其中运输 21%、分拣 14%、仓储 12%；田间损耗占 23%。",
        "采后损耗是温室草莓的最大损耗项")[1]
    check("r6 part_whole: 嵌套分解里有名字的构成项仍是主数（47% 不被拆成 21%）",
          _nested and _nested["raw"] == "47%", str(_nested))
    _u = intel.understand("改造钱花在哪 设备 3200 万元，产线改造 1100 万元，"
                          "数字化 500 万元，合计 4800 万元。")
    _anchor = intel.number_roles(
        "设备 3200 万元，产线改造 1100 万元，数字化 500 万元，合计 4800 万元。",
        "改造钱花在哪")
    _focus = intel.focus_of("prove", _u, {"text": "改造钱花在哪", "source": "extracted"},
                            number_anchor=_anchor, page_text="")
    check("r6 part_whole: 焦点理由把总计 4800 写成分母（不是次一个部分）",
          "4800 万元" in _focus["why"] and "3200 万元" in _focus["why"], _focus["why"])
    _weight = intel.information_weight(
        _u, {"text": "改造钱花在哪", "source": "extracted"},
        "设备 3200 万元，产线改造 1100 万元，数字化 500 万元，合计 4800 万元。",
        number_anchor=_anchor)
    check("r6 part_whole: 权重不把普通构成项写成总体/分母",
          not any("总体/分母" in w and not w.startswith("4800") for w in _weight["weaken"]),
          str(_weight["weaken"]))

    # ── 3 · 收束请求：落点是请求本身，不是支撑证据的项目分期 ─────────
    _brief = {"audience": "投资委员会", "decision": "批准仓网改造二期 1,800 万元"}
    _world = intel.derive_world(_brief, [])
    _page = intel.page_judgment(
        {"id": "s06", "title": "请委员会批准二期仓网改造",
         "content": "一期 3 仓跑通，单均成本降 0.42 元；二期需 1,800 万元。"},
        index=2, total=3, brief=_brief, world=_world, warnings=[])
    check("r6 closing: 请求页焦点是 decision_ask，不是 structure_map",
          _page["focus"]["element_role"] == "decision_ask",
          str(_page["focus"]["element_role"]))
    check("r6 closing: 请求页构图不被分期词拖成 linear_structure",
          _page["composition"]["chosen"] != "linear_structure",
          str(_page["composition"]["chosen"]))
    check("r6 closing: 请求页权重不把顺序形状列为必须最大化",
          not any("关系的形状" in m for m in _page["information_weight"]["maximize"]),
          str(_page["information_weight"]["maximize"]))
    check("r6 closing: 请求句眉标缺省（不输出被「会」劈开的「委员」）",
          intel.section_label({}, {"text": "请委员会批准二期仓网改造"}, "prove") == "",
          repr(intel.section_label({}, {"text": "请委员会批准二期仓网改造"}, "prove")))

    # ── 4 · progress 刻度：行级键不丢、产物与预览同一比例 ────────────
    _rows = chart_rows({"data": [{"label": "A", "value": 150, "max": 200},
                                  {"label": "B", "value": 50}]})
    check("r6 progress: chart_rows 保留行级 max（不再静默丢弃渲染键）",
          _rows[0].get("max") == 200, str(_rows[0]))
    check("r6 progress: progress_ratio 行级优先、元素级回落、超刻度夹紧",
          progress_ratio(_rows[0], {"max": 100}) == 0.75
          and progress_ratio(_rows[1], {"max": 200}) == 0.25
          and progress_ratio({"value": 250}, {"max": 200}) == 1.0)
    _spec = {"canvas": {"width": 1280, "height": 720},
             "theme": {"colors": {"background": "#FFFFFF", "ink": "#111111",
                                  "muted": "#777777", "primary": "#222222",
                                  "secondary": "#333333", "accent": "#AA0000"}},
             "slides": [{"id": "s1", "elements": [
                 {"id": "c1", "type": "chart", "role": "evidence",
                  "x": 80, "y": 100, "width": 1000, "height": 400,
                  "chart_kind": "progress_bar", "max": 200,
                  "source": "s", "unit": "%", "period": "p", "basis": "b",
                  "data": [{"label": "A", "value": 150},
                           {"label": "B", "value": 50}]}]}]}
    _bad = dict(_spec["slides"][0]["elements"][0])
    _bad["data"] = [{"label": "A", "value": 10, "max": -1}]
    _bad_spec = {**_spec, "slides": [{"id": "s1", "elements": [_bad]}]}
    check("r6 progress: 行级 max 非正数被 guard 拦下（不静默回落）",
          any(c["rule"] == "data_integrity" for c in check_spec(_bad_spec)["checks"]
              if c.get("level") == "error"),
          str([c["rule"] for c in check_spec(_bad_spec)["checks"]]))
    import compiler
    from pptx import Presentation
    _out = tmp / "progress.pptx"
    _rep = compiler.compile_deck(_spec, str(_out))
    _shapes = {s.name: s for s in Presentation(str(_out)).slides[0].shapes} \
        if _rep.get("passed") else {}
    _fill_ratio = (_shapes["c1__fill_0"].width / _shapes["c1__track_0"].width) \
        if _shapes else 0
    check("r6 progress: 产物填充 = 150/200 = 75% 轨道宽",
          _rep.get("passed") and abs(_fill_ratio - 0.75) < 0.02,
          f"ratio={_fill_ratio}")
    import ghost as _ghost
    _img = _ghost.ghost_page(_spec["slides"][0], _spec, scale=1,
                             supersample=1).convert("RGB")
    # 第一行中心 y≈元素 y+行高/2；扫这一行上强调色像素的最右端对轨道的相对位置
    _cy = 100 + 100
    _accent_xs = [x for x in range(_img.width)
                  if (lambda c: c[0] > 150 and c[1] < 60 and c[2] < 60)(
                      tuple(_img.getpixel((x, _cy))))]
    _track_left = 80 + int(1000 * 0.24) + 24
    _track_right = 80 + 1000 - 78
    _ghost_ratio = (max(_accent_xs) - _track_left) / max(_track_right - _track_left, 1) \
        if _accent_xs else 0
    check("r6 progress: 预览填充与产物同一刻度（75%，不冲出轨道到 150%）",
          bool(_accent_xs) and abs(_ghost_ratio - 0.75) < 0.05,
          f"ghost_ratio={_ghost_ratio:.3f}")


if __name__ == "__main__":
    raise SystemExit(main())
