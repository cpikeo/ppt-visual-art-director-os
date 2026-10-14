"""8-page same-input regression and A/B timing; no Office or aesthetic score.
Reuses the previous font fixture/CLI/audit helpers. Run without concurrent selftests:
python evaluation/relationship_labels.py --before /path/to/9.15.1 --after . --out /fresh/path --repeat 5
"""
from pathlib import Path
import argparse
import json
import statistics
import hashlib
import platform
from importlib.metadata import version
import subprocess
import sys
from font_fidelity.reproduce import fixture, spec, call, audit, write

GOVERNANCE = {"id":"governance", "title":"决策、预算与执行各有归属",
              "content":"董事会决定战略；管理层负责预算；区域负责人协调资源；当地负责人执行计划。",
              "asset":"none"}


def build_spec(titles, contents, aid):
    deck=spec(titles,contents,aid)
    # The same author-built specimen is used on both versions: no adaptive templates.
    for slide in deck["slides"]:
        for e in slide["elements"]:
            if e.get("id")=="chart":
                e.update(label_size=24,value_size=28)
    def text(eid,value,x,y,w,h,size=28,**kw):
        return dict(id=eid,type="text",text=value,x=x,y=y,width=w,height=h,size=size,
                    color="ink",role="body",line_height=1.5,**kw)
    elements=[text("title",GOVERNANCE["title"],48,88,1184,112,44,bold=True),
              text("qualifier","职责分工不等于汇报层级；未声明的关系不补画。",48,216,1184,40,22),
              text("source","合成治理样本 · 只核对原句职责 / 非客户组织事实",48,664,1184,40,18),
              text("folio","04 / 08",1096,32,136,40,18)]
    elements[0]["role"]="title"; elements[2]["role"]="source"
    for n,(owner,duty) in enumerate((("董事会","决定战略"),("管理层","负责预算"),
                                    ("区域负责人","协调资源"),("当地负责人","执行计划"))):
        elements.extend([text(f"owner{n}",owner,64,288+n*76,344,52),
                         text(f"duty{n}",duty,456,288+n*76,680,52)])
    deck["slides"].insert(3,{"id":"governance", "page_intent":{"insight":GOVERNANCE["title"],"focus":"title"},
                            "source_zone":{"x":48,"y":660,"width":1184,"height":48},"elements":elements})
    for n,slide in enumerate(deck["slides"],1):
        slide["id"]=f"s{n:02}"
        for e in slide["elements"]:
            if e["id"]=="folio": e["text"]=f"{n:02} / 08"
    return deck


def chart_audit(path):
    from pptx import Presentation
    found=[]
    for slide in Presentation(path).slides:
        for shape in slide.shapes:
            if not shape.has_chart: continue
            c=shape.chart; labels=c.plots[0].has_data_labels
            font=c.category_axis.tick_labels.font
            found.append({"show_values":bool(labels and c.plots[0].data_labels.show_value),
                          "category_pt":font.size.pt if font.size else None,
                          "category_ea":[n.get("typeface") for n in font._element.xpath("a:ea")],
                          "value_pt":c.plots[0].data_labels.font.size.pt if labels else None,
                          "values":[list(s.values) for s in c.series]})
    return found


def main():
    parser=argparse.ArgumentParser()
    for key in ("before","after","out"): parser.add_argument("--"+key,type=Path,required=True)
    parser.add_argument("--repeat",type=int,default=5); a=parser.parse_args()
    if a.repeat < 2: raise SystemExit("repeat >= 2 keeps before-0 intact for visual comparison")
    root=a.out.resolve();root.mkdir(parents=True,exist_ok=True)
    if (root/"runs").exists(): raise SystemExit("Use a fresh output directory, not warm caches.")
    write(root/"environment.json",{"platform":platform.platform(),"python":sys.version,
        "dependencies":{name:version(name) for name in ("python-pptx","Pillow","numpy","PyYAML")},
        "repeat":a.repeat,"method":"alternating A/B, new processes/application caches; OS caches not flushed; no Office/model/author timing",
        "sources":{label:{str(f.relative_to(repo)):hashlib.sha256(f.read_bytes()).hexdigest()
                           for f in list((repo/"scripts").glob("*.py"))+[repo/"pyproject.toml",repo/"SKILL.md"]}
                   for label,repo in (("before",a.before.resolve()),("after",a.after.resolve()))}})
    titles,contents=fixture(root)
    brief=json.loads((root/"brief.json").read_text());brief["slides"].insert(3,GOVERNANCE)
    for n,item in enumerate(brief["slides"],1): item["id"]=f"s{n:02}"
    write(root/"brief.json",brief)
    cases=Path(__file__).with_name("relationship-cases.json").resolve()
    probe='''import sys,json; sys.path.insert(0,sys.argv[1]+"/scripts")
import intelligence as i
cases=json.load(open(sys.argv[2],encoding="utf-8")); result=[]
for case in cases:
 j=i.page_judgment({**case,"asset":"none"},index=2,total=7,brief={},world={},warnings=[])
 result.append({"case":case,"judgment":j})
print(json.dumps(result,ensure_ascii=False))'''
    rows=[]
    for label in ("before","after"):
        p=subprocess.run([sys.executable,"-c",probe,str(getattr(a,label).resolve()),str(cases)],
                         capture_output=True,text=True,encoding="utf-8",check=True)
        write(root/f"{label}-judgments.json",json.loads(p.stdout))
    for i in range(a.repeat):
        for label in (("before","after") if i%2==0 else ("after","before")):
            repo=getattr(a,label).resolve();d=root/"runs"/f"{label}-{i}";d.mkdir(parents=True)
            plan=call(repo,["plan",root/"brief.json","--out",d/"plan.json","--skeleton",d/"skeleton.py","--json"],d/"plan-call.json")
            manifest=d/"asset_manifest.json";aid=json.loads(manifest.read_text())["assets"][0]["asset_id"]
            write(d/"build.json",build_spec(titles,contents,aid))
            args=["check",d/"build.json",d/"deck.pptx","--mode","release","--speed","strict","--assets-manifest",manifest,"--json"]
            cold=call(repo,args,d/"cold-call.json");warm=call(repo,args,d/"warm-call.json")
            rows.append({"label":label,"iteration":i,"plan":plan,"cold":cold,"warm":warm,
                         "automation_wall_ms":plan["wall_ms"]+cold["wall_ms"]})
            write(root/"measurements.json",rows)
            if i==0:
                write(d/"pptx-audit.json",audit(d/"deck.pptx"))
                write(d/"chart-audit.json",chart_audit(d/"deck.pptx"))
    summary={}
    for label in ("before","after"):
        batch=[r for r in rows if r["label"]==label]
        summary[label]={"automation_wall_ms":statistics.median(r["automation_wall_ms"] for r in batch)}
        for stage in ("plan","cold","warm"):
            summary[label][stage]={"wall_ms":statistics.median(r[stage]["wall_ms"] for r in batch),
                "timing":{k:statistics.median(r[stage]["timing"].get(k,0) for r in batch)
                          for k,v in batch[0][stage]["timing"].items() if isinstance(v,(int,float)) and not isinstance(v,bool)}}
    write(root/"summary.json",summary)
    # Compiler changed: old native chart without labels must not survive an upgrade cache hit.
    d=root/"runs"/f"before-{a.repeat-1}";args=["check",d/"build.json",d/"deck.pptx","--mode","release",
        "--speed","strict","--assets-manifest",d/"asset_manifest.json","--json"]
    upgraded=call(a.after.resolve(),args,root/"upgrade-call.json")["timing"]
    assert upgraded["compile_ms"]>0 and upgraded["render_pages_drawn"]==8
    assert chart_audit(d/"deck.pptx")[0]["show_values"]
    rewarmed=call(a.after.resolve(),args,root/"rewarm-call.json")["timing"]
    assert rewarmed["compile_ms"]==0 and rewarmed["render_pages_cached"]==8
    write(root/"cache-assertions.json",{"upgrade_invalidates_compiler_and_preview":True,"warm_reuses_all":True,
        "note":"last before-N output intentionally upgraded AFTER timed measurement; before-0 audit is the original when repeat>1"})
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=="__main__": main()
