"""Offline audit evidence, not a second production entry. All production goes through vao.py.
Usage: python audit/reproduce.py --before /path/to/baseline --rounds 5
"""
from __future__ import annotations
import argparse
import copy
import hashlib
import json
import os
import platform
from pathlib import Path
import runpy
import shutil
import statistics
import subprocess
import sys
import time
import zipfile

ROOT = Path(__file__).resolve().parents[1]


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fixtures(work, count, images=False):
    from PIL import Image, ImageDraw
    scope = runpy.run_path(str(ROOT / 'acceptance/build.py'))
    deck = scope['SPEC']
    deck['slides'] = [dict(copy.deepcopy(deck['slides'][i % 7]), id=f's{i+1:02d}') for i in range(count)]
    manifest = None
    if images:
        image = work / 'fixture.png'
        # Original deterministic illustration; not a photograph, measurement or generated witness.
        pic = Image.new('RGB', (1920, 1080), (244, 239, 230))
        draw = ImageDraw.Draw(pic)
        draw.ellipse((1000, 150, 1750, 900), fill=(120, 90, 60))
        draw.line((1100, 850, 1650, 220), fill=(180, 160, 130), width=8)
        pic.save(image)
        ids = []
        for i, page in enumerate(deck['slides']):
            if i != 1:
                continue
            ids.append(page['id'])
            page['elements'] = [e for e in page['elements'] if e.get('role') in ('eyebrow', 'page_number', 'source')]
            page['elements'] += [
                scope['text']('title', '图像是插图，\n不是事实证词。', 64, 176, 448, 152, 40, bold=True),
                scope['text']('caption', '原始合成测试图\n用于资产绑定与裁切验证', 64, 400, 448, 100, 27),
                dict(id='image', type='image', asset_id='fixture', x=560, y=180,
                     width=640, height=360, fit='cover')]
            page['page_intent'] = dict(insight='测试资产链', focus='title')
        manifest = work / 'assets.json'

    brief = dict(audience='仓库维护者', decision='核验结构而非签署视觉', tension='同任务公平测量', slides=[
        dict(id=p['id'], title=p['page_intent']['insight'], content=' '.join(e.get('text', '') for e in p['elements']), asset='none')
        for p in deck['slides']])
    if images:
        for p in brief['slides']:
            if p['id'] in ids:
                p.update(asset='required', asset_role='background', asset_source=dict(
                    kind='original', path=str(image), source='Original synthetic audit illustration; not evidence'))
    write(work / 'brief.json', brief)
    write(work / 'spec.json', deck)
    return deck, manifest


def cli(repo, args, log):
    cmd = [sys.executable, str(repo / 'scripts/vao.py'), *map(str, args)]
    start = time.perf_counter()
    p = subprocess.run(cmd, cwd=ROOT, capture_output=True, encoding='utf-8', env={**os.environ, 'PYTHONDONTWRITEBYTECODE':'1'})
    elapsed = (time.perf_counter() - start) * 1000
    log.write_text(p.stdout + p.stderr, encoding='utf-8')
    if p.returncode:
        raise RuntimeError(f'{cmd} returned {p.returncode}: {log}')
    return elapsed, json.loads(p.stdout)


def inspect_pptx(path, expected_count, report_path):
    # Native object read/edit/save/reopen with python-pptx is NOT a PowerPoint UI editing test.
    from pptx import Presentation
    prs = Presentation(path)
    charts = []; names = []; native_texts = []
    for page in prs.slides:
        current = []
        for shape in page.shapes:
            current.append(shape.name)
            if shape.has_text_frame:
                native_texts.append(shape.text)
            if shape.has_chart:
                charts.append(dict(name=shape.name, values=[list(s.values) for s in shape.chart.series],
                                   categories=[c.label for c in shape.chart.plots[0].categories]))
        assert len(current) == len(set(current)), current
        names += current
    assert len(prs.slides) == expected_count
    assert charts and charts[0]['values'] == [[7.0, 12.0]], charts
    assert charts[0]['categories'] == ['修改前', '修改后']
    assert any('186\u00a0公里' in t for t in native_texts)
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        workbooks = [n for n in z.namelist() if n.startswith('ppt/embeddings/')]
        assert workbooks
    first = next(s for s in prs.slides[0].shapes if s.has_text_frame and '架构简单' in s.text)
    first.text = '原生编辑结构探针'
    edited = path.with_name('edit-roundtrip.pptx'); prs.save(edited)
    assert any('原生编辑结构探针' in s.text for s in Presentation(edited).slides[0].shapes if s.has_text_frame)
    write(report_path, dict(slides=len(prs.slides), shapes=len(names), text_frames=len(native_texts),
          charts=charts, embedded_workbooks=workbooks, unique_names_per_slide=True,
          python_pptx_read_edit_save_reopen='passed', target_powerpoint='not_verified',
          native_editing_ui='not_verified', visual_review='not_performed'))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--before', type=Path, required=True)
    parser.add_argument('--rounds', type=int, default=5)
    args = parser.parse_args()
    repos = dict(before=args.before.resolve(), after=ROOT)
    evidence = ROOT / 'audit/performance'
    evidence.mkdir(exist_ok=True)
    records = []
    # Equal bytecode state: do not compare baseline source compilation against patched .pyc.
    for repo in repos.values():
        subprocess.run([sys.executable, '-m', 'compileall', '-q', str(repo/'scripts')], check=True)
    # Five interleaved paired rounds, reversing order each round. Cold means empty app cache,
    # not flushed OS page cache. Each command has a fresh interpreter, so imports are charged.
    for task, count, images in [('text5',5,False), ('text15',15,False), ('text30',30,False), ('image15',15,True)]:
        for rnd in range(args.rounds):
            for version in (('before','after') if rnd % 2 == 0 else ('after','before')):
                work = evidence / f'{task}-{rnd}-{version}'
                if work.exists(): shutil.rmtree(work)
                work.mkdir()
                deck, manifest = fixtures(work, count, images)
                plan_args = ['plan', work/'brief.json', '--out',work/'plan.json', '--json']
                if manifest: plan_args += ['--assets-out',manifest]
                plan_ms, _ = cli(repos[version], plan_args, work/'plan.log')
                if manifest:
                    entries = json.loads(manifest.read_text())['assets']
                    for p in deck['slides']:
                        for e in p['elements']:
                            if e['type'] == 'image':
                                e['asset_id'] = next(a['asset_id'] for a in entries if p['id'] in a.get('slide_ids', []))
                    write(work/'spec.json', deck)
                out = work / 'out.pptx'
                args_check = ['check',work/'spec.json',out,'--mode','release','--speed','strict','--json']
                if manifest: args_check += ['--assets-manifest',manifest]
                cold_hash = None
                for phase in ('cold','warm','iteration'):
                    if phase == 'iteration':
                        deck['slides'][0]['elements'][2]['text'] = deck['slides'][0]['elements'][2]['text'].replace('架构简单', '执行简单')
                        write(work/'spec.json', deck)
                    wall, result = cli(repos[version],args_check,work/f'{phase}.log')
                    # CLI JSON uses verdict + timing; require structural PASS and all pages.
                    man = json.loads(out.with_suffix('.manifest.json').read_text())
                    assert man['status'] == 'PASS', man
                    ghost = man['ghost_preview']
                    assert ghost['scope'] == 'full' and ghost['count'] == count
                    digest = sha(out)
                    if phase == 'cold': cold_hash = digest
                    if phase == 'warm': assert digest == cold_hash
                    records.append(dict(task=task,round=rnd,version=version,phase=phase,
                        wall_ms=wall,plan_ms=plan_ms,plan_plus_check_ms=plan_ms+wall,
                        timing=result.get('timing'),output_sha256=digest,
                        preview_sha256=sha(Path(ghost['contact_sheet']))))
                print(f'{task} round {rnd+1} {version} PASS', flush=True)
    summaries = []
    for task in sorted({r['task'] for r in records}):
        for phase in ('cold','warm','iteration'):
            row = dict(task=task, phase=phase)
            for version in repos:
                rows = [r for r in records if (r['task'],r['phase'],r['version']) == (task,phase,version)]
                row[version] = dict(median_ms=round(statistics.median(r['wall_ms'] for r in rows),2),
                    plan_median_ms=round(statistics.median(r['plan_ms'] for r in rows),2),
                    plan_plus_check_median_ms=round(statistics.median(r['plan_plus_check_ms'] for r in rows),2))
            row['check_delta_percent'] = round((row['after']['median_ms']/row['before']['median_ms']-1)*100,2)
            summaries.append(row)
    # Artifact equality pairs on same fixture; image file paths differ but bytes must agree.
    pairs = []
    for r in records:
        if r['version'] != 'before': continue
        other = next(v for v in records if (v['task'],v['round'],v['phase'],v['version']) == (r['task'],r['round'],r['phase'],'after'))
        pairs.append(dict(task=r['task'],round=r['round'],phase=r['phase'],
                          pptx_equal=r['output_sha256']==other['output_sha256'],
                          preview_equal=r['preview_sha256']==other['preview_sha256']))
    report = dict(environment=dict(python=platform.python_version(), platform=platform.platform(),
                    cpu=os.cpu_count(), rounds=args.rounds, cache='application cold; OS cache not flushed; both versions precompiled equally',
                    speed='strict, all pages, fresh process per command',
                    scope='deterministic authoring fixtures; excludes human/model authoring and external asset generation'),
                  summaries=summaries, equality=pairs, records=records)
    write(ROOT/'audit/performance.json',report)
    assert all(p['pptx_equal'] and p['preview_equal'] for p in pairs)
    assert all(r['timing']['render_pages_drawn'] == 1 for r in records if r['phase'] == 'iteration')
    inspect_pptx(ROOT/'acceptance/out.pptx',7,ROOT/'audit/pptx-structure.json')
    print(json.dumps(summaries,indent=2))


if __name__ == '__main__': main()
